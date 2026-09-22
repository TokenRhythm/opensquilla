"""Offline checks for the opt-in deliverable acceptance harness."""

from __future__ import annotations

import contextlib
import hashlib
import io
import json
import ntpath
import os
import runpy
import subprocess
import sys
import tomllib
from collections.abc import Iterator
from pathlib import Path
from zipfile import BadZipFile

import httpx
import pytest
from pptx import Presentation
from pypdf.errors import PdfReadError
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.pdfgen import canvas

from opensquilla.artifacts import ArtifactStore
from scripts import live_deliverable_acceptance as acceptance
from scripts.live_deliverable_acceptance import (
    MODEL,
    BoundedRelay,
    listed_artifacts,
    load_report,
    restored_sessions,
    validate_bytes,
)
from scripts.live_tokenrhythm_budget import BudgetRejectedError, FunctionalRequestLog

_REPLY = b'{"choices":[{"message":{"content":"offline response"}}]}'


@pytest.fixture
def offline_relay(
    tmp_path: Path,
) -> Iterator[tuple[BoundedRelay, FunctionalRequestLog, list[httpx.Request]]]:
    request_log = FunctionalRequestLog(tmp_path / "requests.sqlite3", enabled=True)
    request_log.select_phase(variant="new", case_id="offline-deliverable")
    sent: list[httpx.Request] = []

    def upstream(request: httpx.Request) -> httpx.Response:
        sent.append(request)
        return httpx.Response(
            200, content=_REPLY, headers={"Content-Type": "application/json"}
        )

    # Exercise the real relay and SQLite recorder without starting a server or
    # creating an HTTP transport capable of reaching a provider.
    relay = BoundedRelay(
        None,
        {},
        api_key="synthetic-offline-secret",
        request_log=request_log,
        transport=httpx.MockTransport(upstream),
    )
    try:
        yield relay, request_log, sent
    finally:
        relay.close()


def _body(**fields: object) -> bytes:
    return json.dumps(
        {"model": MODEL, "messages": [{"role": "user", "content": "offline"}], **fields}
    ).encode()


def _forward(relay: BoundedRelay, body: bytes) -> None:
    with relay.forward(body) as response:
        assert response.status_code == 200
        assert b"".join(response.chunks) == _REPLY


def _assert_rejected(offline_relay, body: bytes, reason: str) -> None:
    relay, request_log, sent = offline_relay
    calls_before = relay.calls
    sent_before = list(sent)
    log_before = request_log.snapshot()
    with pytest.raises(BudgetRejectedError, match=f"^{reason}$"):
        with relay.forward(body):
            pytest.fail("rejected request yielded an upstream response")
    assert relay.calls == calls_before
    assert sent == sent_before
    assert request_log.snapshot() == log_before


@pytest.mark.parametrize("field", ["max_tokens", "max_completion_tokens"])
@pytest.mark.parametrize("limit", [1, 8192])
def test_relay_accepts_each_output_limit_boundary(offline_relay, field, limit) -> None:
    relay, request_log, sent = offline_relay
    body = _body(**{field: limit}, n=1)

    _forward(relay, body)

    assert relay.calls == 1
    assert len(sent) == 1
    assert sent[0].method == "POST"
    assert sent[0].url.path == "/v1/chat/completions"
    assert sent[0].content == body
    rows = request_log.snapshot()["requests"]
    assert len(rows) == 1
    assert rows[0]["model"] == MODEL
    assert rows[0]["status"] == "completed"
    assert rows[0]["request_bytes"] == len(body)
    assert rows[0]["response_bytes"] == len(_REPLY)


@pytest.mark.parametrize("model", [None, "", "different-model", MODEL.upper(), f" {MODEL}"])
def test_relay_rejects_wrong_model_before_dispatch(offline_relay, model) -> None:
    _assert_rejected(
        offline_relay, _body(model=model, max_tokens=1), "acceptance_request_limit"
    )


def test_relay_rejects_missing_model_before_dispatch(offline_relay) -> None:
    _assert_rejected(
        offline_relay,
        json.dumps({"messages": [], "max_tokens": 1}).encode(),
        "acceptance_request_limit",
    )


def test_relay_rejects_missing_output_limit_before_dispatch(offline_relay) -> None:
    _assert_rejected(offline_relay, _body(), "acceptance_output_limit")


@pytest.mark.parametrize("field", ["max_tokens", "max_completion_tokens"])
@pytest.mark.parametrize(
    "limit",
    [None, 0, -1, 8193, "1", "8192", 1.0, [], {}, True, False],
    ids=["null", "zero", "negative", "over-limit", "string-one", "string-max",
         "float", "array", "object", "true", "false"],
)
def test_relay_rejects_invalid_output_limits_before_dispatch(offline_relay, field, limit) -> None:
    _assert_rejected(offline_relay, _body(**{field: limit}), "acceptance_output_limit")


@pytest.mark.parametrize(
    ("max_tokens", "max_completion_tokens"),
    [(1, 8193), (8193, 1), (1, 2)],
)
def test_relay_rejects_conflicting_output_limits_before_dispatch(
    offline_relay, max_tokens, max_completion_tokens
) -> None:
    _assert_rejected(
        offline_relay,
        _body(max_tokens=max_tokens, max_completion_tokens=max_completion_tokens),
        "acceptance_output_limit",
    )


@pytest.mark.parametrize("limit", [1, 8192])
def test_relay_accepts_equal_valid_output_limits(offline_relay, limit) -> None:
    relay, request_log, sent = offline_relay
    body = _body(max_tokens=limit, max_completion_tokens=limit)

    _forward(relay, body)

    assert relay.calls == 1
    assert len(sent) == 1
    assert sent[0].content == body
    assert len(request_log.snapshot()["requests"]) == 1


@pytest.mark.parametrize(
    "completions",
    [None, 0, -1, 2, 8192, "1", 1.0, True, False],
    ids=["null", "zero", "negative", "two", "many", "string", "float", "true", "false"],
)
def test_relay_rejects_non_single_completion_before_dispatch(offline_relay, completions) -> None:
    _assert_rejected(
        offline_relay, _body(max_tokens=8192, n=completions), "acceptance_output_limit"
    )


def test_relay_allows_exactly_sixty_requests_then_rejects_without_counting(offline_relay) -> None:
    relay, request_log, sent = offline_relay
    body = _body(max_tokens=1)

    for expected_calls in range(1, 61):
        _forward(relay, body)
        assert relay.calls == expected_calls
        assert len(sent) == expected_calls

    rows = request_log.snapshot()["requests"]
    assert len(rows) == 60
    assert all(row["status"] == "completed" for row in rows)
    for _ in range(2):
        _assert_rejected(offline_relay, body, "model_call_limit_exhausted")


@pytest.fixture
def relay_at_last_counter_slot(offline_relay):
    relay, _, _ = offline_relay
    # Simulate earlier dispatches to test this independent guard without lifting
    # the real ledger's hard 60-call limit or making additional requests.
    relay.calls = 89
    return offline_relay


def test_relay_ninety_call_guard_rejects_before_ledger_and_dispatch(
    relay_at_last_counter_slot,
) -> None:
    relay, request_log, sent = relay_at_last_counter_slot
    body = _body(max_tokens=1)

    _forward(relay, body)

    assert relay.calls == 90
    assert len(sent) == 1
    assert len(request_log.snapshot()["requests"]) == 1
    for _ in range(2):
        _assert_rejected(relay_at_last_counter_slot, body, "acceptance_request_limit")


def test_relay_does_not_count_a_rejected_in_flight_reservation(offline_relay) -> None:
    _, request_log, _ = offline_relay
    request_log.start_request(model=MODEL, request_bytes=1)

    _assert_rejected(offline_relay, _body(max_tokens=1), "request_already_in_flight")


def test_relay_counts_dispatched_request_when_consumer_rejects(offline_relay) -> None:
    relay, request_log, sent = offline_relay

    with pytest.raises(BudgetRejectedError, match="^synthetic_consumer_rejection$"):
        with relay.forward(_body(max_tokens=1)):
            raise BudgetRejectedError("synthetic_consumer_rejection")

    assert relay.calls == 1
    assert len(sent) == 1
    rows = request_log.snapshot()["requests"]
    assert len(rows) == 1
    assert rows[0]["status"] == "interrupted"


def test_relay_counts_reserved_dispatch_when_transport_fails(offline_relay, monkeypatch) -> None:
    relay, request_log, sent = offline_relay

    def fail_transport(*args, **kwargs):
        raise httpx.ConnectError("synthetic offline transport failure")

    monkeypatch.setattr(relay._client, "stream", fail_transport)
    with pytest.raises(httpx.ConnectError, match="synthetic offline transport failure"):
        _forward(relay, _body(max_tokens=1))

    assert relay.calls == 1
    assert sent == []
    rows = request_log.snapshot()["requests"]
    assert len(rows) == 1
    assert rows[0]["status"] == "interrupted"
    assert rows[0]["reason"] == "transport_error"


def _pdf_bytes() -> bytes:
    buffer = io.BytesIO()
    document = canvas.Canvas(buffer)
    document.drawString(72, 720, "Week one: interfaces completed")
    document.showPage()
    document.drawString(72, 720, "Week two: staged release")
    document.showPage()
    document.save()
    return buffer.getvalue()


def test_validate_pdf_counts_pages_and_extracts_text() -> None:
    result = validate_bytes(_pdf_bytes(), "weekly-report.PDF")

    assert result["pages"] == 2
    assert "Week one: interfaces completed" in result["text"]
    assert "Week two: staged release" in result["text"]


def _pptx_bytes() -> bytes:
    document = Presentation()
    for title in ("Coffee shop goals", "Budget: equipment 20000"):
        slide = document.slides.add_slide(document.slide_layouts[0])
        slide.shapes.title.text = title
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_validate_pptx_counts_slides_and_extracts_text() -> None:
    result = validate_bytes(_pptx_bytes(), "opening-plan.PPTX")

    assert result["slides"] == 2
    assert "Coffee shop goals" in result["text"]
    assert "Budget: equipment 20000" in result["text"]


@pytest.mark.parametrize("encoding", ["utf-8", "utf-8-sig"])
@pytest.mark.parametrize(
    ("name", "content"),
    [
        ("onboarding.md", "# 入职清单\n\n- [ ] 开通账号\n\n| 账号 | 负责人 |\n| --- | --- |\n"),
        ("index.html", "<!doctype html><html><body><h1>北京四季</h1></body></html>\n"),
        ("expenses.csv", "date,category,amount\n2026-09-01,咖啡,28\n"),
    ],
)
def test_validate_text_deliverables_preserves_utf8_content(name, content, encoding) -> None:
    assert validate_bytes(content.encode(encoding), name) == {"text": content}


@pytest.mark.parametrize("name", ["broken.md", "broken.html", "broken.csv"])
def test_validate_text_deliverables_rejects_invalid_utf8(name) -> None:
    with pytest.raises(UnicodeDecodeError):
        validate_bytes(b"\xff\xfeinvalid text", name)


@pytest.mark.parametrize("data", [b"not a PDF", b"%PDF-1.7\ntruncated"])
def test_validate_pdf_rejects_corrupt_data(data) -> None:
    with pytest.raises(PdfReadError):
        validate_bytes(data, "broken.pdf")


@pytest.mark.parametrize("data", [b"not a presentation", b"PK\x03\x04truncated"])
def test_validate_pptx_rejects_corrupt_data(data) -> None:
    with pytest.raises(BadZipFile):
        validate_bytes(data, "broken.pptx")


def test_listed_artifacts_reads_all_pages_and_uses_visibility_not_source(tmp_path: Path) -> None:
    store = ArtifactStore(tmp_path)
    session_id = "deliverable-session"
    session_key = "agent:main:webchat:deliverable-session"
    materials = [
        ("weekly-report.pdf", "application/pdf", _pdf_bytes(), "publish_artifact"),
        (
            "opening-plan.pptx",
            "application/vnd.openxmlformats-officedocument.presentationml.presentation",
            _pptx_bytes(),
            "workspace-preview",
        ),
        ("onboarding.md", "text/markdown", b"# Onboarding\n", "working_files"),
        ("expenses.csv", "text/csv", b"category,amount\ncoffee,28\n", "publish_artifact"),
    ]
    visible = [
        store.publish_bytes(
            data,
            session_id=session_id,
            session_key=session_key,
            name=name,
            mime=mime,
            source=source,
        )
        for name, mime, data, source in materials
    ]
    internal = [
        store.publish_bytes(
            b"<h1>Internal revision</h1>",
            session_id=session_id,
            session_key=session_key,
            name=f"{source}.html",
            mime="text/html",
            source=source,
            visibility="internal",
        )
        for source in ("workspace-preview", "working_files", "publish_artifact")
    ]
    other_session = store.publish_bytes(
        b"# Private to another session\n",
        session_id="other-session",
        session_key="agent:main:webchat:other-session",
        name="other.md",
        mime="text/markdown",
        source="publish_artifact",
    )

    first_page = store.list_refs(session_id=session_id, limit=2)
    assert len(first_page.refs) == 2
    assert first_page.has_more is True
    result = listed_artifacts(tmp_path, session_id, session_key, page_size=2)

    assert result == [
        ref.to_dict() for ref in sorted(visible, key=lambda ref: (ref.created_at, ref.id))
    ]
    assert {row["name"] for row in result} == {
        "weekly-report.pdf", "opening-plan.pptx", "onboarding.md", "expenses.csv"
    }
    assert {row["source"] for row in result} == {
        "publish_artifact", "workspace-preview", "working_files"
    }
    assert not {ref.id for ref in internal}.intersection(row["id"] for row in result)
    assert other_session.id not in {row["id"] for row in result}


def test_load_report_resume_preserves_cases_and_archives_first_failure(tmp_path: Path) -> None:
    cases = [
        {
            "id": "html-create",
            "session_key": "agent:main:webchat:html-session",
            "failures": [],
            "downloads": [],
        }
    ]
    rejection = {
        "case_id": "html-edit",
        "status": 409,
        "response": {"error": "request already running"},
    }
    original = {
        "model": MODEL,
        "cases": cases,
        "source": "/original/checkout",
        "source_commit": "a" * 40,
        "provider_calls": 12,
        "provider_probe_status": 200,
        "runner_error": "HTTP 409 while submitting html-edit",
        "http_rejection": rejection,
        "blocked": "gateway_rejected_request",
    }
    path = tmp_path / "report.json"
    path.write_text(json.dumps(original), encoding="utf-8")
    bytes_before = path.read_bytes()

    report = load_report(tmp_path, resume=True)

    assert report["cases"] == cases
    assert report["model"] == MODEL
    assert len(report["attempts"]) == 1
    archived = report["attempts"][0]
    assert archived["runner_error"] == original["runner_error"]
    assert archived["http_rejection"] == rejection
    assert archived["blocked"] == original["blocked"]
    assert archived["case_ids"] == ["html-create"]
    assert archived["source"] == original["source"]
    assert archived["source_commit"] == original["source_commit"]
    assert archived["provider_calls"] == 12
    assert archived["provider_probe_status"] == 200
    assert isinstance(archived["archived_at"], float)
    assert archived["archived_at"] > 0
    assert "runner_error" not in report
    assert "http_rejection" not in report
    assert "blocked" not in report
    assert path.read_bytes() == bytes_before

    # A later restart must retain the first failure, not overwrite its evidence.
    report["runner_error"] = "second attempt interrupted"
    report["provider_calls"] = 13
    path.write_text(json.dumps(report), encoding="utf-8")
    resumed_again = load_report(tmp_path, resume=True)

    assert resumed_again["cases"] == cases
    assert len(resumed_again["attempts"]) == 2
    assert resumed_again["attempts"][0] == archived
    assert resumed_again["attempts"][1]["runner_error"] == "second attempt interrupted"
    assert resumed_again["attempts"][1]["provider_calls"] == 13
    assert "http_rejection" not in resumed_again["attempts"][1]


def test_load_report_without_resume_rejects_existing_report_without_overwriting(
    tmp_path: Path,
) -> None:
    path = tmp_path / "report.json"
    path.write_text(json.dumps({"model": MODEL, "cases": []}), encoding="utf-8")
    bytes_before = path.read_bytes()

    with pytest.raises(ValueError, match="^run already has a report; use --resume$"):
        load_report(tmp_path, resume=False)

    assert path.read_bytes() == bytes_before


def test_restored_sessions_recovers_groups_from_legacy_cases() -> None:
    report = {
        "cases": [
            {"id": "html-create", "session_key": "legacy-html"},
            {"id": "html-edit", "session_key": "legacy-html"},
            {"id": "html-multipage", "session_key": "legacy-multi"},
            {"id": "pdf-create", "session_key": "legacy-pdf"},
            {"id": "pptx-create", "session_key": "legacy-pptx"},
            {"id": "markdown-create", "session_key": "legacy-md"},
            {"id": "csv-create", "session_key": "legacy-csv"},
            {"id": "mixed-delivery", "session_key": "legacy-mixed"},
        ]
    }

    assert restored_sessions(report) == {
        "html": "legacy-html",
        "multi": "legacy-multi",
        "pdf": "legacy-pdf",
        "pptx": "legacy-pptx",
        "md": "legacy-md",
        "csv": "legacy-csv",
        "mixed": "legacy-mixed",
    }
    assert "sessions" not in report


def test_restored_sessions_preserves_saved_groups_and_pending_session() -> None:
    report = {
        "sessions": {"html": "old-html", "pdf": "saved-pdf"},
        "cases": [{"id": "html-create", "session_key": "completed-html"}],
        "pending_case": {
            "id": "html-edit",
            "session_key": "pending-html",
            "client_request_id": "original-request-id",
            "before": {"session": []},
        },
    }

    assert restored_sessions(report) == {"html": "pending-html", "pdf": "saved-pdf"}
    assert report["sessions"] == {"html": "old-html", "pdf": "saved-pdf"}
    assert report["pending_case"]["client_request_id"] == "original-request-id"


def test_load_report_resume_keeps_pending_case_for_session_recovery(tmp_path: Path) -> None:
    pending = {
        "id": "pdf-create",
        "session_key": "pending-pdf-session",
        "client_request_id": "existing-provider-request",
        "before": {"session": [], "files": []},
        "accepted": {"runId": "accepted-run"},
    }
    original = {"model": MODEL, "cases": [], "pending_case": pending}
    (tmp_path / "report.json").write_text(json.dumps(original), encoding="utf-8")

    report = load_report(tmp_path, resume=True)

    assert report["pending_case"] == pending
    assert restored_sessions(report) == {"pdf": "pending-pdf-session"}


@pytest.fixture
def forbid_live_validation(monkeypatch: pytest.MonkeyPatch) -> None:
    def forbidden(*args, **kwargs):
        pytest.fail("offline content validation attempted live execution or credential access")

    monkeypatch.setattr(acceptance.httpx, "Client", forbidden)
    monkeypatch.setattr(acceptance.httpx, "AsyncClient", forbidden)
    monkeypatch.setattr(acceptance.getpass, "getpass", forbidden)
    monkeypatch.setattr(acceptance.subprocess, "Popen", forbidden)


def _content_case(case_id: str) -> dict:
    return {
        "id": case_id,
        "session_key": f"agent:main:webchat:{case_id}",
        "task": {"status": "succeeded"},
        "before": {"files": [], "sources": [], "artifacts": [], "publications": []},
        "after": {
            "workspace": None,
            "files": [],
            "sources": [],
            "artifacts": [],
            "publications": [],
        },
        "downloads": [],
        "failures": [],
    }


def _download_case(
    run_root: Path,
    case_id: str,
    name: str,
    data: bytes,
    *,
    parsed: dict | None = None,
) -> dict:
    path = run_root / "downloads" / case_id / name
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(data)
    case = _content_case(case_id)
    case["downloads"] = [{
        "name": name,
        "path": str(path),
        "sha256": hashlib.sha256(data).hexdigest(),
        "bytes": len(data),
        "parsed": parsed or {},
    }]
    return case


def _workspace_case(run_root: Path, case_id: str, files: dict[str, str]) -> dict:
    workspace = run_root / "workspace" / case_id
    case = _content_case(case_id)
    case["after"]["workspace"] = {"root": str(workspace)}
    for name, text in files.items():
        path = workspace / name
        path.parent.mkdir(parents=True, exist_ok=True)
        data = text.encode("utf-8")
        path.write_bytes(data)
        case["after"]["files"].append({
            "path": name, "size": len(data), "sha256": hashlib.sha256(data).hexdigest()
        })
    return case


def _check(result: dict, name: str) -> dict:
    matches = [check for check in result["checks"] if check["name"] == name]
    assert len(matches) == 1, f"Expected exactly one {name!r} check, got {result['checks']!r}"
    return matches[0]


def _chinese_pdf_bytes(*pages: str) -> bytes:
    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    buffer = io.BytesIO()
    document = canvas.Canvas(buffer)
    for text in pages:
        document.setFont("STSong-Light", 14)
        document.drawString(40, 720, text)
        document.showPage()
    document.save()
    return buffer.getvalue()


@pytest.mark.parametrize(("case_id", "release"), [("pdf-create", "上线"), ("pdf-edit", "灰度发布")])
def test_report_validation_rechecks_pdf_bytes_instead_of_cached_page_count(
    tmp_path: Path, forbid_live_validation, case_id, release
) -> None:
    content = f"海风计划 本周完成 接口联调 首页开发 下周计划 测试 {release}"
    data = _chinese_pdf_bytes(content)
    case = _download_case(
        tmp_path, case_id, "weekly-report.pdf", data, parsed={"pages": 2, "text": content}
    )
    report_path = tmp_path / "report.json"
    report_path.write_text(
        json.dumps({"model": MODEL, "cases": [case]}, ensure_ascii=False, indent=4) + "\n",
        encoding="utf-8",
    )
    before = report_path.read_bytes()
    mtime_before = report_path.stat().st_mtime_ns

    result = acceptance.validate_report(report_path)

    assert result["source_report"] == str(report_path)
    assert result["source_sha256"] == hashlib.sha256(before).hexdigest()
    assert result["status"] == "failed"
    assert result["summary"] == {"passed": 0, "failed": 1, "inconclusive": 0}
    assert len(result["cases"]) == 1
    validated_case = result["cases"][0]
    assert validated_case["id"] == case_id
    assert validated_case["status"] == "failed"
    page_count = _check(validated_case, "pdf_page_count")
    assert page_count["status"] == "failed"
    assert page_count["expected"] == 2
    assert page_count["actual"] == 1
    assert report_path.read_bytes() == before
    assert report_path.stat().st_mtime_ns == mtime_before
    assert Path(case["downloads"][0]["path"]).read_bytes() == data


@pytest.mark.parametrize(("case_id", "release"), [("pdf-create", "上线"), ("pdf-edit", "灰度发布")])
def test_pdf_content_accepts_two_pages_with_the_requested_plan(
    tmp_path: Path, forbid_live_validation, case_id, release
) -> None:
    data = _chinese_pdf_bytes(
        "海风计划 本周完成 接口联调 首页开发", f"海风计划 下周计划 测试 {release}"
    )
    case = _download_case(tmp_path, case_id, "weekly-report.pdf", data)

    result = acceptance.validate_case_content(case, tmp_path)

    assert result["status"] == "passed"
    assert _check(result, "pdf_page_count")["status"] == "passed"


def test_pdf_edit_rejects_unchanged_release_plan(
    tmp_path: Path, forbid_live_validation
) -> None:
    data = _chinese_pdf_bytes(
        "海风计划 本周完成 接口联调 首页开发", "海风计划 下周计划 测试 上线"
    )
    case = _download_case(tmp_path, "pdf-edit", "weekly-report.pdf", data)

    result = acceptance.validate_case_content(case, tmp_path)

    assert _check(result, "pdf_page_count")["status"] == "passed"
    assert result["status"] == "failed"


_EXPECTED_CSV = (
    "date,category,amount\n"
    "2026-09-01,coffee,28\n"
    "2026-09-02,lunch,35\n"
    "2026-09-03,transport,12\n"
)


@pytest.mark.parametrize(
    ("content", "expected_status"),
    [(_EXPECTED_CSV, "passed"), (_EXPECTED_CSV.replace("coffee,28", "coffee,29"), "failed")],
    ids=["exact-requested-rows", "wrong-amount"],
)
def test_csv_content_checks_actual_rows(
    tmp_path: Path, forbid_live_validation, content, expected_status
) -> None:
    case = _download_case(
        tmp_path, "csv-create", "expenses.csv", content.encode(), parsed={"text": _EXPECTED_CSV}
    )

    result = acceptance.validate_case_content(case, tmp_path)

    assert result["status"] == expected_status
    assert _check(result, "csv_rows")["status"] == expected_status


_EXPECTED_MARKDOWN = (
    "# 新同事入职清单\n\n"
    "- [ ] 开通账号\n"
    "- [ ] 配置电脑\n"
    "- [ ] 加入团队\n\n"
    "| 账号 | 负责人 |\n"
    "| --- | --- |\n"
    "| 邮箱 | 王同事 |\n"
)


@pytest.mark.parametrize(
    ("content", "checkbox_status", "table_status"),
    [
        (_EXPECTED_MARKDOWN, "passed", "passed"),
        (_EXPECTED_MARKDOWN.replace("- [ ] 加入团队\n", ""), "failed", "passed"),
        (_EXPECTED_MARKDOWN.split("| 账号")[0], "passed", "failed"),
        (
            _EXPECTED_MARKDOWN.replace("| 负责人 |", "| 负责人 | 状态 |")
            .replace("| --- | --- |", "| --- | --- | --- |")
            .replace("| 王同事 |", "| 王同事 | 待办理 |"),
            "passed",
            "failed",
        ),
    ],
    ids=["complete-checklist", "only-two-checkboxes", "missing-table", "three-column-table"],
)
def test_markdown_content_requires_three_checkboxes_and_two_column_table(
    tmp_path: Path, forbid_live_validation, content, checkbox_status, table_status
) -> None:
    case = _download_case(
        tmp_path,
        "markdown-create",
        "onboarding.md",
        content.encode("utf-8-sig"),
        parsed={"text": _EXPECTED_MARKDOWN},
    )

    result = acceptance.validate_case_content(case, tmp_path)

    assert _check(result, "markdown_checkboxes")["status"] == checkbox_status
    assert _check(result, "markdown_table")["status"] == table_status
    expected = "passed" if checkbox_status == table_status == "passed" else "failed"
    assert result["status"] == expected


def _cafe_files() -> dict[str, str]:
    pages = [
        ("welcome.html", "首页", "欢迎来到咖啡店"),
        ("drinks.html", "菜单", "拿铁 28 元，美式 22 元"),
        ("reach-us.html", "联系方式", "地址：城市中心 电话：010-12345678"),
    ]
    navigation = "".join(f'<a href="{name}">{title}</a>' for name, title, _ in pages)
    return {
        **{
            f"cafe/{name}": (
                f"<!doctype html><html><head><title>{title}</title>"
                '<link rel="stylesheet" href="assets/common.css"></head>'
                f"<body><nav>{navigation}</nav><h1>{title}</h1><p>{content}</p></body></html>"
            )
            for name, title, content in pages
        },
        "cafe/assets/common.css": "body { color: #321; background: #fff; }\n",
    }


def test_html_multipage_accepts_arbitrary_filenames_with_links_and_shared_css(
    tmp_path: Path, forbid_live_validation
) -> None:
    case = _workspace_case(tmp_path, "html-multipage", _cafe_files())

    result = acceptance.validate_case_content(case, tmp_path)

    assert result["status"] == "passed"
    assert _check(result, "html_page_links")["status"] == "passed"
    assert _check(result, "html_shared_css")["status"] == "passed"


def test_html_multipage_accepts_branded_homepage_heading(
    tmp_path: Path, forbid_live_validation
) -> None:
    files = _cafe_files()
    files["cafe/welcome.html"] = files["cafe/welcome.html"].replace(
        "<title>首页</title>", "<title>海风咖啡</title>"
    ).replace("<h1>首页</h1>", "<h1>海风咖啡</h1>")
    case = _workspace_case(tmp_path, "html-multipage", files)

    result = acceptance.validate_case_content(case, tmp_path)

    assert result["status"] == "passed"
    assert _check(result, "html_page_links")["status"] == "passed"


def test_html_multipage_rejects_missing_cross_page_link(
    tmp_path: Path, forbid_live_validation
) -> None:
    files = _cafe_files()
    files["cafe/welcome.html"] = files["cafe/welcome.html"].replace(
        '<a href="reach-us.html">联系方式</a>', ""
    )
    case = _workspace_case(tmp_path, "html-multipage", files)

    result = acceptance.validate_case_content(case, tmp_path)

    assert result["status"] == "failed"
    assert _check(result, "html_page_links")["status"] == "failed"
    assert _check(result, "html_shared_css")["status"] == "passed"


def test_html_multipage_rejects_missing_css_dependency(
    tmp_path: Path, forbid_live_validation
) -> None:
    files = _cafe_files()
    files.pop("cafe/assets/common.css")
    case = _workspace_case(tmp_path, "html-multipage", files)

    result = acceptance.validate_case_content(case, tmp_path)

    assert result["status"] == "failed"
    assert _check(result, "html_shared_css")["status"] == "failed"


def test_changed_historical_html_snapshot_is_inconclusive_instead_of_wrong_title(
    tmp_path: Path, forbid_live_validation
) -> None:
    original = (
        "<!doctype html><html><head><title>北京四季</title></head>"
        "<body><h1>北京四季</h1><p>春季 夏季 秋季 冬季</p></body></html>"
    )
    case = _workspace_case(tmp_path, "html-create", {"beijing-site/index.html": original})
    current = original.replace("北京四季", "北京四季出行备忘")
    source_path = Path(case["after"]["workspace"]["root"]) / "beijing-site/index.html"
    source_path.write_text(current, encoding="utf-8")

    result = acceptance.validate_case_content(case, tmp_path)

    assert result["status"] == "inconclusive"
    assert _check(result, "source_snapshot_available")["status"] == "inconclusive"
    assert not [check for check in result["checks"] if check["status"] == "failed"]
    assert source_path.read_text(encoding="utf-8") == current


def _plan_docx(fixture, *, include_year=True, disclaimer_only=False):
    from docx import Document

    document = Document()
    document.add_heading(fixture["title"], 0)
    document.add_paragraph("虚构教学数据；" + "；".join(
        (fixture["year"] + "年" if include_year else "")
        + f"{row[0]}{row[3]}为{row[1]}{row[2]}" for row in fixture["rows"]
    ) + "。")
    table = document.add_table(rows=1, cols=5)
    for cell, text in zip(table.rows[0].cells, ["城市", "人口", "单位", "统计口径", "数据年份"]):
        cell.text = text
    for row in fixture["rows"]:
        values = row + [fixture["year"] if include_year else ""]
        for cell, text in zip(table.add_row().cells, values):
            cell.text = text
    if disclaimer_only:
        document.add_paragraph("部分数据可能为 " + fixture["year"] + " 年，以最新公报为准。")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_plan_suite_repeats_core_cases_without_reusing_session_groups() -> None:
    cases = acceptance.plan_cases()
    assert [case["id"] for case in cases if case["kind"] == "plan-first"] == [
        "plan-first-1", "plan-first-2", "plan-first-3",
    ]
    assert [case["id"] for case in cases if case["kind"] == "plan-repair"] == [
        "plan-repair-1", "plan-repair-2", "plan-repair-3",
    ]
    assert len({case["id"] for case in cases}) == len(cases)


@pytest.mark.parametrize("count", [0, -1, 11, True, 1.5])
def test_plan_suite_rejects_unbounded_repetition(count) -> None:
    with pytest.raises(ValueError, match="plan repetitions"):
        acceptance.plan_cases(count)


def test_docx_structure_pass_does_not_satisfy_missing_fixture_year(tmp_path: Path) -> None:
    case = {"kind": "plan-repair", "number": 2}
    fixture = acceptance.plan_fixture(case)
    draft = acceptance.seed_plan_draft(tmp_path, fixture, tmp_path / "evidence")
    result = acceptance.check_plan_document(Path(draft["path"]).read_bytes(), fixture)
    assert result["structure_passed"] is True
    assert result["fixture_requirements_passed"] is False
    assert result["rows_missing_required_fact"] == [row[0] for row in fixture["rows"]]
    assert result["body_missing_required_fact"] == [row[0] for row in fixture["rows"]]
    saved = (tmp_path / "evidence" / "before.docx").read_bytes()
    assert hashlib.sha256(saved).hexdigest() == draft["sha256"]


def test_docx_generic_disclaimer_cannot_replace_per_row_year() -> None:
    fixture = acceptance.plan_fixture({"kind": "plan-repair", "number": 1})
    result = acceptance.check_plan_document(
        _plan_docx(fixture, include_year=False, disclaimer_only=True), fixture,
    )
    assert result["structure_passed"] is True
    assert result["fixture_requirements_passed"] is False


@pytest.mark.parametrize("number", [1, 2, 3])
def test_docx_content_check_uses_current_fixture_facts_not_a_fixed_year(number) -> None:
    fixture = acceptance.plan_fixture({"kind": "plan-repair", "number": number})
    result = acceptance.check_plan_document(_plan_docx(fixture), fixture)
    assert result["fixture_requirements_passed"]
    wrong_fixture = {**fixture, "year": "1999"}
    wrong_result = acceptance.check_plan_document(_plan_docx(fixture), wrong_fixture)
    assert not wrong_result["fixture_requirements_passed"]


def test_unverifiable_case_does_not_supply_an_answer_in_prompt_or_source(tmp_path: Path) -> None:
    case = {"kind": "plan-unverifiable", "number": 1}
    fixture = acceptance.plan_fixture(case)
    assert fixture["year"] not in acceptance.plan_prompt(case, fixture)
    acceptance.seed_plan_draft(tmp_path, fixture, tmp_path / "evidence")
    assert "year" not in json.loads((tmp_path / "source.json").read_text(encoding="utf-8"))


@pytest.mark.parametrize("kind", ["plan-repair", "plan-flexible", "plan-unverifiable"])
def test_draft_cases_disclose_input_handoff_before_planning(kind: str) -> None:
    case = {"kind": kind, "number": 1}
    prompt = acceptance.plan_prompt(case, acceptance.plan_fixture(case))
    assert "当前工作区已提供 report.docx 草稿及 source.json" in prompt
    assert "先检查草稿和资料，再规划如何依据要求修改" in prompt
    assert "缺少年份" not in prompt  # The Agent must discover the fixture defect.


def test_plan_fixture_does_not_overwrite_an_existing_output(tmp_path: Path) -> None:
    existing = tmp_path / "report.docx"
    existing.write_bytes(b"keep original work")
    with pytest.raises(ValueError, match="do not overwrite"):
        acceptance.seed_plan_draft(
            tmp_path, acceptance.plan_fixture({"kind": "plan-repair", "number": 1}),
            tmp_path / "evidence",
        )
    assert existing.read_bytes() == b"keep original work"


def test_plan_fixture_does_not_overwrite_existing_source_or_partially_create_draft(
    tmp_path: Path,
) -> None:
    source = tmp_path / "source.json"
    source.write_text('{"keep": "original source"}', encoding="utf-8")
    with pytest.raises(ValueError, match="do not overwrite"):
        acceptance.seed_plan_draft(
            tmp_path, acceptance.plan_fixture({"kind": "plan-repair", "number": 1}),
            tmp_path / "evidence",
        )
    assert json.loads(source.read_text(encoding="utf-8")) == {"keep": "original source"}
    assert not (tmp_path / "report.docx").exists()


def test_plan_public_input_preflight_requires_both_readable_files(tmp_path: Path) -> None:
    def http(request):
        assert request.headers["x-opensquilla-session-key"] == "owned-session"
        if request.method == "POST":
            assert set(json.loads(request.content)["paths"]) == {"report.docx", "source.json"}
            return httpx.Response(200, json={"files": [{
                "path": "report.docx",
                "contentUrl": "/api/v1/workspace-files/content?path=report.docx",
            }]})
        return httpx.Response(200, content=b"document bytes")

    with httpx.Client(base_url="http://127.0.0.1:1", transport=httpx.MockTransport(http)) as client:
        with pytest.raises(RuntimeError, match="unavailable through the public file API"):
            acceptance.read_plan_inputs(client, "owned-session", tmp_path, tmp_path, "input")
    assert (tmp_path / "input-report.docx").read_bytes() == b"document bytes"


def test_plan_offline_validation_rechecks_download_not_success_or_cached_structure(
    tmp_path: Path, forbid_live_validation,
) -> None:
    fixture = acceptance.plan_fixture({"kind": "plan-repair", "number": 1})
    workspace = tmp_path / "workspace.docx"
    download = tmp_path / "download.docx"
    workspace.write_bytes(_plan_docx(fixture))
    download.write_bytes(_plan_docx(fixture, include_year=False))
    case = {"id": "plan-repair-1", "kind": "plan-repair", "fixture": fixture,
            "task": {"status": "succeeded"}, "failures": [],
            "outputs": [{"kind": kind, "evidence_path": path.name,
                         "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
                         "cached_pass": True}
                        for kind, path in [("workspace", workspace), ("download", download)]]}
    report = tmp_path / "report.json"
    report.write_text(json.dumps({"suite": "plan", "cases": [case]}), encoding="utf-8")
    before = report.read_bytes()
    result = acceptance.validate_report(report)
    assert result["status"] == "failed"
    assert all(check["structure_passed"] for check in result["cases"][0]["checks"])
    assert "fixture_content_not_satisfied" in result["cases"][0]["transport_failures"]
    assert report.read_bytes() == before


def test_plan_correct_fixture_still_requires_review_of_actual_completion_claim(
    tmp_path: Path,
) -> None:
    fixture = acceptance.plan_fixture({"kind": "plan-repair", "number": 1})
    path = tmp_path / "correct.docx"
    path.write_bytes(_plan_docx(fixture))
    case = {"id": "plan-repair-1", "kind": "plan-repair", "fixture": fixture,
            "failures": [], "outputs": [{"kind": kind, "evidence_path": path.name,
                                          "sha256": hashlib.sha256(path.read_bytes()).hexdigest()}
                                         for kind in ("workspace", "download")]}
    result = acceptance.validate_plan_case(case, tmp_path)
    assert result["status"] == "inconclusive"
    assert result["review_required"]
    assert acceptance.plan_report_exit_code({"cases": [{**case, "assessment": result}]}) == 2


def test_custom_model_relay_preserves_exact_model_binding(tmp_path: Path) -> None:
    request_log = FunctionalRequestLog(tmp_path / "model-requests.sqlite3", enabled=True)
    request_log.select_phase(variant="new", case_id="custom-model")
    with contextlib.ExitStack() as stack:
        relay = BoundedRelay(None, {}, model="deepseek-flash", api_key="offline-only",
                             request_log=request_log,
                             transport=httpx.MockTransport(
                                 lambda _request: httpx.Response(200, content=_REPLY),
                             ))
        stack.callback(relay.close)
        _forward(relay, _body(model="deepseek-flash", max_tokens=1))
        with pytest.raises(BudgetRejectedError, match="acceptance_request_limit"):
            _forward(relay, _body(model=MODEL, max_tokens=1))
    assert [item["model"] for item in request_log.snapshot()["requests"]] == ["deepseek-flash"]


def test_resume_rejects_a_different_model(tmp_path: Path) -> None:
    (tmp_path / "report.json").write_text(json.dumps({"model": "deepseek-flash", "cases": []}))
    with pytest.raises(ValueError, match="model or cases"):
        acceptance.load_report(tmp_path, resume=True)
    resumed = acceptance.load_report(tmp_path, resume=True, model="deepseek-flash")
    assert resumed["model"] == "deepseek-flash"


@pytest.mark.parametrize("repair", [False, True])
@pytest.mark.parametrize("kind", ["plan-repair", "plan-flexible", "plan-unverifiable"])
def test_plan_runner_approves_via_public_rpc_and_retains_failed_content(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, repair: bool, kind: str,
) -> None:
    from opensquilla.gateway.adapters.turn_admission import GatewayTurnAdmissionAdapter

    case_id = kind + "-1"
    workspace = tmp_path / "profile" / "tasks" / "fixture"
    workspace.mkdir(parents=True)
    fixture = acceptance.plan_fixture({"kind": kind, "number": 1})
    calls = []
    implemented = False
    planning = False
    clarified = False

    def rpc(_client, method, params):
        nonlocal implemented, planning, clarified
        calls.append((method, params))
        if method == "sessions.create":
            assert params == {"agentId": "main", "kind": "webchat"}
            assert not (workspace / "report.docx").exists()
            return {"key": "agent:main:webchat:created", "sessionId": "created-session"}
        if method == "plans.setMode":
            assert params == {"sessionKey": "agent:main:webchat:created", "mode": "plan",
                              "expectedRevision": 0}
            return {"collaboration": {"mode": "plan", "revision": 1}}
        if method == "chat.send":
            assert params["sessionKey"] == "agent:main:webchat:created"
            assert params["intent"] == "continue"
            assert "collaborationMode" not in params
            assert GatewayTurnAdmissionAdapter._initial_collaboration_mode(params) is None
            assert (workspace / "report.docx").is_file()
            source = json.loads((workspace / "source.json").read_text(encoding="utf-8"))
            assert ("year" in source) == fixture["verifiable"]
            planning = True
            return {"taskId": "planning-turn"}
        if method == "chat.clarify_submit":
            assert params["requestId"] == "input-location"
            assert "当前工作区" in params["fields"]["source"]
            if kind == "plan-unverifiable":
                assert "没有额外来源或可核实年份" in params["fields"]["source"]
                assert fixture["year"] not in params["fields"]["source"]
            clarified = True
            return {"resolved": True}
        if method == "plans.implement":
            assert params["planRevisionId"] == "real-rpc-proposal"
            assert (workspace / "report.docx").is_file()
            assert clarified
            if repair:
                (workspace / "report.docx").write_bytes(_plan_docx(fixture))
            implemented = True
            return {"turn_id": "implementation-turn"}
        assert method == "sessions.bootstrap"
        if not planning:
            return {"session": {"workspace": str(workspace)},
                    "collaboration": {"mode": "default", "revision": 0}, "tasks": []}
        if not clarified:
            return {"session": {"workspace": str(workspace), "pendingUserInputs": [{
                "request_id": "input-location", "clarify_schema": {"fields": [{
                    "name": "source", "type": "enum", "allow_other": True,
                    "choices": ["另行上传", "使用工作区资料"],
                }]},
            }]}, "tasks": [{"task_id": "planning-turn", "status": "running"}]}
        return {"session": {"workspace": str(workspace)},
                "currentPlan": {"revisionId": "real-rpc-proposal"},
                "tasks": [{"task_id": "implementation-turn" if implemented else "planning-turn",
                           "status": "succeeded"}], "history": {"messages": ["full evidence"]}}

    def http(request):
        if request.method == "POST":
            assert request.url.path == "/api/v1/workspace-files/resolve"
            return httpx.Response(200, json={"files": [{"path": name,
                "contentUrl": "/api/v1/workspace-files/content?path=" + name,
            } for name in json.loads(request.content)["paths"]]})
        name = request.url.params.get("path", "report.docx")
        return httpx.Response(200, content=(workspace / name).read_bytes())

    monkeypatch.setattr(acceptance, "plan_rpc", rpc)
    monkeypatch.setattr(acceptance.time, "sleep", lambda _duration: None)
    monkeypatch.setattr(acceptance, "snapshot", lambda *_: {"artifacts": [{
        "id": "delivered-report", "name": "report.docx",
        "sha256": hashlib.sha256((workspace / "report.docx").read_bytes()).hexdigest(),
    }]})
    log = FunctionalRequestLog(tmp_path / "calls.sqlite3", enabled=True)
    report = {"suite": "plan", "model": MODEL, "repetitions": 3, "cases": []}
    with httpx.Client(base_url="http://127.0.0.1:1", transport=httpx.MockTransport(http)) as client:
        code = acceptance.run_plan_cases(client, tmp_path, report, {case_id}, log)
        before_calls = list(calls)
        assert acceptance.run_plan_cases(client, tmp_path, report, {case_id}, log) == code
    assert calls == before_calls  # Resume does not silently rerun a failed attempt.
    assert [method for method, _ in calls] == [
        "sessions.create", "sessions.bootstrap", "plans.setMode", "chat.send",
        "sessions.bootstrap", "chat.clarify_submit", "sessions.bootstrap",
        "plans.implement", "sessions.bootstrap",
    ]
    entry = report["cases"][0]
    assert entry["inputs_before_planning"]["files"]["report.docx"]["sha256"] == (
        entry["seeded_draft"]["sha256"]
    )
    before_hashes = {
        name: item["sha256"] for name, item in entry["inputs_before_planning"]["files"].items()
    }
    assert before_hashes == {
        name: item["sha256"] for name, item in entry["inputs_after_planning"]["files"].items()
    }
    assert entry["seeded_draft"]["parsed"]["tables"]
    assert not acceptance.check_plan_document(
        (tmp_path / "plan-evidence" / case_id / "before.docx").read_bytes(), fixture,
    )["fixture_requirements_passed"]
    assert len(entry["outputs"]) == 2
    assert entry["stages"][1]["bootstrap"]["history"]["messages"] == ["full evidence"]
    needs_review = repair or kind == "plan-unverifiable"
    assert entry["assessment"]["status"] == ("inconclusive" if needs_review else "failed")
    assert code == (2 if needs_review else 1)


def test_plan_first_original_prompt_answers_structured_questions_and_stops_at_submit(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    sessions = {}
    calls = []

    def rpc(_client, method, params):
        calls.append((method, params))
        if method == "chat.send":
            assert params["message"] == "帮我写个文档介绍华北"
            key = params["sessionKey"]
            assert key not in sessions
            sessions[key] = False
            return {"taskId": key + "-turn"}
        if method == "chat.clarify_submit":
            sessions[params["sessionKey"]] = True
            assert "2500" in params["fields"]["preferences"]
            assert "内蒙古" in params["fields"]["preferences"]
            return {"resolved": True}
        assert method == "sessions.bootstrap"  # No plans.implement or artifact request.
        key = params["key"]
        answered = sessions[key]
        pending = [] if answered else [{"request_id": key + "-question",
                                       "clarify_schema": {"fields": [
                                           {"name": "preferences", "type": "string"},
                                       ]}}]
        return {"session": {"pendingUserInputs": pending},
                "tasks": [{"task_id": key + "-turn",
                           "status": "succeeded" if answered else "running"}],
                "currentPlan": {"revisionId": key + "-proposal"} if answered else None}

    monkeypatch.setattr(acceptance, "plan_rpc", rpc)
    monkeypatch.setattr(acceptance.time, "sleep", lambda _duration: None)
    log = FunctionalRequestLog(tmp_path / "calls.sqlite3", enabled=True)
    report = {"suite": "plan", "model": MODEL, "repetitions": 3, "cases": []}
    selected = {"plan-first-1", "plan-first-2", "plan-first-3"}
    code = acceptance.run_plan_cases(object(), tmp_path, report, selected, log)
    assert code == 0
    assert len(sessions) == 3
    assert all(case["assessment"]["status"] == "passed" for case in report["cases"])
    assert all(len(case["stages"]) == 1 for case in report["cases"])
    assert all(len(case["stages"][0]["clarifications"]) == 1 for case in report["cases"])
    assert "plans.implement" not in [method for method, _ in calls]


def test_correct_table_does_not_hide_a_body_missing_required_year() -> None:
    from docx import Document

    fixture = acceptance.plan_fixture({"kind": "plan-repair", "number": 1})
    document = Document(io.BytesIO(_plan_docx(fixture)))
    document.paragraphs[1].text = document.paragraphs[1].text.replace(fixture["year"] + "年", "")
    buffer = io.BytesIO()
    document.save(buffer)
    result = acceptance.check_plan_document(buffer.getvalue(), fixture)
    assert result["rows_missing_required_fact"] == []
    assert result["body_missing_required_fact"]
    assert not result["fixture_requirements_passed"]


def test_plan_first_without_proposal_is_not_accepted_as_success(tmp_path: Path) -> None:
    result = acceptance.validate_plan_case({
        "id": "plan-first-1", "kind": "plan-first", "failures": [],
        "stages": [{"bootstrap": {"currentPlan": None, "tasks": [{"status": "succeeded"}]}}],
    }, tmp_path)
    assert result["status"] == "failed"
    assert result["transport_failures"] == ["first_turn_missing_proposal"]


def test_windows_child_home_supports_expanduser_without_real_profile_access(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    original_environment = dict(os.environ)
    env = acceptance.child_environment(
        "tokenrhythm", {"TOKENRHYTHM_API_KEY": "synthetic-placeholder"},
        base_environment={"USERPROFILE": r"C:\real-user", "APPDATA": r"C:\real-roaming",
                          "LOCALAPPDATA": r"C:\real-local"},
    )
    env.update(acceptance.isolated_windows_home(tmp_path, platform="nt"))
    assert os.environ == original_environment
    assert all(Path(env[name]).is_dir() and Path(env[name]).is_relative_to(tmp_path)
               for name in ("USERPROFILE", "APPDATA", "LOCALAPPDATA"))
    with monkeypatch.context() as patch:
        patch.delenv("HOME", raising=False)
        patch.delenv("HOMEDRIVE", raising=False)
        patch.delenv("HOMEPATH", raising=False)
        for name, value in env.items():
            patch.setenv(name, value)
        assert ntpath.expanduser("~") == str(tmp_path / "host-home")
        if os.name == "nt":
            assert Path.home() == tmp_path / "host-home"


def test_windows_home_override_does_not_change_other_platforms(tmp_path: Path) -> None:
    assert acceptance.isolated_windows_home(tmp_path, platform="posix") == {}
    assert not (tmp_path / "host-home").exists()


@pytest.mark.parametrize("old_corrupt", [False, True])
def test_republished_correct_docx_preserves_old_failure_without_rejecting_final_repair(
    tmp_path: Path, old_corrupt: bool,
) -> None:
    fixture = acceptance.plan_fixture({"kind": "plan-repair", "number": 1})
    original = (b"invalid initial document" if old_corrupt
                else _plan_docx(fixture, include_year=False))
    corrected = _plan_docx(fixture)
    outputs = []
    for kind, name, data in [
        ("workspace", "workspace.docx", corrected),
        ("download", "first.docx", original),
        ("download", "last.docx", corrected),
    ]:
        (tmp_path / name).write_bytes(data)
        outputs.append({"kind": kind, "artifact_id": name, "evidence_path": name,
                        "sha256": hashlib.sha256(data).hexdigest()})
    case = {"id": "plan-repair-1", "kind": "plan-repair", "fixture": fixture,
            "failures": [], "outputs": outputs}
    result = acceptance.validate_plan_case(case, tmp_path)
    assert result["status"] == "inconclusive"  # Final-reply version still needs review.
    assert result["transport_failures"] == []
    assert len(result["checks"]) == 3
    first = result["checks"][1]
    assert first["fixture_requirements_passed"] is False
    assert first["final"] is False
    assert first["artifact_id"] == "first.docx"
    last = result["checks"][2]
    assert last["fixture_requirements_passed"] is True
    assert last["final"] is True
    assert last["sha256"] == result["checks"][0]["sha256"]


def test_correct_historical_docx_cannot_hide_an_incomplete_final_delivery(tmp_path: Path) -> None:
    fixture = acceptance.plan_fixture({"kind": "plan-repair", "number": 1})
    corrected = _plan_docx(fixture)
    incomplete = _plan_docx(fixture, include_year=False)
    outputs = []
    for kind, name, data in [
        ("workspace", "workspace.docx", corrected),
        ("download", "first.docx", corrected),
        ("download", "last.docx", incomplete),
    ]:
        (tmp_path / name).write_bytes(data)
        outputs.append({"kind": kind, "evidence_path": name,
                        "sha256": hashlib.sha256(data).hexdigest()})
    result = acceptance.validate_plan_case({
        "id": "plan-repair-1", "kind": "plan-repair", "fixture": fixture,
        "failures": [], "outputs": outputs,
    }, tmp_path)
    assert result["status"] == "failed"
    assert "fixture_content_not_satisfied" in result["transport_failures"]
    assert "final_download_workspace_mismatch" in result["transport_failures"]
    assert result["checks"][1]["final"] is False
    assert result["checks"][2]["final"] is True


def test_semantically_correct_but_different_final_bytes_require_explicit_version_match(
    tmp_path: Path,
) -> None:
    from docx import Document

    fixture = acceptance.plan_fixture({"kind": "plan-repair", "number": 1})
    workspace = _plan_docx(fixture)
    document = Document(io.BytesIO(workspace))
    document.add_paragraph("这是另一版本，不能暗中当成最终工作区版本。")
    buffer = io.BytesIO()
    document.save(buffer)
    download = buffer.getvalue()
    outputs = []
    for kind, data in [("workspace", workspace), ("download", download)]:
        path = tmp_path / (kind + ".docx")
        path.write_bytes(data)
        outputs.append({"kind": kind, "evidence_path": path.name,
                        "sha256": hashlib.sha256(data).hexdigest()})
    result = acceptance.validate_plan_case({
        "id": "plan-repair-1", "kind": "plan-repair", "fixture": fixture,
        "failures": [], "outputs": outputs,
    }, tmp_path)
    assert all(check["fixture_requirements_passed"] for check in result["checks"])
    assert result["status"] == "failed"
    assert result["transport_failures"] == ["final_download_workspace_mismatch"]


@pytest.mark.parametrize("plan_suite", [False, True])
def test_plan_budget_removes_internal_smoke_limits_but_preserves_external_bounds(
    plan_suite: bool,
) -> None:
    base = (
        'agent_max_iterations = 16\nagent_runtime_timeout_seconds = 240\n'
        'agent_max_provider_retries = 0\n'
        'llm_request_timeout_seconds = 90\n[task_runtime]\nturn_hard_deadline_s = 270\n'
    )
    budgets = acceptance.acceptance_budgets(plan_suite)
    rendered = acceptance.configure_acceptance_budgets(base, budgets)
    config = tomllib.loads(rendered)
    assert config["agent_max_iterations"] == (0 if plan_suite else 16)
    assert config.get("agent_runtime_timeout_seconds") == (None if plan_suite else 240)
    assert config.get("agent_max_provider_retries") == (None if plan_suite else 0)
    assert config["task_runtime"].get("turn_hard_deadline_s") == (None if plan_suite else 270)
    assert config["llm_request_timeout_seconds"] == 90
    assert budgets["harness_turn_deadline_s"] == (600 if plan_suite else 285)
    if not plan_suite:
        assert rendered == base


def test_plan_provider_retry_report_matches_unmodified_product_default() -> None:
    from opensquilla.engine.types import AgentConfig
    from opensquilla.gateway.config import GatewayConfig

    budgets = acceptance.acceptance_budgets(True)
    config = GatewayConfig.model_validate(tomllib.loads(
        acceptance.configure_acceptance_budgets('agent_max_provider_retries = 0\n', budgets),
    ))
    assert config.agent_max_provider_retries is None
    assert budgets["agent_max_provider_retries"] is None
    assert budgets["provider_retry_policy"] == "product_default"
    assert budgets["effective_agent_max_provider_retries"] == AgentConfig().max_provider_retries
    smoke_budgets = acceptance.acceptance_budgets(False)
    assert smoke_budgets["effective_agent_max_provider_retries"] == 0
    assert smoke_budgets["provider_retry_policy"] == "smoke_override"


def test_acceptance_deadline_aborts_owned_task_and_does_not_start_another_case(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    calls = []

    def rpc(_client, method, params):
        calls.append((method, params))
        if method == "chat.send":
            return {"task_id": "deadline-task"}
        assert method == "chat.abort"
        assert params["runId"] == "deadline-task"
        return {"aborted": True}

    ticks = iter([0, 601])
    monkeypatch.setattr(acceptance, "plan_rpc", rpc)
    monkeypatch.setattr(acceptance.time, "monotonic", lambda: next(ticks))
    log = FunctionalRequestLog(tmp_path / "calls.sqlite3", enabled=True)
    report = {"suite": "plan", "model": MODEL, "repetitions": 3, "cases": [],
              "budget_config": acceptance.acceptance_budgets(True)}
    code = acceptance.run_plan_cases(
        object(), tmp_path, report, {"plan-first-1", "plan-first-2"}, log,
    )
    assert code == 2
    assert [method for method, _params in calls] == ["chat.send", "chat.abort"]
    assert len(report["cases"]) == 1
    case = report["cases"][0]
    assert case["status"] == "acceptance_limit"
    assert case["budget_config"]["agent_max_iterations"] == 0
    assert case["budget_config"]["harness_turn_deadline_s"] == 600
    assert case["stages"][0]["acceptance_limit"] == {"kind": "wall_time", "seconds": 600}
    assert case["assessment"]["status"] == "inconclusive"
    assert case["assessment"]["reason"] == "acceptance_limit"


def test_resume_archives_budget_settings_without_rewriting_old_cases(tmp_path: Path) -> None:
    old_case = {"id": "plan-first-2", "budget_config": acceptance.acceptance_budgets(False),
                "failures": ["first_turn_missing_proposal"]}
    before = {"model": MODEL, "cases": [old_case],
              "budget_config": acceptance.acceptance_budgets(False)}
    (tmp_path / "report.json").write_text(json.dumps(before), encoding="utf-8")
    report = acceptance.load_report(tmp_path, resume=True)
    report["budget_config"] = acceptance.acceptance_budgets(True)
    assert report["attempts"][-1]["budget_config"]["agent_max_iterations"] == 16
    assert report["cases"] == [old_case]
    assert report["cases"][0]["budget_config"]["agent_max_iterations"] == 16


@pytest.mark.parametrize("wrong_number", ["3110", "1311", "311.5", "311,000", "-311"])
def test_plan_docx_rejects_population_substrings_in_body_and_table(wrong_number: str) -> None:
    from docx import Document

    fixture = acceptance.plan_fixture({"kind": "plan-repair", "number": 1})
    document = Document(io.BytesIO(_plan_docx(fixture)))
    document.paragraphs[1].text = document.paragraphs[1].text.replace("311", wrong_number)
    document.tables[0].rows[1].cells[1].text = wrong_number
    buffer = io.BytesIO()
    document.save(buffer)
    result = acceptance.check_plan_document(buffer.getvalue(), fixture)
    assert result["rows_missing_required_fact"] == ["海岬市"]
    assert result["body_missing_required_fact"] == ["海岬市"]


def test_plan_docx_accepts_shared_paragraph_year_and_same_table_header() -> None:
    from docx import Document

    fixture = acceptance.plan_fixture({"kind": "plan-repair", "number": 1})
    document = Document(io.BytesIO(_plan_docx(fixture)))
    document.paragraphs[1].text = (
        f"{fixture['year']}年年末常住人口：海岬市311.0万人，青湾市421万人。"
    )
    document.tables[0].rows[0].cells[1].text = f"{fixture['year']}年年末常住人口（万人）"
    for row in document.tables[0].rows[1:]:
        row.cells[2].text = ""
        row.cells[3].text = ""
        row.cells[4].text = ""
    buffer = io.BytesIO()
    document.save(buffer)
    result = acceptance.check_plan_document(buffer.getvalue(), fixture)
    assert result["fixture_requirements_passed"] is True


def test_plan_docx_cannot_borrow_year_from_body_or_a_different_table() -> None:
    from docx import Document

    fixture = acceptance.plan_fixture({"kind": "plan-repair", "number": 1})
    document = Document(io.BytesIO(_plan_docx(fixture)))
    for row in document.tables[0].rows[1:]:
        row.cells[4].text = ""
    other_table = document.add_table(rows=1, cols=1)
    other_table.cell(0, 0).text = f"其他资料，年份{fixture['year']}"
    buffer = io.BytesIO()
    document.save(buffer)
    result = acceptance.check_plan_document(buffer.getvalue(), fixture)
    assert result["body_missing_required_fact"] == []
    assert result["rows_missing_required_fact"] == ["海岬市", "青湾市"]


def test_plan_docx_cannot_borrow_year_from_another_city_row() -> None:
    from docx import Document

    fixture = acceptance.plan_fixture({"kind": "plan-repair", "number": 1})
    document = Document(io.BytesIO(_plan_docx(fixture)))
    document.tables[0].rows[2].cells[4].text = ""
    buffer = io.BytesIO()
    document.save(buffer)
    result = acceptance.check_plan_document(buffer.getvalue(), fixture)
    assert result["rows_missing_required_fact"] == ["青湾市"]


def test_plan_docx_rejects_year_substrings_instead_of_matching_another_number() -> None:
    from docx import Document

    fixture = acceptance.plan_fixture({"kind": "plan-repair", "number": 1})
    document = Document(io.BytesIO(_plan_docx(fixture)))
    document.paragraphs[1].text = document.paragraphs[1].text.replace("2021", "20210")
    for row in document.tables[0].rows[1:]:
        row.cells[4].text = "20210"
    buffer = io.BytesIO()
    document.save(buffer)
    result = acceptance.check_plan_document(buffer.getvalue(), fixture)
    assert not result["fixture_requirements_passed"]
    assert result["body_missing_required_fact"] == ["海岬市", "青湾市"]
    assert result["rows_missing_required_fact"] == ["海岬市", "青湾市"]


def test_gateway_bootstrap_installs_relay_before_importing_cli(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    from scripts import live_tokenrhythm_transport

    events = []
    launcher, pythonpath = acceptance.prepare_gateway_bootstrap(tmp_path)
    monkeypatch.setattr(live_tokenrhythm_transport, "install_from_env",
                        lambda: events.append("relay_installed"))
    monkeypatch.setattr(runpy, "run_module",
                        lambda name, **kwargs: events.append((name, kwargs)))
    runpy.run_path(str(launcher))
    assert events == ["relay_installed", ("opensquilla.cli.main", {"run_name": "__main__"})]
    assert str(tmp_path / "transport") not in pythonpath.split(os.pathsep)


def test_gateway_bootstrap_fails_closed_before_cli_when_relay_installation_fails(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    from scripts import live_tokenrhythm_transport

    def reject():
        raise RuntimeError("synthetic invalid relay")

    launcher, _pythonpath = acceptance.prepare_gateway_bootstrap(tmp_path)
    monkeypatch.setattr(live_tokenrhythm_transport, "install_from_env", reject)
    monkeypatch.setattr(runpy, "run_module", lambda *_args, **_kwargs: pytest.fail("CLI started"))
    with pytest.raises(SystemExit, match="acceptance transport unavailable"):
        runpy.run_path(str(launcher))


def test_tool_python_does_not_inherit_legacy_sitecustomize_or_transport_import(
    tmp_path: Path,
) -> None:
    legacy = tmp_path / "transport"
    legacy.mkdir()
    marker = legacy / "sitecustomize.py"
    marker.write_text("raise SystemExit('legacy injection must not run')\n", encoding="utf-8")
    _launcher, pythonpath = acceptance.prepare_gateway_bootstrap(tmp_path)
    env = acceptance.child_environment("tokenrhythm", {}, base_environment=os.environ)
    env.update(acceptance.isolated_windows_home(tmp_path))
    env["PYTHONPATH"] = pythonpath
    env["OPENSQUILLA_LIVE_TRANSPORT"] = "1"
    result = subprocess.run(
        [sys.executable, "-c",
         "import sys; print('scripts.live_tokenrhythm_transport' in sys.modules)"],
        cwd=tmp_path, env=env, capture_output=True, text=True, timeout=10, check=True,
    )
    assert result.stdout.strip() == "False"
    assert result.stderr == ""
    assert marker.is_file()  # Historical startup evidence is preserved.
