"""Generic authoring baseline, independent of channel execution authorization.

The trusted CLI generates files through the real generic tool dispatcher.
Publication and a fake channel exercise the shared artifact pipeline. Channel
attestation and OS sandbox enforcement have separate security regressions.
"""

from __future__ import annotations

import csv
import json
from io import BytesIO, StringIO
from pathlib import Path
from types import SimpleNamespace
from zipfile import ZipFile

import pytest
from openpyxl import load_workbook
from pptx import Presentation
from pypdf import PdfReader

from opensquilla.artifacts import ArtifactNotFoundError, ArtifactStore
from opensquilla.channels.artifact_delivery import deliver_artifacts_as_channel_files
from opensquilla.channels.contract import ChannelCapabilityProfile, ChannelSendResult
from opensquilla.channels.types import ChannelArtifactDeliveryRequest, IncomingMessage
from opensquilla.engine.types import ToolCall
from opensquilla.sandbox.config import SandboxSettings
from opensquilla.sandbox.integration import configure_runtime, reset_runtime
from opensquilla.tools.dispatch import build_tool_handler
from opensquilla.tools.registry import get_default_registry
from opensquilla.tools.types import CallerKind, ToolContext

_PROGRAMS = {
    "csv": '''import csv
with open("result.csv", "w", encoding="utf-8-sig", newline="") as stream:
    csv.writer(stream).writerows([["名称", "数量"], ["测试", 42]])
''',
    "xlsx": '''from openpyxl import Workbook
wb = Workbook()
ws = wb.active
ws.title = "Summary"
ws.append(["名称", "数量"])
ws.append(["测试", 42])
ws["B3"] = "=SUM(B2:B2)"
wb.save("result.xlsx")
''',
    "pptx": '''from pptx import Presentation
prs = Presentation()
slide = prs.slides.add_slide(prs.slide_layouts[1])
slide.shapes.title.text = "测试报告"
slide.placeholders[1].text = "Generic authoring works"
prs.save("result.pptx")
''',
    "pdf": '''from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.pdfgen import canvas
pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
document = canvas.Canvas("result.pdf")
document.setFont("STSong-Light", 16)
document.drawString(72, 720, "中文测试报告")
document.save()
''',
}


def _check_material(extension: str, material: bytes) -> None:
    if extension == "csv":
        assert material.startswith(b"\xef\xbb\xbf")
        assert list(csv.reader(StringIO(material.decode("utf-8-sig")))) == [
            ["名称", "数量"], ["测试", "42"],
        ]
    elif extension == "xlsx":
        with ZipFile(BytesIO(material)) as archive:
            assert archive.testzip() is None
            assert "xl/workbook.xml" in archive.namelist()
        workbook = load_workbook(BytesIO(material), data_only=False)
        assert workbook["Summary"]["A2"].value == "测试"
        assert workbook["Summary"]["B2"].value == 42
        assert workbook["Summary"]["B3"].value == "=SUM(B2:B2)"
        workbook.close()
    elif extension == "pptx":
        with ZipFile(BytesIO(material)) as archive:
            assert archive.testzip() is None
            assert "ppt/presentation.xml" in archive.namelist()
        deck = Presentation(BytesIO(material))
        assert deck.slides[0].shapes.title.text == "测试报告"
        assert "Generic authoring works" in deck.slides[0].placeholders[1].text
    else:
        document = PdfReader(BytesIO(material))
        assert len(document.pages) == 1
        assert "中文测试报告" in document.pages[0].extract_text()


@pytest.mark.parametrize("extension", list(_PROGRAMS))
async def test_generic_tools_generate_publish_and_deliver(
    tmp_path: Path, extension: str,
) -> None:
    workspace = tmp_path / "workspace"
    workspace.mkdir()
    media_root = tmp_path / "media"
    ctx = ToolContext(
        is_owner=True,
        caller_kind=CallerKind.CLI,
        workspace_dir=str(workspace),
        artifact_media_root=str(media_root),
        artifact_session_id="synthetic-authoring-session",
        session_key="agent:main:feishu:group:synthetic-room:sender:synthetic-user",
        run_mode="full",
    )
    registry = get_default_registry()
    names = {item.name for item in registry.to_tool_definitions()}
    assert {"create_csv", "create_xlsx", "create_pptx", "create_pdf_report"}.isdisjoint(names)
    handler = build_tool_handler(registry, ctx)
    configure_runtime(SandboxSettings(sandbox=False), workspace=workspace)
    try:
        written = await handler(ToolCall(
            tool_use_id="write-script", tool_name="write_file",
            arguments={"path": "author.py", "content": _PROGRAMS[extension]},
        ))
        assert not written.is_error, written.content
        executed = await handler(ToolCall(
            tool_use_id="run-script", tool_name="execute_code",
            arguments={"code": 'import runpy\nrunpy.run_path("author.py")'},
        ))
        assert not executed.is_error, executed.content
        execution = json.loads(executed.content)
        assert execution["exit_code"] == 0, execution
        # The channel publication contract hides host paths and delegates
        # transmission to the adapter rather than to the generation script.
        ctx.caller_kind = CallerKind.CHANNEL
        ctx.is_owner = False
        for attempt in range(2):
            published = await handler(ToolCall(
                tool_use_id=f"publish-{attempt}", tool_name="publish_artifact",
                arguments={"path": f"result.{extension}", "bundle": "none"},
            ))
            assert not published.is_error, published.content
            payload = json.loads(published.content)
            assert payload["status"] == ("published" if attempt == 0 else "already_published")
            assert "local_path" not in payload["artifact"]
        assert len(ctx.published_artifacts) == 1

        store = ArtifactStore(media_root)
        artifact = payload["artifact"]
        _, stored = store.resolve_for_download(
            artifact["id"], session_id=ctx.artifact_session_id,
        )
        material = stored.read_bytes()
        _check_material(extension, material)
        with pytest.raises(ArtifactNotFoundError):
            store.resolve_for_download(artifact["id"], session_id="other-session")

        outside = tmp_path / "outside.txt"
        outside.write_text("outside workspace", encoding="utf-8")
        denied = await handler(ToolCall(
            tool_use_id="outside-publish", tool_name="publish_artifact",
            arguments={"path": str(outside)},
        ))
        assert denied.is_error

        inbound = IncomingMessage(
            sender_id="synthetic-user", channel_id="synthetic-room",
            content="Generate a file", message_id="synthetic-inbound",
        )
        deliveries = []

        class FakeChannel:
            capability_profile = ChannelCapabilityProfile(
                channel_type="fake", artifact_delivery=True,
            )

            async def deliver_artifact(self, request: ChannelArtifactDeliveryRequest):
                assert request.inbound is inbound
                assert request.artifact_id == artifact["id"]
                assert Path(request.file_path).read_bytes() == material
                deliveries.append(request.name)
                return ChannelSendResult.sent(
                    capability="artifact_delivery", target_id=inbound.channel_id,
                )

        undelivered = await deliver_artifacts_as_channel_files(
            FakeChannel(), inbound, ctx.published_artifacts,
            SimpleNamespace(attachments=SimpleNamespace(media_root=str(media_root))),
            expected_session_id=ctx.artifact_session_id,
        )
        assert undelivered == []
        assert deliveries == [f"result.{extension}"]
    finally:
        reset_runtime()
