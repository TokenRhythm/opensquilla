"""Ordinary channel authoring must carry runtime authority, never route claims."""

from __future__ import annotations

import json
import socket
import stat
import subprocess
import sys
import zipfile
from pathlib import Path
from types import SimpleNamespace

import pytest

from opensquilla.channels.types import IncomingMessage
from opensquilla.execution_workspaces import prepare_managed_workspace
from opensquilla.gateway.project_workspace_runtime import apply_run_context_route_metadata
from opensquilla.gateway.routing import build_channel_route_envelope, tool_context_from_envelope
from opensquilla.safety.permission_matrix import Principal, is_tool_allowed
from opensquilla.safety.tool_tiers import RiskTier, get_tier
from opensquilla.sandbox.backend import NoopBackend, SeatbeltBackend, UnavailableBackend
from opensquilla.sandbox.config import SandboxSettings
from opensquilla.sandbox.integration import (
    active_file_system_profile,
    configure_runtime,
    reset_runtime,
)
from opensquilla.sandbox.permissions import FileSystemAccess
from opensquilla.sandbox.run_context import MountGrant, RunContext, get_run_context
from opensquilla.sandbox.run_mode import RunMode
from opensquilla.sandbox.types import NetworkMode, SandboxBackendError, SandboxResult
from opensquilla.tools.builtin import code_exec, filesystem
from opensquilla.tools.dispatch import build_tool_handler
from opensquilla.tools.registry import ToolRegistry
from opensquilla.tools.types import (
    ToolContext,
    ToolSpec,
    WorkspaceAccessError,
    current_tool_context,
)
from opensquilla.tools.workspace_authoring import (
    WORKSPACE_AUTHORING_TOOLS,
    workspace_authoring_attested,
)


class AvailableSeatbelt(SeatbeltBackend):
    def __init__(self) -> None:
        super().__init__()
        self.requests = []
        self.failure = False

    def available(self) -> bool:
        return True

    async def run(self, request):
        self.requests.append(request)
        if self.failure:
            raise SandboxBackendError("synthetic backend failure")
        return SandboxResult(
            returncode=0,
            stdout="ok",
            stderr="",
            wall_time_s=0,
            backend_used="seatbelt",
            policy_used=request.policy.summary(),
        )


def _directory_link(link: Path, target: Path) -> None:
    """Exercise real directory aliases without requiring Windows symlink privilege."""
    try:
        link.symlink_to(target, target_is_directory=True)
    except OSError as exc:
        if sys.platform != "win32" or getattr(exc, "winerror", None) != 1314:
            raise
        result = subprocess.run(
            ["cmd.exe", "/d", "/c", "mklink", "/J", str(link), str(target)],
            capture_output=True,
            text=True,
            check=False,
        )
        assert result.returncode == 0, result.stderr or result.stdout
        assert link.is_junction()
        assert link.lstat().st_reparse_tag == stat.IO_REPARSE_TAG_MOUNT_POINT
    else:
        assert link.is_symlink()
    assert link.resolve(strict=True) == target.resolve(strict=True)
    assert link.samefile(target)


@pytest.fixture
def channel_workspace(tmp_path):
    prepared = prepare_managed_workspace(tmp_path / "profile")
    workspace = Path(prepared.binding["root"])
    runtime = configure_runtime(
        SandboxSettings(sandbox=True, backend="noop", security_grading=False),
        workspace=workspace,
    )
    backend = AvailableSeatbelt()
    runtime.backend = backend
    yield workspace, runtime, backend
    reset_runtime()


def _envelope(workspace: Path, *, kind="managed", mounts=()):
    envelope = build_channel_route_envelope(
        IncomingMessage(sender_id="member-a", channel_id="group-a", content="make a report"),
        session_key="agent:main:feishu:group-a:member-a",
        session_id="session-a",
        session_prefix="feishu",
    )
    apply_run_context_route_metadata(
        envelope,
        RunContext(
            run_mode=RunMode.SAFE,
            workspace=str(workspace),
            workspace_binding_kind=kind,
            mounts=mounts,
        ),
        principal_is_owner=False,
    )
    return envelope


def _context(workspace: Path, *, kind="managed", mounts=()) -> ToolContext:
    return tool_context_from_envelope(
        _envelope(workspace, kind=kind, mounts=mounts),
        workspace_dir=str(workspace),
        workspace_strict=True,
    )


def test_managed_real_backend_attests_bounded_workspace(channel_workspace):
    workspace, _runtime, _backend = channel_workspace
    ctx = _context(workspace)
    assert workspace_authoring_attested(ctx)
    assert ctx.workspace_lockdown
    assert ctx.environment is None
    token = current_tool_context.set(ctx)
    try:
        profile = active_file_system_profile(workspace)
    finally:
        current_tool_context.reset(token)
    assert profile is not None
    assert profile.default_access is FileSystemAccess.DENY
    assert profile.resolve(workspace / "report.xlsx") is FileSystemAccess.WRITE
    assert (
        profile.resolve(workspace.parent / "other-session" / "secret.txt") is FileSystemAccess.DENY
    )


@pytest.mark.parametrize("kind", [None, "configured"])
def test_legacy_or_configured_workspace_cannot_attest(channel_workspace, kind):
    workspace, _runtime, _backend = channel_workspace
    assert not workspace_authoring_attested(_context(workspace, kind=kind))


def test_mount_grants_cannot_attest(channel_workspace):
    workspace, _runtime, _backend = channel_workspace
    assert not workspace_authoring_attested(
        _context(workspace, mounts=(MountGrant(path=str(workspace.parent)),))
    )


@pytest.mark.parametrize("setting", ["extra_ro_mounts", "extra_rw_mounts"])
def test_runtime_host_mounts_cannot_attest(channel_workspace, setting):
    workspace, runtime, _backend = channel_workspace
    setattr(runtime.settings, setting, [str(workspace.parent)])
    assert not workspace_authoring_attested(_context(workspace))


def test_pinned_operator_restrictions_only_narrow_authoring(channel_workspace):
    from types import SimpleNamespace

    from opensquilla.sandbox.policy_models import SandboxPolicy
    from opensquilla.tools.workspace_authoring import channel_workspace_file_system

    workspace, runtime, _backend = channel_workspace
    hidden = workspace / "hidden.txt"
    runtime.settings.denied_read_roots = [str(hidden)]
    ctx = _context(workspace)
    outside = workspace.parent / "other-task"
    readonly = workspace / "readonly"
    ctx.sandbox_gateway_config = SimpleNamespace(state_dir=str(workspace.parent / "state"))
    ctx.sandbox_policy = SandboxPolicy.model_validate(
        {"files": {"custom_deny_write_paths": [str(readonly), str(outside)]}}
    )
    profile = channel_workspace_file_system(ctx)
    assert profile.resolve(workspace / "report.txt") is FileSystemAccess.WRITE
    assert profile.resolve(readonly / "report.txt") is FileSystemAccess.READ
    assert profile.resolve(hidden) is FileSystemAccess.DENY
    assert profile.resolve(outside / "data.txt") is FileSystemAccess.DENY
    ctx.sandbox_policy = SandboxPolicy.model_validate(
        {"files": {"custom_deny_write_paths": [str(workspace.parent)]}}
    )
    assert (
        channel_workspace_file_system(ctx).resolve(workspace / "report.txt")
        is FileSystemAccess.READ
    )


@pytest.mark.parametrize(
    ("denied_relative", "readonly_relative", "target_relative"),
    [
        ("private", "private", "private/secret.txt"),
        ("private", "private/readonly", "private/readonly/secret.txt"),
        ("private/secret.txt", "private", "private/secret.txt"),
        ("..", "readonly", "readonly/secret.txt"),
        ("..", None, "secret.txt"),
    ],
)
def test_operator_readonly_cannot_reopen_runtime_denial(
    channel_workspace,
    denied_relative,
    readonly_relative,
    target_relative,
):
    from opensquilla.sandbox.policy_models import SandboxPolicy
    from opensquilla.tools.workspace_authoring import channel_workspace_file_system

    workspace, runtime, _backend = channel_workspace
    runtime.settings.denied_read_roots = [str((workspace / denied_relative).resolve())]
    ctx = _context(workspace)
    if readonly_relative is not None:
        ctx.sandbox_policy = SandboxPolicy.model_validate(
            {
                "files": {"custom_deny_write_paths": [str(workspace / readonly_relative)]},
            }
        )
    profile = channel_workspace_file_system(ctx)
    assert profile.resolve(workspace / target_relative) is FileSystemAccess.DENY
    # Backends must not receive a more specific grant that could reopen a
    # denied ancestor, even if a renderer ignores the high-level resolver.
    denied_root = (workspace / denied_relative).resolve()
    assert not any(
        entry.access is not FileSystemAccess.DENY and entry.path.is_relative_to(denied_root)
        for entry in profile.effective_entries
    )


def test_metadata_cannot_supply_workspace_authority(channel_workspace):
    workspace, _runtime, _backend = channel_workspace
    envelope = _envelope(workspace, kind=None)
    envelope.metadata.update(
        {
            "sandboxed_workspace_authoring": True,
            "workspace_binding_kind": "managed",
            "sandbox_backend_ready": True,
        }
    )
    envelope.metadata["sandbox_run_context"]["workspace_binding_kind"] = "managed"
    ctx = tool_context_from_envelope(
        envelope,
        workspace_dir=str(workspace),
        workspace_strict=True,
    )
    assert ctx.sandboxed_workspace_authoring is None


@pytest.mark.parametrize("backend", [NoopBackend(), UnavailableBackend("test")])
def test_unavailable_or_noop_backend_cannot_attest(channel_workspace, backend):
    workspace, runtime, _backend = channel_workspace
    runtime.backend = backend
    assert not workspace_authoring_attested(_context(workspace))


def test_windows_authoring_stays_closed_until_process_boundary_supported(channel_workspace):
    from opensquilla.sandbox.backend import WindowsDefaultBackend

    workspace, runtime, _backend = channel_workspace
    runtime.backend = WindowsDefaultBackend()
    assert not workspace_authoring_attested(_context(workspace))


@pytest.mark.parametrize("change", ["workspace", "session", "backend", "mode", "lockdown"])
def test_changed_execution_facts_revoke_attestation(channel_workspace, change):
    workspace, runtime, _backend = channel_workspace
    ctx = _context(workspace)
    assert workspace_authoring_attested(ctx)
    if change == "workspace":
        ctx.workspace_dir = str(workspace.parent)
    elif change == "session":
        ctx.session_id = "session-b"
    elif change == "backend":
        runtime.backend = NoopBackend()
    elif change == "mode":
        ctx.run_mode = "full"
    else:
        ctx.workspace_lockdown = False
    assert not workspace_authoring_attested(ctx)


def test_replaced_workspace_revokes_attestation(channel_workspace):
    workspace, _runtime, _backend = channel_workspace
    ctx = _context(workspace)
    workspace.rename(workspace.with_name("old-root"))
    workspace.mkdir(mode=0o700)
    assert not workspace_authoring_attested(ctx)


def test_directory_link_workspace_cannot_attest(channel_workspace):
    workspace, _runtime, _backend = channel_workspace
    assert workspace_authoring_attested(_context(workspace))
    link = workspace.with_name("alias")
    _directory_link(link, workspace)
    assert not workspace_authoring_attested(_context(link))


def test_permission_matrix_only_relaxes_exact_workspace_tools():
    for name in WORKSPACE_AUTHORING_TOOLS - {"read_file"}:
        assert get_tier(name) is RiskTier.ADMIN_ONLY
        assert not is_tool_allowed(name, "group", Principal()).allowed
        assert is_tool_allowed(
            name,
            "group",
            Principal(),
            workspace_authoring_attested=True,
        ).allowed
    for name in ("exec_command", "background_process", "git_commit", "git_push"):
        assert not is_tool_allowed(
            name,
            "group",
            Principal(),
            workspace_authoring_attested=True,
        ).allowed


def test_catalog_and_search_use_the_same_authoring_capability(channel_workspace):
    workspace, _runtime, _backend = channel_workspace
    registry = ToolRegistry()

    async def handler():
        return "ok"

    for name in (*WORKSPACE_AUTHORING_TOOLS, "exec_command", "background_process", "tool_search"):
        registry.register(ToolSpec(name=name, description=name, parameters={}), handler)
    ctx = _context(workspace)
    definitions = registry.to_tool_definitions(ctx)
    names = {item.name for item in definitions}
    assert WORKSPACE_AUTHORING_TOOLS <= names
    assert not {"exec_command", "background_process"} & names
    registry.to_model_tool_definitions(definitions, ctx)
    assert ctx.authorized_tool_names == names
    assert ctx.tool_search_index is not None
    denied = _context(workspace, kind="configured")
    assert not (WORKSPACE_AUTHORING_TOOLS - {"read_file"}) & {
        item.name for item in registry.to_tool_definitions(denied)
    }


@pytest.mark.asyncio
async def test_dispatch_rejects_elevation_even_when_tool_is_visible(channel_workspace):
    from opensquilla.tool_boundary import ToolCall

    workspace, _runtime, _backend = channel_workspace
    registry = ToolRegistry()
    called = False

    async def handler(**kwargs):
        nonlocal called
        called = True
        return "ok"

    registry.register(
        ToolSpec(
            name="execute_code",
            description="code",
            parameters={"sandbox_permissions": {"type": "string"}},
        ),
        handler,
    )
    dispatch = build_tool_handler(registry)
    token = current_tool_context.set(_context(workspace))
    try:
        result = await dispatch(
            ToolCall(
                tool_use_id="call-a",
                tool_name="execute_code",
                arguments={"sandbox_permissions": "require_escalated"},
            )
        )
    finally:
        current_tool_context.reset(token)
    assert result.is_error
    assert not called


@pytest.mark.asyncio
async def test_channel_code_uses_clean_env_closed_profile_and_no_network(
    channel_workspace, monkeypatch
):
    workspace, _runtime, backend = channel_workspace
    monkeypatch.setenv("SYNTHETIC_PROVIDER_SECRET", "must-not-leak")
    token = current_tool_context.set(_context(workspace))
    try:
        output = json.loads(await code_exec.execute_code("print('ok')"))
    finally:
        current_tool_context.reset(token)
    assert output["exit_code"] == 0
    request = backend.requests[0]
    assert request.policy.network is NetworkMode.NONE
    assert request.policy.file_system.default_access is FileSystemAccess.DENY
    assert request.policy.file_system.resolve(workspace.parent / "secrets") is FileSystemAccess.DENY
    assert request.env["HOME"] == str(workspace)
    assert "SYNTHETIC_PROVIDER_SECRET" not in request.env
    assert set(request.env) == set(request.policy.env_allowlist)
    assert request.argv[1:3] == ("-I", "-c")


@pytest.mark.asyncio
async def test_channel_code_preserves_future_imports_and_module_docstrings(
    channel_workspace, monkeypatch
):
    workspace, _runtime, backend = channel_workspace
    code = (
        '"""Authoring script."""\n'
        "from __future__ import annotations\n"
        "def render(value: MissingDependency):\n"
        "    return value\n"
    )
    token = current_tool_context.set(_context(workspace))
    try:
        await code_exec.execute_code(code)
    finally:
        current_tool_context.reset(token)
    request = backend.requests[0]
    limits = []
    monkeypatch.setitem(
        sys.modules, "resource",
        SimpleNamespace(
            RLIMIT_CPU=0, RLIM_INFINITY=-1,
            getrlimit=lambda kind: (-1, -1),
            setrlimit=lambda kind, value: limits.append((kind, value)),
        ),
    )
    namespace = {"__name__": "__main__"}
    exec(compile(request.argv[-1], "<authoring-bootstrap>", "exec"), namespace)
    assert namespace["__doc__"] == "Authoring script."
    assert namespace["render"].__annotations__ == {"value": "MissingDependency"}
    assert limits == [(0, (request.policy.limits.cpu_seconds, request.policy.limits.cpu_seconds))]


@pytest.mark.asyncio
@pytest.mark.skipif(sys.platform == "win32", reason="POSIX inherited resource limit required")
@pytest.mark.parametrize("inherited_limit", [5, 60])
async def test_channel_code_respects_inherited_cpu_ceiling(channel_workspace, inherited_limit):
    """Execute the emitted bootstrap under a real lower or higher CPU ceiling."""
    import resource

    workspace, runtime, backend = channel_workspace
    runtime.settings.cpu_seconds = 30
    original_limits = resource.getrlimit(resource.RLIMIT_CPU)
    inherited_hard = original_limits[1]
    if inherited_hard != resource.RLIM_INFINITY:
        inherited_limit = min(inherited_limit, inherited_hard)
    code = (
        "from __future__ import annotations\n"
        "import json, resource\n"
        "print(json.dumps(resource.getrlimit(resource.RLIMIT_CPU)))\n"
    )
    token = current_tool_context.set(_context(workspace))
    try:
        await code_exec.execute_code(code)
    finally:
        current_tool_context.reset(token)
    request = backend.requests[0]
    # Change the ceiling in a separate interpreter, then exec the exact tool
    # argv. The pytest process keeps its original resource limits.
    launcher = (
        "import os, resource\n"
        f"resource.setrlimit(resource.RLIMIT_CPU, ({inherited_limit}, {inherited_limit}))\n"
        f"os.execv({request.argv[0]!r}, {request.argv!r})\n"
    )
    result = subprocess.run(
        [sys.executable, "-I", "-c", launcher],
        cwd=workspace, env=request.env,
        text=True, capture_output=True, timeout=10, check=False,
    )
    assert result.returncode == 0, result.stderr
    expected = min(runtime.settings.cpu_seconds, inherited_limit)
    assert json.loads(result.stdout) == [expected, expected]
    assert resource.getrlimit(resource.RLIMIT_CPU) == original_limits


@pytest.mark.asyncio
@pytest.mark.parametrize("denial", [None, "parent", "root", "alias_parent", "canonical_parent"])
async def test_channel_runtime_mounts_preserve_denial_through_linux_planner(
    channel_workspace, tmp_path, monkeypatch, denial
):
    """Compile the actual tool request; a filtered grant must not reappear as a mount."""
    from opensquilla.sandbox.backend.bubblewrap import build_bwrap_plan
    from opensquilla.sandbox.backend.linux_permissions import compile_linux_permissions

    workspace, runtime, backend = channel_workspace
    parent = tmp_path / "operator-private"
    runtime_root = parent / "python-runtime"
    runtime_root.mkdir(parents=True)
    alias = tmp_path / "runtime-alias"
    if denial in {"alias_parent", "canonical_parent"}:
        _directory_link(alias, parent)
    exposed_root = alias / runtime_root.name if denial == "canonical_parent" else runtime_root
    denied = (
        alias if denial == "alias_parent" else runtime_root if denial == "root" else parent
    )
    if denial is not None:
        runtime.settings.denied_read_roots = [str(denied)]
    monkeypatch.setattr(
        code_exec, "_current_python_runtime_roots", lambda **kwargs: (exposed_root,)
    )
    token = current_tool_context.set(_context(workspace))
    try:
        await code_exec.execute_code("print('ok')")
    finally:
        current_tool_context.reset(token)
    request = backend.requests[0]
    profile = request.policy.file_system
    assert profile is not None
    assert profile.resolve(workspace / "report.csv") is FileSystemAccess.WRITE
    assert profile.resolve(runtime_root / "module.py") is (
        FileSystemAccess.READ if denial is None else FileSystemAccess.DENY
    )
    permissions = compile_linux_permissions(request.policy)
    assert any(root.host_path == workspace for root in permissions.write_roots)
    plan = build_bwrap_plan(request)
    try:
        bind_targets = [
            Path(plan.argv[index + 2])
            for index, argument in enumerate(plan.argv[:-2])
            if argument in {"--bind", "--ro-bind", "--bind-try", "--ro-bind-try"}
        ]
        if denial is None:
            assert any(root.host_path == runtime_root for root in permissions.read_roots)
            assert runtime_root in bind_targets
        else:
            denied_root = denied.resolve()
            assert not any(
                root.host_path.resolve().is_relative_to(denied_root)
                for root in (*permissions.read_roots, *permissions.write_roots)
            )
            assert not any(target.resolve().is_relative_to(denied_root) for target in bind_targets)
            assert any(
                plan.argv[index : index + 2] == ["--tmpfs", str(denied_root)]
                for index in range(len(plan.argv) - 1)
            )
    finally:
        for stream in plan.preserved_files:
            stream.close()


@pytest.mark.asyncio
@pytest.mark.skipif(not sys.platform.startswith("linux"), reason="native Bubblewrap required")
async def test_real_bubblewrap_channel_runtime_denial_is_not_reopened(tmp_path, monkeypatch):
    from opensquilla.sandbox.backend.linux_readiness import probe_bwrap

    readiness = probe_bwrap()
    if not readiness.available:
        pytest.skip(readiness.message)
    workspace = Path(prepare_managed_workspace(tmp_path / "profile").binding["root"])
    denied_parent = tmp_path / "private-install"
    runtime_root = denied_parent / "python-runtime"
    runtime_root.mkdir(parents=True)
    secret = runtime_root / "data.txt"
    secret.write_text("synthetic runtime data", encoding="utf-8")
    roots = code_exec._current_python_runtime_roots(workspace=workspace)
    monkeypatch.setattr(
        code_exec, "_current_python_runtime_roots", lambda **kwargs: (*roots, runtime_root)
    )

    async def forbidden(*args, **kwargs):
        raise AssertionError("host fallback or approval must not run")

    monkeypatch.setattr(code_exec, "create_owned_subprocess_exec", forbidden)
    monkeypatch.setattr(code_exec, "escalate_unavailable_backend_in_managed_mode", forbidden)
    code = (
        "from pathlib import Path\n"
        "Path('output.txt').write_text('workspace write works')\n"
        "try:\n"
        f"    Path({str(secret)!r}).read_text()\n"
        "    print('READ_ALLOWED')\n"
        "except (PermissionError, FileNotFoundError):\n"
        "    print('READ_DENIED')\n"
    )
    try:
        for denied in (False, True):
            configure_runtime(
                SandboxSettings(
                    sandbox=True,
                    backend="bubblewrap",
                    security_grading=False,
                    denied_read_roots=[str(denied_parent)] if denied else [],
                ),
                workspace=workspace,
            )
            token = current_tool_context.set(_context(workspace))
            try:
                output = json.loads(await code_exec.execute_code(code))
            finally:
                current_tool_context.reset(token)
            assert output["exit_code"] == 0, output
            assert output["stdout"].strip() == ("READ_DENIED" if denied else "READ_ALLOWED")
            assert (workspace / "output.txt").read_text() == "workspace write works"
    finally:
        reset_runtime()


@pytest.mark.asyncio
async def test_channel_backend_failure_has_no_host_retry(channel_workspace, monkeypatch):
    workspace, _runtime, backend = channel_workspace
    backend.failure = True

    async def forbidden(*args, **kwargs):
        raise AssertionError("host fallback or approval must not run")

    monkeypatch.setattr(code_exec, "create_owned_subprocess_exec", forbidden)
    monkeypatch.setattr(code_exec, "escalate_unavailable_backend_in_managed_mode", forbidden)
    token = current_tool_context.set(_context(workspace))
    try:
        output = json.loads(await code_exec.execute_code("print('ok')"))
    finally:
        current_tool_context.reset(token)
    assert output["exit_code"] == -1
    assert "Sandbox execution unavailable" in output["stderr"]


@pytest.mark.asyncio
async def test_direct_file_handler_cannot_write_other_session(channel_workspace):
    workspace, _runtime, _backend = channel_workspace
    token = current_tool_context.set(_context(workspace))
    try:
        with pytest.raises(WorkspaceAccessError):
            await filesystem.write_file(str(workspace.parent / "other-session.txt"), "blocked")
    finally:
        current_tool_context.reset(token)
    assert not (workspace.parent / "other-session.txt").exists()


@pytest.mark.asyncio
async def test_get_run_context_binding_fact_is_runtime_only(tmp_path):
    from types import SimpleNamespace

    prepared = prepare_managed_workspace(tmp_path / "profile")
    node = SimpleNamespace(
        execution_workspace=prepared.binding,
        workspace_id=None,
        origin={},
    )
    config = SimpleNamespace(sandbox=SandboxSettings(sandbox=True))
    context = await get_run_context(
        None,
        "session-a",
        config=config,
        workspace=None,
        session_node=node,
        include_user_grants=False,
    )
    assert context.workspace_binding_kind == "managed"
    assert context.workspace == prepared.binding["root"]
    assert "workspace_binding_kind" not in context.to_origin_payload()
    node.execution_workspace = None
    node.origin = {
        "sandbox_run_context": {
            **context.to_origin_payload(),
            "workspace_binding_kind": "managed",
        }
    }
    context = await get_run_context(
        None,
        "session-a",
        config=config,
        workspace=None,
        session_node=node,
        include_user_grants=False,
    )
    assert context.workspace_binding_kind is None


def test_channel_metadata_cannot_impersonate_webui(channel_workspace):
    workspace, _runtime, _backend = channel_workspace
    envelope = _envelope(workspace, kind="configured")
    envelope.metadata.update(
        {
            "tool_source_kind": "webui",
            "tool_source_name": "webui",
            "principal_is_owner": True,
            "sandboxed_workspace_authoring": True,
        }
    )
    ctx = tool_context_from_envelope(
        envelope,
        is_owner=True,
        workspace_dir=str(workspace),
        workspace_strict=True,
    )
    assert ctx.source_kind == "channel"
    assert ctx.source_name == "feishu"
    assert ctx.is_owner is False
    assert ctx.sandboxed_workspace_authoring is None


@pytest.mark.asyncio
async def test_forged_owner_and_webui_claims_cannot_bypass_matrix(channel_workspace):
    from opensquilla.tool_boundary import ToolCall
    from opensquilla.tools.types import CallerKind

    _workspace, _runtime, _backend = channel_workspace
    registry = ToolRegistry()
    called = False

    async def handler():
        nonlocal called
        called = True
        return "ok"

    registry.register(ToolSpec(name="write_file", description="write", parameters={}), handler)
    ctx = ToolContext(
        caller_kind=CallerKind.CHANNEL,
        is_owner=True,
        source_kind="webui",
        channel_kind="webui",
        allowed_tools={"write_file"},
    )
    assert not registry.to_tool_definitions(ctx)
    token = current_tool_context.set(ctx)
    try:
        result = await build_tool_handler(registry)(
            ToolCall(
                tool_use_id="forged-call",
                tool_name="write_file",
                arguments={},
            )
        )
    finally:
        current_tool_context.reset(token)
    assert result.is_error
    assert not called


def test_false_backend_readiness_cannot_attest(channel_workspace, monkeypatch):
    workspace, _runtime, backend = channel_workspace
    monkeypatch.setattr(backend, "available", lambda: False)
    assert not workspace_authoring_attested(_context(workspace))


@pytest.mark.asyncio
@pytest.mark.parametrize("inherited", [False, True])
async def test_production_task_dispatch_attests_managed_channel_session(
    channel_workspace,
    tmp_path,
    inherited,
):
    from types import SimpleNamespace

    from opensquilla.engine.types import DoneEvent
    from opensquilla.gateway.boot import dispatch_task_runtime_turn
    from opensquilla.gateway.config import GatewayConfig
    from opensquilla.gateway.execution_workspaces import build_execution_workspace_factory
    from opensquilla.session.manager import SessionManager
    from opensquilla.session.storage import SessionStorage

    _workspace, _runtime, _backend = channel_workspace
    config = GatewayConfig(
        sandbox={"run_mode": "safe"},
        state_dir=str(tmp_path / "state"),
        attachments={"media_root": str(tmp_path / "media")},
        agent_stream_heartbeat_interval_seconds=0,
    )
    assert config.workspace_strict is None
    captured = []

    class Runner:
        async def run(self, message, session_key, **kwargs):
            captured.append(kwargs["tool_context"])
            yield DoneEvent()

    async def emit(*args):
        pass

    async with SessionStorage(tmp_path / "sessions.db") as storage:
        manager = SessionManager(
            storage,
            execution_workspace_factory=build_execution_workspace_factory(
                config,
                profile_home=tmp_path / "managed-profile",
            ),
        )
        parent = await manager.create("agent:main:feishu:group-a:member-a")
        session = (
            await manager.branch(parent.session_key, "agent:main:feishu:group-a:branch-a")
            if inherited
            else parent
        )
        if inherited:
            assert session.parent_session_key == parent.session_key
            assert session.execution_workspace == parent.execution_workspace
            assert session.execution_workspace is not parent.execution_workspace
        workspace = Path(session.execution_workspace["root"])
        envelope = build_channel_route_envelope(
            IncomingMessage(
                sender_id="member-a",
                channel_id="group-a",
                content="make a report",
                metadata={"is_group": True},
            ),
            session_key=session.session_key,
            session_id=session.session_id,
            session_epoch=session.epoch,
            session_prefix="feishu",
        )
        run = SimpleNamespace(
            agent_id="main",
            task_id="task-a",
            session_key=session.session_key,
            session_id=session.session_id,
            session_epoch=session.epoch,
            message="make a report",
            envelope=envelope,
            attachments=[],
            input_provenance={},
            run_kind="interactive",
            no_memory_capture=False,
            ingress_pipeline_steps=[],
            semantic_message=None,
            stream_event_sink=None,
        )
        await dispatch_task_runtime_turn(
            run,
            config=config,
            session_manager=manager,
            turn_runner=Runner(),
            event_emitter=emit,
        )
    assert len(captured) == 1
    assert captured[0].workspace_dir == str(workspace)
    assert captured[0].workspace_strict is True
    assert workspace_authoring_attested(captured[0]) is not inherited


@pytest.mark.asyncio
async def test_direct_channel_code_rejects_elevation(channel_workspace):
    workspace, _runtime, backend = channel_workspace
    token = current_tool_context.set(_context(workspace))
    try:
        with pytest.raises(Exception, match="cannot request host execution"):
            await code_exec.execute_code(
                "print('ok')",
                sandbox_permissions="require_escalated",
                justification="synthetic request",
            )
    finally:
        current_tool_context.reset(token)
    assert backend.requests == []


@pytest.mark.asyncio
async def test_attested_read_cannot_access_other_workspace(channel_workspace):
    workspace, _runtime, _backend = channel_workspace
    other = workspace.parent / "other-task.txt"
    other.write_text("other task data")
    token = current_tool_context.set(_context(workspace))
    try:
        with pytest.raises(WorkspaceAccessError):
            await filesystem.read_file(str(other))
    finally:
        current_tool_context.reset(token)


@pytest.mark.asyncio
@pytest.mark.skipif(sys.platform != "darwin", reason="native Seatbelt required")
async def test_real_seatbelt_channel_authoring_and_process_boundary(tmp_path, monkeypatch):
    """Exercise the gateway proof and actual kernel boundary, including children."""
    from opensquilla.sandbox.policy_models import SandboxPolicy

    if not SeatbeltBackend().available():
        pytest.skip("requires macOS sandbox-exec")
    for library in ("openpyxl", "pptx", "reportlab"):
        pytest.importorskip(library)
    workspace = Path(prepare_managed_workspace(tmp_path / "profile").binding["root"])
    outside = tmp_path / "other-session.txt"
    outside.write_text("private neighboring task", encoding="utf-8")
    outside_write = tmp_path / "outside-output.txt"
    denied_file = workspace / "private.txt"
    denied_file.write_text("runtime denied data", encoding="utf-8")
    monkeypatch.setenv("SYNTHETIC_CHANNEL_HOST_SECRET", "must-not-leak")
    configure_runtime(
        SandboxSettings(
            sandbox=True,
            backend="seatbelt",
            security_grading=False,
            denied_read_roots=[str(denied_file)],
        ),
        workspace=workspace,
    )
    ctx = _context(workspace)
    ctx.sandbox_policy = SandboxPolicy.model_validate(
        {
            "files": {"custom_deny_write_paths": [str(denied_file)]},
        }
    )
    assert workspace_authoring_attested(ctx)
    listener = socket.socket()
    listener.bind(("127.0.0.1", 0))
    listener.listen(1)
    port = listener.getsockname()[1]
    code = f"""
import csv, json, os, socket, subprocess, sys
from pathlib import Path
from openpyxl import Workbook
from pptx import Presentation
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.pdfgen.canvas import Canvas

with open("table.csv", "w", encoding="utf-8-sig", newline="") as f:
    writer = csv.writer(f)
    writer.writerow(["项目", "数量"])
    writer.writerow(["样例", 3])
book = Workbook()
book.active.append(["数值", "公式"])
book.active.append([3, "=A2*2"])
book.save("table.xlsx")
slides = Presentation()
slide = slides.slides.add_slide(slides.slide_layouts[1])
slide.shapes.title.text = "渠道报告"
slides.save("slides.pptx")
pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
pdf = Canvas("report.pdf")
pdf.setFont("STSong-Light", 16)
pdf.drawString(72, 720, "中文报告")
pdf.save()
checks = {{"host_secret_absent": "SYNTHETIC_CHANNEL_HOST_SECRET" not in os.environ}}
try:
    Path({str(denied_file)!r}).read_text()
except PermissionError:
    checks["runtime_denial_preserved"] = True
try:
    Path({str(outside)!r}).read_text()
except PermissionError:
    checks["outside_read_denied"] = True
try:
    Path({str(outside_write)!r}).write_text("escape")
except PermissionError:
    checks["outside_write_denied"] = True
try:
    socket.create_connection(("127.0.0.1", {port}), timeout=2)
except PermissionError:
    checks["network_denied"] = True
child = subprocess.run(
    [sys.executable, "-I", "-c", "from pathlib import Path; Path({str(outside)!r}).read_text()"],
    capture_output=True, text=True,
)
checks["child_read_denied"] = child.returncode != 0 and "PermissionError" in child.stderr
print(json.dumps(checks))
"""
    token = current_tool_context.set(ctx)
    try:
        result = json.loads(await code_exec.execute_code(code, timeout=30))
    finally:
        current_tool_context.reset(token)
        listener.close()
        reset_runtime()
    assert result["exit_code"] == 0, result["stderr"]
    assert json.loads(result["stdout"]) == {
        "host_secret_absent": True,
        "outside_read_denied": True,
        "outside_write_denied": True,
        "network_denied": True,
        "child_read_denied": True,
        "runtime_denial_preserved": True,
    }
    assert not outside_write.exists()
    assert "样例" in (workspace / "table.csv").read_text(encoding="utf-8-sig")
    assert zipfile.is_zipfile(workspace / "table.xlsx")
    with zipfile.ZipFile(workspace / "slides.pptx") as archive:
        assert "渠道报告" in archive.read("ppt/slides/slide1.xml").decode()
    assert (workspace / "report.pdf").read_bytes().startswith(b"%PDF-")
