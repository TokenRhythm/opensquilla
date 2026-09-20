"""Offline packaged smoke: generate and publish documents through the real tools."""

from __future__ import annotations

import asyncio
import contextlib
import csv
import io
import json
import os
import sys
from pathlib import Path
from typing import Any

_TITLE = "Packaged Python tool smoke"
_PDF_TEXT = "中文文件验收 样本 42"
_DOCUMENT_CODE = '''import csv
import json
import sys
from openpyxl import Workbook
from pptx import Presentation
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.pdfgen import canvas

with open("python-tool-smoke.csv", "w", newline="", encoding="utf-8-sig") as output:
    csv.writer(output).writerows([["名称", "数量"], ["验收样本", "42"]])
workbook = Workbook()
workbook.active.append(["名称", "数量"])
workbook.active.append(["验收样本", 42])
workbook.save("python-tool-smoke.xlsx")

presentation = Presentation()
slide = presentation.slides.add_slide(presentation.slide_layouts[0])
slide.shapes.title.text = "Packaged Python tool smoke"
presentation.save("python-tool-smoke.pptx")
reopened = Presentation("python-tool-smoke.pptx")
assert len(reopened.slides) == 1
assert reopened.slides[0].shapes.title.text == "Packaged Python tool smoke"
pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
document = canvas.Canvas("python-tool-smoke.pdf")
document.setFont("STSong-Light", 18)
document.drawString(50, 780, "中文文件验收 样本 42")
document.save()
print(json.dumps({"executable": sys.executable, "frozen": bool(getattr(sys, "frozen", False))}))
'''


def _execution_result(raw: str, expected_exit: int) -> dict[str, Any]:
    result = json.loads(raw)
    if (
        not isinstance(result, dict)
        or result.get("exit_code") != expected_exit
        or result.get("timed_out") is not False
    ):
        raise RuntimeError(f"execute_code expected exit {expected_exit}, received {result!r}")
    return result


async def _probe() -> dict[str, Any]:
    from opensquilla.artifacts import ArtifactStore
    from opensquilla.sandbox.config import SandboxSettings
    from opensquilla.sandbox.integration import configure_runtime, reset_runtime
    from opensquilla.tools.builtin.artifacts import publish_artifact
    from opensquilla.tools.builtin.code_exec import execute_code
    from opensquilla.tools.types import CallerKind, ToolContext, current_tool_context

    workspace = Path(os.environ["OPENSQUILLA_STATE_DIR"]) / "workspace" / "code-execution"
    media_root = workspace.parent / "media"
    workspace.mkdir(parents=True, exist_ok=True)
    # Full grants host access; network_default is not an enforced Full-mode
    # network boundary. Both fixed snippets are local and make no network calls.
    configure_runtime(
        SandboxSettings(run_mode="full", network_default="none"),
        workspace=workspace,
    )
    context = ToolContext(
        is_owner=True,
        caller_kind=CallerKind.CLI,
        session_key="packaged-code-execution-smoke",
        run_mode="full",
        workspace_dir=str(workspace),
        artifact_media_root=str(media_root),
        artifact_session_id="packaged-code-execution-smoke",
    )
    token = current_tool_context.set(context)
    try:
        success = _execution_result(await execute_code(_DOCUMENT_CODE, timeout=15), 0)
        child = json.loads(success["stdout"])
        if (
            Path(child["executable"]).resolve() != Path(sys.executable).resolve()
            or child["frozen"] is not bool(getattr(sys, "frozen", False))
        ):
            raise RuntimeError("execute_code selected a different Python runtime")
        failed = _execution_result(await execute_code("raise SystemExit(7)", timeout=15), 7)
        store = ArtifactStore(media_root)
        stored: dict[str, bytes] = {}
        for suffix in ("csv", "xlsx", "pptx", "pdf"):
            filename = f"python-tool-smoke.{suffix}"
            published = json.loads(await publish_artifact(filename))
            repeated = json.loads(await publish_artifact(filename))
            if (
                published.get("status") != "published"
                or repeated.get("status") != "already_published"
                or published["artifact"]["id"] != repeated["artifact"]["id"]
            ):
                raise RuntimeError(f"Document publication was not idempotent: {filename}")
            reference = store.get_ref(
                session_id="packaged-code-execution-smoke",
                artifact_id=published["artifact"]["id"],
            )
            stored[suffix] = store.path_for(reference).read_bytes()
            if stored[suffix] != (workspace / filename).read_bytes():
                raise RuntimeError(f"Publication changed document bytes: {filename}")
        if len(context.published_artifacts) != 4:
            raise RuntimeError("Repeated publication created duplicate artifacts")
    finally:
        current_tool_context.reset(token)
        reset_runtime()

    import pdfplumber
    from openpyxl import load_workbook
    from pptx import Presentation

    rows = list(csv.reader(io.StringIO(stored["csv"].decode("utf-8-sig"))))
    if rows != [["名称", "数量"], ["验收样本", "42"]]:
        raise RuntimeError("CSV content did not survive publication")
    workbook = load_workbook(io.BytesIO(stored["xlsx"]))
    try:
        if workbook.active["A2"].value != "验收样本" or workbook.active["B2"].value != 42:
            raise RuntimeError("XLSX content did not survive publication")
    finally:
        workbook.close()
    with pdfplumber.open(io.BytesIO(stored["pdf"])) as document:
        if len(document.pages) != 1 or document.pages[0].extract_text() != _PDF_TEXT:
            raise RuntimeError("Chinese PDF content did not survive publication")
    presentation = Presentation(io.BytesIO(stored["pptx"]))
    pages = len(presentation.slides)
    title = presentation.slides[0].shapes.title.text if pages == 1 else None
    if pages != 1 or title != _TITLE:
        raise RuntimeError("execute_code did not produce the expected one-slide presentation")
    return {
        "probe": "opensquilla-desktop-code-execution",
        "frozen": bool(getattr(sys, "frozen", False)),
        "pythonExit": success["exit_code"],
        "errorExit": failed["exit_code"],
        "pages": pages,
        "title": title,
        "documents": {"csvRows": len(rows), "xlsxValue": 42, "pdfText": _PDF_TEXT},
        "published": len(context.published_artifacts),
    }


async def main() -> None:
    with contextlib.redirect_stdout(sys.stderr):
        result = await _probe()
    print(json.dumps(result))


if __name__ == "__main__":
    asyncio.run(main())
