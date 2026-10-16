"""Generate synthetic Workbench preview fixtures without touching existing files.

Usage (from the repository root)::

    .venv/Scripts/python scripts/testing/create_workbench_preview_fixtures.py \
        --output .pytest-tmp/workbench-trial/workspace/preview-gallery

The destination must not exist. Optional document/image dependencies and FFmpeg
are used only when already installed; missing producers are recorded in manifest.
This script does not read user documents, fetch assets, or install dependencies.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import importlib.util
import io
import json
import math
import shutil
import struct
import subprocess
import wave
import zipfile
from datetime import UTC, datetime
from pathlib import Path
from typing import Any


class Gallery:
    def __init__(self, directory: Path) -> None:
        self.directory = directory
        self.entries: list[dict[str, Any]] = []
        self.skipped: list[dict[str, str]] = []

    def add(self, name: str, preview: str, mime: str, scenario: str, *, limits: str = "") -> Path:
        path = self.directory / name
        content = path.read_bytes()
        self.entries.append(
            {
                "path": name,
                "sizeBytes": len(content),
                "sha256": hashlib.sha256(content).hexdigest(),
                "expectedPreview": preview,
                "expectedMime": mime,
                "scenario": scenario,
                "limits": limits or "Below the current preview size limit.",
            }
        )
        return path

    def text(self, name: str, text: str, preview: str, mime: str, scenario: str, **kwargs) -> Path:
        (self.directory / name).write_text(text, encoding="utf-8", newline="")
        return self.add(name, preview, mime, scenario, **kwargs)

    def dependency(self, name: str) -> bool:
        if importlib.util.find_spec(name) is not None:
            return True
        self.skipped.append({"producer": name, "reason": "Optional dependency not installed"})
        return False


def make_text(g: Gallery) -> None:
    g.text(
        "01-welcome.md",
        "# Workbench preview gallery\n\nAll content is synthetic. 所有内容均为合成测试数据。\n\n"
        "## Markdown\n\n**Bold**, *italic*, `inline code`, and Unicode: "
        "中文 / 日本語 / café / 🌱.\n\n"
        "| Item | Quantity | Notes |\n| --- | ---: | --- |\n| Alpha | 12 | table cell |\n"
        "| 测试 | 3 | Unicode ✓ |\n\n- List item\n  - Nested item\n\n"
        "> A quotation with a [relative link](03-data.json).\n\n"
        "```python\ndef greet(name):\n    return f'Hello, {name}'\n```\n",
        "markdown",
        "text/markdown",
        "Short rendered Markdown with table, code, Unicode, and relative link.",
    )
    sections = ["# Long synthetic report\n"]
    for index in range(1, 91):
        sections.append(
            f"## Section {index:03d} — 分页测试\n\n"
            f"Synthetic paragraph {index}; café, 日本語, emoji 🌱.\n\n"
            "| Metric | Value |\n| --- | ---: |\n"
            f"| Sample | {index * 3} |\n\n"
        )
    g.text(
        "02-long-report.md",
        "\n".join(sections),
        "source-paged",
        "text/markdown",
        "Long Markdown tests source pagination and Unicode across page boundaries.",
        limits="More than 200 lines; current UI uses paged source instead of partial Markdown.",
    )
    g.text(
        "03-data.json",
        json.dumps(
            {
                "title": "合成数据",
                "enabled": True,
                "items": [{"id": n, "value": n / 3} for n in range(8)],
            },
            ensure_ascii=False,
            indent=2,
        )
        + "\n",
        "source",
        "application/json",
        "Nested JSON and non-ASCII property values.",
    )
    for name, content, mime, scenario in [
        (
            "04-settings.yaml",
            "title: 测试\nenabled: true\nitems:\n  - alpha\n  - beta\n",
            "application/yaml",
            "YAML source",
        ),
        (
            "05-settings.toml",
            '[preview]\ntitle = "合成数据"\nenabled = true\nports = [8080, 9000]\n',
            "application/toml",
            "TOML source",
        ),
        (
            "06-catalog.xml",
            '<?xml version="1.0" encoding="UTF-8"?>\n'
            '<catalog><item id="1">中文 &amp; café</item></catalog>\n',
            "application/xml",
            "XML source with escaped entities",
        ),
        (
            "07-example.py",
            '"""Synthetic source fixture."""\n\ndef greet(name: str) -> str:\n'
            '    return f"你好, {name}"\n\nprint(greet("Workbench"))\n',
            "text/x-python",
            "Python highlighting",
        ),
        (
            "08-example.ts",
            "type Entry = { id: number; label: string };\n"
            'export const entries: Entry[] = [{ id: 1, label: "中文" }];\n',
            "text/typescript",
            "TypeScript must be source, not Windows MPEG media",
        ),
        (
            "09-component.tsx",
            "export function Preview() {\n"
            '  return <section aria-label="合成内容"><h1>Hello</h1></section>;\n}\n',
            "text/typescript",
            "TSX source",
        ),
        (
            "10-style.css",
            ":root { --accent: #d8662b; }\n"
            ".fixture { display: grid; gap: 16px; color: var(--accent); }\n",
            "text/css",
            "CSS highlighting",
        ),
    ]:
        g.text(name, content, "source", mime, scenario)
    rows = [
        ["id", "name", "notes", "literal"],
        ["1", "中文 / café", "comma, quoted", "=1+1"],
        ["2", 'He said "hello"', "first line\nsecond line", "+SUM(1,2)"],
        ["3", "emoji 🌱", "embedded\ttab", "@SUM(A1:A3)"],
        ["4", "empty cell", "", "-42"],
    ]
    for name, delimiter, mime in [
        ("11-table.csv", ",", "text/csv"),
        ("12-table.tsv", "\t", "text/tab-separated-values"),
    ]:
        stream = io.StringIO(newline="")
        csv.writer(stream, delimiter=delimiter).writerows(rows)
        g.text(
            name,
            stream.getvalue(),
            "csv",
            mime,
            "Quoted multiline cells, Unicode, delimiters; "
            "formula-looking cells remain literal text.",
        )
    stream = io.StringIO(newline="")
    writer = csv.writer(stream)
    writer.writerow(["row", "label", "amount", "category"])
    writer.writerows((n, f"row {n}", n / 10, "synthetic") for n in range(1, 10_501))
    g.text(
        "13-large-table.csv",
        stream.getvalue(),
        "csv",
        "text/csv",
        "Bounded CSV worker preview.",
        limits="10,500 data rows; preview must show truncation after 10,000 rows.",
    )
    g.text("14-empty.txt", "", "source", "text/plain", "Valid empty text; no loading loop.")
    g.text(
        "15-long-line.txt",
        "x" * (256 * 1024 + 1) + "\n",
        "unavailable",
        "text/plain",
        "Explicit source line limit feedback.",
        limits="Single UTF-8 line exceeds the backend 256 KiB line limit.",
    )
    g.text(
        "16-preview.html",
        '<!doctype html><html><head><meta charset="utf-8"><style>'
        "body{font:18px system-ui;padding:24px;color:#303030}table{border-collapse:collapse}"
        "td,th{padding:8px;border:1px solid #ddd}</style></head>"
        "<body><h1>合成 HTML preview</h1><p>Offline styled document.</p>"
        "<table><tr><th>Item</th><th>Count</th></tr>"
        "<tr><td>Alpha</td><td>12</td></tr></table></body></html>",
        "html",
        "text/html",
        "Offline HTML layout, Unicode and table.",
    )
    make_html_probe(g)


def make_html_probe(g: Gallery) -> None:
    g.text(
        "17-isolation-probe.html",
        '<!doctype html><html><head><meta charset="utf-8"></head><body>'
        '<h1>HTML isolation probe</h1><p id="script">Waiting for local JavaScript.</p>'
        '<p id="parent">Waiting for parent access check.</p>'
        '<p id="network">Waiting for network access check.</p><script>'
        'document.getElementById("script").textContent="PASS: local JavaScript works";'
        'try{void parent.document.body;document.getElementById("parent").textContent='
        '"FAIL: parent document accessible"}catch(error){'
        'document.getElementById("parent").textContent="PASS: parent access blocked"}'
        'fetch("https://example.invalid/synthetic-preview-probe").then(()=>{'
        'document.getElementById("network").textContent="FAIL: network request succeeded"'
        '}).catch(()=>{document.getElementById("network").textContent='
        '"Network request blocked or failed; inspect CSP for policy evidence"})'
        "</script></body></html>",
        "html",
        "text/html",
        "SECURITY PROBE: local JS may execute in an opaque sandbox; parent access "
        "and network must be blocked. The network failure alone does not prove CSP.",
    )


def make_images(g: Gallery) -> None:
    if not g.dependency("PIL"):
        return
    from PIL import Image, ImageDraw

    image = Image.new("RGB", (960, 600), "#f4f4f4")
    draw = ImageDraw.Draw(image)
    draw.text((36, 24), "SYNTHETIC WORKBENCH PREVIEW", fill="#252525")
    colors = ["#dc682b", "#246fbd", "#31845b", "#8561b4"]
    for i, color in enumerate(colors):
        x = 60 + 220 * i
        draw.rounded_rectangle((x, 100, x + 175, 500), radius=20, fill=color)
        draw.text((x + 20, 530), f"Sample {i + 1}", fill="#252525")
    for name, fmt, mime in [
        ("18-chart.png", "PNG", "image/png"),
        ("19-chart.jpg", "JPEG", "image/jpeg"),
        ("20-chart.webp", "WEBP", "image/webp"),
    ]:
        image.save(g.directory / name, format=fmt)
        g.add(name, "image", mime, "Synthetic color chart, aspect ratio and scaling.")
    frames = []
    for i in range(8):
        frame = Image.new("RGB", (480, 240), "#f4f4f4")
        painter = ImageDraw.Draw(frame)
        painter.ellipse((30 + i * 45, 80, 100 + i * 45, 150), fill=colors[i % 4])
        painter.text((20, 20), f"Animated fixture: frame {i + 1}", fill="#303030")
        frames.append(frame)
    frames[0].save(
        g.directory / "21-animation.gif",
        save_all=True,
        append_images=frames[1:],
        duration=160,
        loop=0,
    )
    g.add("21-animation.gif", "image", "image/gif", "Animated GIF playback.")
    transparent = Image.new("RGBA", (480, 320), (0, 0, 0, 0))
    ImageDraw.Draw(transparent).ellipse((80, 30, 400, 290), fill=(220, 104, 43, 155))
    transparent.save(g.directory / "22-transparent.png")
    g.add("22-transparent.png", "image", "image/png", "Alpha transparency is preserved.")
    g.text(
        "23-vector.svg",
        '<svg xmlns="http://www.w3.org/2000/svg" width="800" height="400" '
        'viewBox="0 0 800 400"><rect width="800" height="400" fill="#fafafa"/>'
        '<circle cx="180" cy="200" r="100" fill="#dc682b"/>'
        '<text x="320" y="210" font-size="28" fill="#333">SVG / 中文 / café</text></svg>',
        "image",
        "image/svg+xml",
        "SVG rendered through an isolated image Blob, not inline DOM.",
    )
    g.text(
        "24-isolation-probe.svg",
        '<svg xmlns="http://www.w3.org/2000/svg" width="800" height="300">'
        '<rect width="800" height="300" fill="#fafafa"/>'
        '<text id="result" x="25" y="80" font-size="24">SVG scripts must not execute</text>'
        '<script>document.getElementById("result").textContent="FAIL: script executed";'
        'document.documentElement.setAttribute("data-script-executed", "true")</script>'
        '<foreignObject x="25" y="120" width="700" height="80">'
        '<div xmlns="http://www.w3.org/1999/xhtml">Isolation fixture; no network or secrets.'
        "</div></foreignObject></svg>",
        "image",
        "image/svg+xml",
        "SECURITY PROBE: img-context SVG must not execute script or access parent DOM.",
    )


def make_pdf(g: Gallery) -> None:
    if not g.dependency("reportlab"):
        return
    from reportlab.lib.colors import HexColor
    from reportlab.pdfgen import canvas

    path = g.directory / "25-multipage.pdf"
    pdf = canvas.Canvas(str(path), pagesize=(595, 842), invariant=1)
    pdf.setTitle("Synthetic multi-page preview fixture")
    for page in range(1, 4):
        pdf.setFont("Helvetica-Bold", 22)
        pdf.drawString(48, 770, f"Synthetic report - page {page} / 3")
        pdf.setFont("Helvetica", 12)
        pdf.drawString(
            48, 735, "All data is synthetic. Test page navigation, zoom and text selection."
        )
        for row in range(12):
            y = 680 - row * 37
            pdf.setFillColor(HexColor("#eeeeee" if row % 2 == 0 else "#ffffff"))
            pdf.rect(48, y - 8, 495, 32, stroke=0, fill=1)
            pdf.setFillColor(HexColor("#303030"))
            pdf.drawString(60, y, f"Record {page * 100 + row}: value {row * 7} units")
        pdf.setFillColor(HexColor("#dc682b"))
        pdf.rect(48, 100, 150 * page, 40, stroke=0, fill=1)
        pdf.showPage()
    pdf.save()
    g.add(
        path.name,
        "pdf",
        "application/pdf",
        "Three pages with selectable text, zebra table and vector graphic.",
        limits="Offline PDF.js current-page rendering with pagination; "
        "visual fidelity is verified separately. Download remains available.",
    )


def make_office(g: Gallery) -> None:
    if g.dependency("docx"):
        from docx import Document
        from docx.shared import Inches

        document = Document()
        document.add_heading("Synthetic preview report / 合成报告", 0)
        document.add_paragraph("This document contains no real user data. 中文、日本語、café.")
        document.add_heading("Table and styles", 1)
        table = document.add_table(rows=1, cols=3)
        table.style = "Light Shading Accent 1"
        for cell, text in zip(table.rows[0].cells, ["Item", "Quantity", "Comment"], strict=True):
            cell.text = text
        for n in range(1, 9):
            for cell, text in zip(
                table.add_row().cells, [f"Sample {n}", str(n * 5), "Unicode 中文 ✓"], strict=True
            ):
                cell.text = text
        document.add_page_break()
        document.add_heading("Second page: image and lists", 1)
        for text in ["First synthetic item", "Second item with café", "Third item with 中文"]:
            document.add_paragraph(text, style="List Bullet")
        chart = g.directory / "18-chart.png"
        if chart.exists():
            document.add_picture(str(chart), width=Inches(5.5))
        document.sections[0].header.paragraphs[0].text = "Workbench synthetic fixture"
        document.sections[0].footer.paragraphs[
            0
        ].text = "Preview fidelity test - not a user document"
        document.save(g.directory / "26-report.docx")
        g.add(
            "26-report.docx",
            "office",
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            "Two-page DOCX: styled table, headings, Unicode, list, image, header/footer.",
            limits="Preview fidelity may differ from Microsoft Word; no editing or writeback here.",
        )
    if g.dependency("pptx"):
        from pptx import Presentation
        from pptx.dml.color import RGBColor
        from pptx.util import Inches, Pt

        presentation = Presentation()
        presentation.slide_width = Inches(12)
        presentation.slide_height = Inches(6.75)
        slide = presentation.slides.add_slide(presentation.slide_layouts[0])
        slide.shapes.title.text = "Synthetic slide deck / 合成演示"
        slide.placeholders[1].text = "Four slides: text, table, image, shapes"
        slide = presentation.slides.add_slide(presentation.slide_layouts[5])
        slide.shapes.title.text = "Table and Unicode"
        table = slide.shapes.add_table(
            5, 3, Inches(0.7), Inches(1.7), Inches(10.6), Inches(3.5)
        ).table
        for row in range(5):
            for col in range(3):
                table.cell(row, col).text = (
                    ["Item", "Quantity", "Notes"][col]
                    if row == 0
                    else [f"Sample {row}", str(row * 8), "中文 / café"][col]
                )
        slide = presentation.slides.add_slide(presentation.slide_layouts[5])
        slide.shapes.title.text = "Image scaling"
        chart = g.directory / "18-chart.png"
        if chart.exists():
            slide.shapes.add_picture(str(chart), Inches(2.5), Inches(1.4), width=Inches(7))
        slide = presentation.slides.add_slide(presentation.slide_layouts[5])
        slide.shapes.title.text = "Shapes and typography"
        for n, rgb in enumerate([(220, 104, 43), (36, 111, 189), (49, 132, 91)]):
            shape = slide.shapes.add_shape(
                5, Inches(0.7 + n * 3.8), Inches(2), Inches(3.2), Inches(2.5)
            )
            shape.fill.solid()
            shape.fill.fore_color.rgb = RGBColor(*rgb)
            shape.text = f"Stage {n + 1}"
            shape.text_frame.paragraphs[0].font.size = Pt(28)
        presentation.save(g.directory / "27-deck.pptx")
        g.add(
            "27-deck.pptx",
            "office",
            "application/vnd.openxmlformats-officedocument.presentationml.presentation",
            "Four 16:9 slides: title, table, embedded image and colored shapes.",
            limits="Slide navigation must expose all four slides; animation is not required.",
        )
    if g.dependency("openpyxl"):
        from openpyxl import Workbook
        from openpyxl.chart import BarChart, Reference
        from openpyxl.styles import Font, PatternFill

        workbook = Workbook()
        summary = workbook.active
        summary.title = "Summary"
        summary.append(["Category", "Q1", "Q2", "Total formula"])
        for n, label in enumerate(["Alpha", "Beta", "Gamma", "中文"], 2):
            summary.append([label, n * 10, n * 13, f"=SUM(B{n}:C{n})"])
        summary.freeze_panes = "B2"
        summary.auto_filter.ref = "A1:D5"
        for cell in summary[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="DC682B")
        for column in "ABCD":
            summary.column_dimensions[column].width = 23
        chart = BarChart()
        chart.title = "Synthetic comparison"
        chart.add_data(
            Reference(summary, min_col=2, max_col=3, min_row=1, max_row=5), titles_from_data=True
        )
        chart.set_categories(Reference(summary, min_col=1, min_row=2, max_row=5))
        summary.add_chart(chart, "F2")
        details = workbook.create_sheet("Details")
        details.append(["ID", "Date", "Amount", "Notes"])
        for n in range(1, 81):
            details.append([n, datetime(2026, 1, (n % 28) + 1), n * 1.25, "Synthetic café"])
            details.cell(n + 1, 2).number_format = "yyyy-mm-dd"
            details.cell(n + 1, 3).number_format = "#,##0.00"
        details.freeze_panes = "A2"
        notes = workbook.create_sheet("中文说明")
        notes.merge_cells("A1:D1")
        notes["A1"] = "合成多工作表测试 — Synthetic fixture"
        notes["A3"] = "Cells beginning with = below are intentional spreadsheet formulas."
        notes["A4"] = "Formula caching and recalculation depend on the preview library."
        workbook.save(g.directory / "28-workbook.xlsx")
        g.add(
            "28-workbook.xlsx",
            "office",
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            "Three sheets: styled table, formulas, chart, dates, freeze panes "
            "and Unicode sheet name.",
            limits=(
                "Formula cached results are absent in generated openpyxl files; "
                "do not require Excel calculation parity."
            ),
        )


def make_media(g: Gallery, ffmpeg: str | None) -> None:
    path = g.directory / "29-tone.wav"
    sample_rate = 22050
    with wave.open(str(path), "wb") as audio:
        audio.setnchannels(1)
        audio.setsampwidth(2)
        audio.setframerate(sample_rate)
        # Low-amplitude three-second tone with fade prevents loud clicks.
        frames = bytearray()
        for index in range(sample_rate * 3):
            fade = min(1, index / 1000, (sample_rate * 3 - index) / 1000)
            frames.extend(
                struct.pack(
                    "<h", int(2000 * fade * math.sin(2 * math.pi * 440 * index / sample_rate))
                )
            )
        audio.writeframes(frames)
    g.add(
        path.name,
        "audio",
        "audio/wav",
        "Three-second low-amplitude synthetic tone; manual play, pause and seek.",
    )
    if not ffmpeg:
        g.skipped.append(
            {
                "producer": "ffmpeg",
                "reason": "Not installed; compressed audio and video not generated",
            }
        )
        return
    jobs = [
        (
            "30-tone.mp3",
            "audio",
            "audio/mpeg",
            ["-i", str(path), "-codec:a", "libmp3lame", "-q:a", "6"],
        ),
        ("31-tone.flac", "audio", "audio/flac", ["-i", str(path), "-codec:a", "flac"]),
        ("32-tone.aac", "audio", "audio/aac", ["-i", str(path), "-codec:a", "aac", "-b:a", "64k"]),
        (
            "33-motion.mp4",
            "video",
            "video/mp4",
            [
                "-f",
                "lavfi",
                "-i",
                "testsrc2=size=640x360:rate=24",
                "-t",
                "3",
                "-an",
                "-codec:v",
                "libx264",
                "-preset",
                "ultrafast",
                "-pix_fmt",
                "yuv420p",
                "-movflags",
                "+faststart",
            ],
        ),
        (
            "34-motion.webm",
            "video",
            "video/webm",
            [
                "-f",
                "lavfi",
                "-i",
                "testsrc2=size=640x360:rate=24",
                "-t",
                "3",
                "-an",
                "-codec:v",
                "libvpx-vp9",
                "-deadline",
                "realtime",
                "-cpu-used",
                "8",
                "-b:v",
                "350k",
            ],
        ),
    ]
    for name, preview, mime, args in jobs:
        output = g.directory / name
        try:
            subprocess.run(
                [
                    ffmpeg,
                    "-hide_banner",
                    "-loglevel",
                    "error",
                    "-nostdin",
                    "-n",
                    *args,
                    str(output),
                ],
                check=True,
                capture_output=True,
                timeout=45,
            )
        except (OSError, subprocess.SubprocessError) as exc:
            output.unlink(missing_ok=True)
            g.skipped.append(
                {
                    "producer": name,
                    "reason": f"FFmpeg unavailable codec or error: {type(exc).__name__}",
                }
            )
            continue
        g.add(
            name,
            preview,
            mime,
            "Short synthetic media; no autoplay, remote assets or user recordings.",
            limits=(
                "Codec playback depends on browser; show useful media error "
                "and download fallback if unsupported."
            ),
        )


def make_downloads(g: Gallery) -> None:
    with zipfile.ZipFile(g.directory / "35-archive.zip", "w", zipfile.ZIP_DEFLATED) as archive:
        info = zipfile.ZipInfo("README.txt", (2026, 1, 1, 0, 0, 0))
        archive.writestr(info, "Synthetic archive fixture; preview should offer download.\n")
    g.add(
        "35-archive.zip",
        "download",
        "application/zip",
        "Unsupported preview format; download/open actions remain available.",
    )
    (g.directory / "36-opaque.bin").write_bytes(b"\x00\xffSYNTHETIC\x00" + bytes(range(256)))
    g.add(
        "36-opaque.bin",
        "download",
        "application/octet-stream",
        "Binary content must not be decoded as text.",
    )
    (g.directory / "37-empty.bin").write_bytes(b"")
    g.add(
        "37-empty.bin",
        "download",
        "application/octet-stream",
        "Zero-byte unknown binary; metadata/download should still work.",
    )
    make_invalid_payloads(g)


def make_invalid_payloads(g: Gallery) -> None:
    for name, preview, mime in [
        ("38-invalid-image.png", "image-error", "image/png"),
        ("39-invalid-document.pdf", "pdf-error", "application/pdf"),
        ("40-invalid-audio.mp3", "audio-error", "audio/mpeg"),
        ("41-invalid-video.mp4", "video-error", "video/mp4"),
    ]:
        (g.directory / name).write_bytes(b"Synthetic invalid payload: not the advertised format.\n")
        g.add(
            name,
            preview,
            mime,
            "NEGATIVE FIXTURE: display a useful preview error "
            "and retain download; never remain stuck loading.",
        )


def generate(output: Path, ffmpeg: str | None = None) -> dict[str, Any]:
    # No overwrite option: a new directory is cheap and cannot erase user data.
    output.mkdir(parents=True, exist_ok=False)
    gallery = Gallery(output)
    make_text(gallery)
    make_images(gallery)
    make_pdf(gallery)
    make_office(gallery)
    make_media(gallery, ffmpeg)
    make_downloads(gallery)
    manifest = {
        "schemaVersion": 1,
        "generatedAt": datetime.now(UTC).isoformat(),
        "syntheticOnly": True,
        "generator": "scripts/testing/create_workbench_preview_fixtures.py",
        "checks": [
            "Select each file through Workspace files; resolve pins its revision.",
            "Expected MIME should match resolve, listing and conditional content response.",
            "Refresh a selected file after modification; stale revision must reject the old read.",
            "SVG image-context scripts must not execute. HTML local JS is allowed "
            "in an opaque sandbox; parent document and network access must be blocked.",
            "No media autoplay; image/Office aspect ratios and navigation must remain usable.",
            "These fixtures verify integration, not full Office rendering fidelity.",
        ],
        "files": gallery.entries,
        "skipped": gallery.skipped,
    }
    (output / "manifest.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    return manifest


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--output",
        type=Path,
        required=True,
        help="New output directory; existing directories are rejected",
    )
    parser.add_argument(
        "--ffmpeg",
        default=shutil.which("ffmpeg"),
        help="Optional FFmpeg executable; detected from PATH by default",
    )
    args = parser.parse_args()
    manifest = generate(args.output.resolve(), args.ffmpeg)
    print(
        json.dumps(
            {
                "output": str(args.output.resolve()),
                "files": len(manifest["files"]),
                "skipped": manifest["skipped"],
            },
            ensure_ascii=False,
        )
    )


if __name__ == "__main__":
    main()
