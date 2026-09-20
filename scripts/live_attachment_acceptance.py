#!/usr/bin/env python3
"""Offline fixture preparation and bounded real Gateway attachment acceptance.

The browser performs ordinary uploads, sends, edits, and queue/restart actions.
This harness never seeds session rows or supplies model responses. One external
functional relay owns all physical calls across both source variants/restarts.
"""
from __future__ import annotations

import argparse
import asyncio
import contextlib
import hashlib
import io
import json
import logging
import os
import platform
import re
import secrets
import sqlite3
import subprocess
import sys
from pathlib import Path
from typing import Any
from unittest.mock import patch

REPO_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(REPO_ROOT / "src"))
sys.path.insert(0, str(REPO_ROOT))

from scripts.live_harness_security import (  # noqa: E402
    minimal_child_environment,
    registry_endpoint,
    require_temporary_report_path,
    restrict_private_file_permissions,
    write_safe_report,
)
from scripts.live_tokenrhythm_budget import (  # noqa: E402
    ATTACHMENT_PHASE_CALL_LIMITS,
    FunctionalRequestLog,
)

TEXT_MODELS = ("deepseek-v4-flash-0731", "deepseek-v4-pro-0813")
VISION_MODEL = "kimi-k2.6"
FILE_TOOLS = frozenset({"read_file", "read_spreadsheet", "edit_file", "write_file", "pdf", "image"})


def digest(value: bytes | str) -> str:
    return hashlib.sha256(value.encode() if isinstance(value, str) else value).hexdigest()


def source_evidence() -> dict[str, Any]:
    def git(*arguments: str) -> str:
        return subprocess.check_output(["git", *arguments], cwd=REPO_ROOT, text=True).strip()
    untracked = {
        name: digest((REPO_ROOT / name).read_bytes())
        for name in git("ls-files", "--others", "--exclude-standard", "--", "src", "scripts")
        .splitlines() if (REPO_ROOT / name).is_file()
    }
    return {"commit": git("rev-parse", "HEAD"),
            "tracked_diff_sha256": digest(git("diff", "HEAD", "--", "src", "scripts")),
            "untracked_source_count": len(untracked),
            "untracked_source_sha256": digest(json.dumps(untracked, sort_keys=True)),
            "harness_sha256": digest(Path(__file__).read_bytes()),
            "platform": platform.system(), "python": platform.python_version(),
            "session_seeded": False, "physical_call_limit": 60}


def tool_outcome(result: Any) -> dict[str, Any]:
    """Classify real handler results without retaining messages or file contents."""
    try:
        payload = json.loads(result) if isinstance(result, str) else result
    except (ValueError, TypeError):
        return {"ok": True}
    if not isinstance(payload, dict):
        return {"ok": True}
    status = payload.get("status")
    status = status if isinstance(status, str) and re.fullmatch(r"[a-z_]{1,80}", status) else None
    rejected = status in {
        "blocked", "denied", "error", "elevation_required", "path_access_required",
        "approval_required", "approval_denied", "approval_pending",
    }
    return {"ok": not (rejected or payload.get("error") or payload.get("ok") is False),
            **({"status": status} if status else {})}


def prepare_fixtures(root: Path, manifest_path: Path) -> dict[str, Any]:
    """Produce only synthetic sources; keep the answer oracle outside all tool roots."""
    from docx import Document
    from openpyxl import Workbook
    from PIL import Image, ImageDraw, ImageFont
    from pptx import Presentation
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfgen.canvas import Canvas

    root = require_temporary_report_path(root / "fixture-marker.json").parent
    manifest_path = require_temporary_report_path(manifest_path)
    if manifest_path.is_relative_to(root):
        raise ValueError("oracle_must_be_outside_fixture_root")
    if root.exists() and any(root.iterdir()):
        raise ValueError("fixture_root_must_be_empty")
    root.mkdir(parents=True, exist_ok=True, mode=0o700)
    cases: list[dict[str, Any]] = []

    def code(kind: str, page: int = 1) -> str:
        return "AMBER-" + digest(f"synthetic-{kind}-{page}")[:8].upper()

    def add(
        name: str, mime: str, answers: list[str], *, vision: bool = False,
        pages: int | None = None, textless_pages: int | None = None,
    ) -> None:
        path = root / name
        restrict_private_file_permissions(path)
        cases.append({"id": path.name, "name": name, "mime": mime,
                      "sha256": digest(path.read_bytes()), "size": path.stat().st_size,
                      "answers": answers, "requires_vision": vision,
                      **({"pages": pages, "textless_pages": textless_pages} if pages else {}),
                      "prompt": "Read the attached file completely. Return every RECORD_CODE "
                      "in page or row order, and report missing or unreadable pages. "
                      "Use the file tools if its contents are not already available. "
                      "For image-only PDF pages use pdf with render=true on the required pages. "
                      "Do not infer a code from the filename."})

    text = "STATUS=DRAFT\nRECORD_CODE=" + code("text") + "\n"
    (root / "record.txt").write_text(text)
    add("record.txt", "text/plain", [code("text")])
    doc = Document()
    doc.add_paragraph("Synthetic attachment record")
    doc.add_paragraph("RECORD_CODE=" + code("docx"))
    doc.save(root / "record.docx")
    add("record.docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        [code("docx")])
    slides = Presentation()
    for index in range(1, 3):
        slide = slides.slides.add_slide(slides.slide_layouts[1])
        slide.shapes.title.text = f"Synthetic record {index}"
        slide.placeholders[1].text = "RECORD_CODE=" + code("pptx", index)
    slides.save(root / "record.pptx")
    add("record.pptx", "application/vnd.openxmlformats-officedocument.presentationml.presentation",
        [code("pptx", index) for index in range(1, 3)])
    book = Workbook()
    sheet = book.active
    sheet.title = "Synthetic records"
    sheet.append(["record", "RECORD_CODE"])
    for index in range(1, 4):
        sheet.append([index, code("xlsx", index)])
    book.save(root / "record.xlsx")
    add("record.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        [code("xlsx", index) for index in range(1, 4)])
    (root / "record.eml").write_text(
        "From: author@example.invalid\nTo: reader@example.invalid\n"
        "Date: Tue, 1 Jan 2030 00:00:00 +0000\nMIME-Version: 1.0\n"
        "Subject: Synthetic record\nContent-Type: text/plain; charset=utf-8\n\n"
        "RECORD_CODE=" + code("eml") + "\n"
    )
    add("record.eml", "message/rfc822", [code("eml")])

    def raster(label: str):
        image = Image.new("RGB", (1400, 800), "white")
        draw = ImageDraw.Draw(image)
        draw.text((60, 180), "RECORD_CODE=" + label, fill="black",
                  font=ImageFont.load_default(size=54))
        return image

    raster(code("image")).save(root / "record.png")
    add("record.png", "image/png", [code("image")], vision=True)
    for mode in ("text", "scan", "mixed"):
        name = f"five-pages-{mode}.pdf"
        canvas = Canvas(str(root / name), pagesize=(700, 400))
        answers = [code("pdf-" + mode, page) for page in range(1, 6)]
        for page, answer in enumerate(answers, 1):
            if mode == "scan" or (mode == "mixed" and page % 2 == 0):
                canvas.drawImage(ImageReader(raster(answer)), 0, 0, width=700, height=400)
            else:
                canvas.setFont("Helvetica", 22)
                canvas.drawString(30, 240, f"Page {page}: RECORD_CODE={answer}")
            canvas.showPage()
        canvas.save()
        add(name, "application/pdf", answers, vision=mode != "text", pages=5,
            textless_pages=5 if mode == "scan" else 2 if mode == "mixed" else 0)
    for index in range(1, 6):
        name = f"receipt-{index}.pdf"
        answer = code("pdf-receipt", index)
        canvas = Canvas(str(root / name), pagesize=(700, 400))
        canvas.setFont("Helvetica", 22)
        canvas.drawString(30, 240, "RECORD_CODE=" + answer)
        canvas.showPage()
        canvas.save()
        add(name, "application/pdf", [answer], pages=1, textless_pages=0)
    pressure = []
    for index in range(1, 4):
        path = root / f"pressure-{index:03d}.txt"
        lines = [f"SYNTHETIC_PRESSURE_{index:03d}_{line:04d}: " +
                 "Disposable background about colored paper, with no additional task. " * 3
                 for line in range(240)]
        path.write_text("\n".join(lines) + "\n")
        restrict_private_file_permissions(path)
        pressure.append({"name": path.name, "sha256": digest(path.read_bytes()),
                         "prompt": f"Read {path.name} with read_file, then reply BACKGROUND_READ. "
                         "Keep the earlier record's exact code and edit status for continuation."})
    manifest = {"schema_version": 1, "synthetic_only": True, "fixtures": cases,
                "pressure_files": pressure,
                "edit_prompt": "Change STATUS=DRAFT to STATUS=FINAL in the available working file. "
                "Preserve its RECORD_CODE and report the file path and that exact code.",
                "consumption_batches": {
                    "ten_document_uploads": ["record.txt", "record.docx", "record.pptx",
                                             "record.xlsx", "record.eml",
                                             *[f"receipt-{index}.pdf" for index in range(1, 6)]],
                    "vision_and_pagination": ["record.png", "five-pages-text.pdf",
                                              "five-pages-scan.pdf", "five-pages-mixed.pdf"],
                },
                "phase_limits": ATTACHMENT_PHASE_CALL_LIMITS,
                "models": {"text_candidates": TEXT_MODELS, "vision_candidate": VISION_MODEL},
                "coverage_rules": {
                    "native_window": "uncovered until reached by actual sends/tools",
                    "configured_pressure": "explicit smaller application window only",
                    "restart": "reuse the same Gateway root and shared relay",
                    "original_bytes": "compare source hashes before and after edits",
                }}
    write_safe_report(manifest_path, manifest, ())
    return manifest


def preflight(root: Path, manifest_path: Path) -> dict[str, Any]:
    """Verify fixtures and parser expectations without constructing a provider."""
    import pdfplumber

    from opensquilla.contracts.attachment_sniff import sniff_mime_from_bytes
    from opensquilla.tools.document_readers import read_document

    manifest = json.loads(require_temporary_report_path(manifest_path).read_text())
    checks: list[dict[str, Any]] = []
    for item in manifest["fixtures"]:
        path = root / item["name"]
        unchanged = digest(path.read_bytes()) == item["sha256"]
        mime_matches = sniff_mime_from_bytes(path.read_bytes()) == item["mime"]
        parser_ok: bool | None = None
        pages = textless = None
        if path.suffix == ".pdf":
            with pdfplumber.open(path) as document:
                texts = [page.extract_text() or "" for page in document.pages]
            pages, textless = len(texts), sum(not text.strip() for text in texts)
            parser_ok = pages == item["pages"] and textless == item["textless_pages"]
            if textless == 0:
                parser_ok = parser_ok and all(answer in "\n".join(texts)
                                              for answer in item["answers"])
        elif path.suffix not in {".png", ".txt"}:
            rendered = json.dumps(read_document(path, offset=1, limit=10), ensure_ascii=False)
            parser_ok = all(answer in rendered for answer in item["answers"])
        checks.append({"case": item["id"], "hash_matches": unchanged,
                       "mime_matches": mime_matches,
                       "parser_check": parser_ok, "pages": pages, "textless_pages": textless})
    return {"ok": all(row["hash_matches"] and row["mime_matches"]
                       and row["parser_check"] is not False for row in checks),
            "physical_calls": 0, "checks": checks, **source_evidence()}


def evaluate(root: Path, fixtures: Path, manifest_path: Path) -> dict[str, Any]:
    """Compare real persisted answers against hidden fixture-only facts, without exporting text."""
    from scripts.live_compaction_gateway import storage_evidence

    manifest = json.loads(require_temporary_report_path(manifest_path).read_text())
    database = root / "state" / "sessions.db"
    answers: list[str] = []
    with contextlib.closing(sqlite3.connect(database.as_uri() + "?mode=ro", uri=True)) as db:
        db.execute("PRAGMA query_only=ON")
        for table in ("transcript_entries", "compacted_transcript_entries"):
            answers.extend(str(row[0] or "") for row in db.execute(
                f"SELECT content FROM {table} WHERE role='assistant'"
            ))
    cases = [{"case": item["id"],
              "answers_present": [any(expected in answer for answer in answers)
                                  for expected in item["answers"]],
              "complete_answer_present": any(
                  all(expected in answer for expected in item["answers"]) for answer in answers
              ),
              "source_unchanged": digest((fixtures / item["name"]).read_bytes()) == item["sha256"]}
             for item in manifest["fixtures"]]
    storage, _ = storage_evidence(database)
    return {"status": "requires_case_and_browser_review", "cases": cases,
            "storage": storage, "working_files": working_file_evidence(root), **source_evidence()}


def working_file_evidence(root: Path) -> list[dict[str, Any]]:
    """Hash durable source/copy bindings without exporting paths or file contents."""
    database = root / "state" / "sessions.db"
    if not database.is_file():
        return []
    configured_workspace = (root / "workspace").resolve()
    managed_workspaces = (root / "state" / "tasks").resolve()
    result: list[dict[str, Any]] = []
    with contextlib.closing(sqlite3.connect(database.as_uri() + "?mode=ro", uri=True)) as db:
        db.execute("PRAGMA query_only=ON")
        for session_id, raw, project, execution in db.execute(
            "SELECT s.session_id,s.origin,w.path,s.execution_workspace FROM sessions s "
            "LEFT JOIN project_workspaces w ON s.workspace_id=w.workspace_id"
        ):
            execution = json.loads(execution) if execution else {}
            workspace = Path(project or execution.get("root") or configured_workspace).resolve()
            origin = json.loads(raw) if raw else {}
            for source, binding in origin.get("attachment_working_files", {}).items():
                original = (workspace / source).resolve()
                working = (workspace / str(binding.get("path", ""))).resolve()
                bounded = ((workspace.is_relative_to(configured_workspace)
                            or workspace.is_relative_to(managed_workspaces))
                           and original.is_relative_to(workspace)
                           and working.is_relative_to(workspace))
                before = digest(original.read_bytes()) if bounded and original.is_file() else None
                after = digest(working.read_bytes()) if bounded and working.is_file() else None
                result.append({"source_path_sha256": digest(source),
                               "working_path_sha256": digest(str(binding.get("path", ""))),
                               "session_matches": session_id == binding.get("session_id"),
                               "has_working_copy": bool(binding.get("path")),
                               "source_unchanged": before is not None
                               and before == binding.get("sha256"),
                               "source_sha256": before, "working_sha256": after,
                               "copy_differs": before is not None and after is not None
                               and before != after})
    return result


def open_relay(ready_path: Path) -> tuple[dict[str, Any], FunctionalRequestLog]:
    from scripts.live_tokenrhythm_transport import RelayTarget

    ready = json.loads(require_temporary_report_path(ready_path).read_text())
    if ready.get("mode") != "functional" or ready.get("enabled") is not True:
        raise ValueError("functional_relay_required")
    RelayTarget(str(ready["base_url"]), str(ready["client_key"]))
    log = FunctionalRequestLog(require_temporary_report_path(Path(ready["request_log"])),
                               enabled=True, max_calls=60)
    if log.phase_limits != ATTACHMENT_PHASE_CALL_LIMITS:
        raise ValueError("attachment_phase_allocations_required")
    return ready, log


async def catalog_evidence(placeholder: str, model: str) -> dict[str, Any]:
    import httpx

    from opensquilla.provider.tokenrhythm_catalog import parse_tokenrhythm_published

    async with httpx.AsyncClient(trust_env=False, follow_redirects=False, timeout=30) as client:
        authorized = await client.get(
            registry_endpoint("tokenrhythm") + "/models",
            headers={"Authorization": "Bearer " + placeholder},
        )
        public = await client.get("https://tokenrhythm.studio/api/models")
    if authorized.status_code != 200 or public.status_code != 200:
        raise ValueError("model_catalog_unavailable")
    rows = authorized.json().get("data", [])
    available = {row.get("id") for row in rows if isinstance(row, dict)}
    published = parse_tokenrhythm_published(public.json())
    entry = published.get(model)
    if model not in available or entry is None:
        raise ValueError("selected_model_not_in_both_catalogs")
    vision = bool(entry.capabilities.vision and "image" in (entry.modalities or ()))
    if model == VISION_MODEL and not vision:
        raise ValueError("vision_model_capability_not_verified")
    return {"model": model, "authorized": True, "published": True, "vision": vision,
            "context_window_tokens": entry.context_window,
            "max_output_tokens": entry.max_output_tokens,
            "authorized_catalog_sha256": digest(authorized.content),
            "public_catalog_sha256": digest(public.content)}


def prepare_desktop_owner(config: Any, root: Path, ready_path: Path) -> dict[str, Any]:
    """Prepare private local launch authority; never include it in public evidence."""
    from opensquilla.gateway.desktop_ownership import desktop_gateway_auth_token
    from opensquilla.recovery.locking import profile_lock_key

    ready_path = require_temporary_report_path(ready_path).resolve()
    root = root.resolve()
    if ready_path.is_relative_to(root):
        raise ValueError("desktop_owner_handoff_must_be_outside_gateway_root")
    profile = root / "desktop-profile"
    profile.mkdir(parents=True, exist_ok=True, mode=0o700)
    config.config_path = str(profile / "config.toml")
    fingerprint = profile_lock_key(profile)
    ownership = ready_path.parent / "desktop-ownership" / fingerprint
    nonce, instance = secrets.token_urlsafe(32), secrets.token_hex(16)
    os.environ.update({
        "OPENSQUILLA_DESKTOP": "1",
        "OPENSQUILLA_DESKTOP_GATEWAY_INSTANCE_NONCE": nonce,
        "OPENSQUILLA_DESKTOP_GATEWAY_INSTANCE_ID": instance,
        "OPENSQUILLA_DESKTOP_GATEWAY_OWNERSHIP_DIR": str(ownership),
    })
    return {
        "schemaVersion": 1, "instanceId": instance, "profileFingerprint": fingerprint,
        "httpUrl": f"http://127.0.0.1:{config.port}",
        "authToken": desktop_gateway_auth_token(nonce), "nonce": nonce,
        "ownershipDir": str(ownership),
    }


def write_private_desktop_handoff(path: Path, handoff: dict[str, Any]) -> None:
    path = require_temporary_report_path(path)
    path.parent.mkdir(parents=True, exist_ok=True, mode=0o700)
    descriptor = os.open(path, os.O_WRONLY | os.O_CREAT | os.O_EXCL, 0o600)
    try:
        restrict_private_file_permissions(path, descriptor=descriptor)
        with os.fdopen(descriptor, "w", encoding="utf-8") as stream:
            descriptor = -1
            json.dump(handoff, stream)
            stream.flush()
            os.fsync(stream.fileno())
    except BaseException:
        if descriptor >= 0:
            os.close(descriptor)
        path.unlink(missing_ok=True)
        raise


async def serve_gateway(args: Any, ready: dict[str, Any], log: FunctionalRequestLog) -> None:
    import opensquilla.engine.runtime
    from opensquilla.engine.cache_break_monitor import add_compaction_listener
    from opensquilla.gateway.boot import start_gateway_server
    from opensquilla.gateway.config import AuthConfig
    from opensquilla.tools.builtin import filesystem, media  # noqa: F401
    from opensquilla.tools.registry import ToolRegistry, get_default_registry
    from opensquilla.tools.types import current_tool_context
    from scripts.live_compaction_gateway import storage_evidence
    from scripts.live_reasoning_replay_e2e import (
        WireObserver,
        _config,
        _usage_report,
        _wire_diagnostics,
    )

    root = require_temporary_report_path(args.gateway_root / "marker.json").parent
    if not Path(opensquilla.engine.runtime.__file__).resolve().is_relative_to(
        REPO_ROOT / "src"
    ):
        raise ValueError("imported_source_root_mismatch")
    report_path = require_temporary_report_path(args.report)
    running_source = source_evidence()
    if report_path.is_relative_to(root):
        raise ValueError("report_must_be_outside_gateway_root")
    root.mkdir(parents=True, exist_ok=True, mode=0o700)
    catalog = await catalog_evidence(ready["client_key"], args.model)
    endpoint = registry_endpoint("tokenrhythm")
    config = _config(root, "tokenrhythm", args.model, endpoint, thinking="off")
    config.host, config.port = "127.0.0.1", args.gateway_port
    config.auth = AuthConfig(mode="none")
    config.llm.api_key_env = "TOKENRHYTHM_API_KEY"
    config.llm.context_window_tokens = args.context_window
    config.llm.max_tokens = args.max_output
    config.sandbox.run_mode = args.run_mode
    config.sandbox.sandbox = args.run_mode != "full"
    config.compaction.enabled = True
    config.log_file_enabled = False
    config.privacy.reliability_diagnostics_enabled = False
    config.privacy.product_analytics_enabled = False
    config.privacy.disable_network_observability = True
    config.tools.profile = "minimal"
    config.tools.allow = sorted(FILE_TOOLS)
    config.tools.deny = ["session_status"]
    desktop_handoff = (
        prepare_desktop_owner(config, root, args.desktop_owner_ready)
        if args.desktop_owner_ready else None
    )
    registry = ToolRegistry()
    tool_observations: list[dict[str, Any]] = []
    roots = (Path(config.workspace_dir).resolve(), Path(config.attachments.media_root).resolve(),
             (root / "state" / "tasks").resolve())
    for name in sorted(FILE_TOOLS):
        item = get_default_registry().get(name)
        if item is None:
            continue

        def guard(spec: Any, handler: Any):
            async def invoke(**arguments: Any):
                path = Path(str(arguments.get("path", "")))
                context = current_tool_context.get()
                workspace = Path(getattr(context, "workspace_dir", None) or roots[0])
                path = (path if path.is_absolute() else workspace / path).resolve()
                if not any(path.is_relative_to(allowed) for allowed in roots):
                    raise RuntimeError("acceptance_path_not_allowlisted")
                before = digest(path.read_bytes()) if path.is_file() else None
                observation = {"name": spec.name, "path_sha256": digest(str(path)),
                               "path_kind": "imported" if ".opensquilla" in path.parts
                               and "attachments" in path.parts
                               else "project", "before_sha256": before, "ok": False}
                tool_observations.append(observation)
                try:
                    result = await handler(**arguments)
                    observation.update(tool_outcome(result))
                except Exception as error:
                    observation["error_type"] = type(error).__name__
                    raise
                finally:
                    observation["after_sha256"] = (
                        digest(path.read_bytes()) if path.is_file() else None
                    )
                return result
            return invoke

        registry.register(item.spec, guard(item.spec, item.handler))
    observer = WireObserver(endpoint, endpoints={"tokenrhythm": endpoint}, max_calls=60)
    compaction_events: list[dict[str, Any]] = []

    def compacted(_key: str, payload: dict[str, Any]) -> None:
        allowed = {"status", "phase", "source", "reason", "tokens_before", "tokens_after",
                   "removed_count", "kept_count", "durability", "effect_status"}
        compaction_events.append({key: value for key, value in payload.items()
                                  if key in allowed and isinstance(value, (str, int, bool))})

    def publish(status: str) -> None:
        storage, _ = storage_evidence(root / "state" / "sessions.db")
        report = {"status": status, **running_source, "model": args.model,
                  "gateway_url": f"http://127.0.0.1:{args.gateway_port}",
                  "catalog": catalog,
                  "window_mode": "configured_pressure" if args.context_window else "native_auto",
                  "context_window_tokens": args.context_window, "run_mode": args.run_mode,
                  "native_window_coverage": "unverified", "storage": storage,
                  "desktop_native_owner_enabled": desktop_handoff is not None,
                  "working_files": working_file_evidence(root),
                  "tool_observations": tool_observations, "compaction_events": compaction_events,
                  "shared_relay": log.snapshot(), **_usage_report(observer.calls),
                  **_wire_diagnostics(observer)}
        private_values = (ready["client_key"],) + (
            (desktop_handoff["nonce"], desktop_handoff["authToken"])
            if desktop_handoff else ()
        )
        write_safe_report(report_path, report, private_values)

    with contextlib.ExitStack() as stack:
        stack.enter_context(observer.observe())
        stack.callback(add_compaction_listener(compacted))
        server = await start_gateway_server(config=config, run=True, tool_registry=registry)
        desktop_handoff_written = False
        try:
            if desktop_handoff is not None:
                write_private_desktop_handoff(args.desktop_owner_ready, desktop_handoff)
                desktop_handoff_written = True
            while not (root / "stop").exists():
                publish("running")
                await asyncio.sleep(1)
        finally:
            await server.close()
            if desktop_handoff is not None:
                from opensquilla.gateway.desktop_ownership import (
                    release_active_desktop_gateway_ownership,
                )

                release_active_desktop_gateway_ownership()
                if desktop_handoff_written:
                    args.desktop_owner_ready.unlink(missing_ok=True)
            publish("stopped")


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("command", choices=("prepare", "preflight", "evaluate", "phase", "serve"))
    parser.add_argument("--fixtures", type=Path)
    parser.add_argument("--manifest", type=Path)
    parser.add_argument("--report", type=Path)
    parser.add_argument("--relay-ready", type=Path)
    parser.add_argument("--phase", choices=ATTACHMENT_PHASE_CALL_LIMITS)
    parser.add_argument("--case-id", default="attachment-acceptance")
    parser.add_argument("--variant", choices=("baseline", "new"), default="new")
    parser.add_argument("--live", action="store_true")
    parser.add_argument("--gateway-root", type=Path)
    parser.add_argument("--gateway-port", type=int, default=18799)
    parser.add_argument("--ui-dist", type=Path)
    parser.add_argument("--desktop-owner-ready", type=Path)
    parser.add_argument("--expected-source-root", type=Path)
    parser.add_argument("--model", choices=(*TEXT_MODELS, VISION_MODEL), default=TEXT_MODELS[0])
    parser.add_argument("--context-window", type=int, default=0)
    parser.add_argument("--max-output", type=int, default=1024)
    parser.add_argument("--run-mode", choices=("full", "safe"), default="full")
    args = parser.parse_args(argv)
    if args.expected_source_root and args.expected_source_root.resolve() != REPO_ROOT:
        parser.error("expected-source-root does not match the harness checkout")
    if args.command == "serve" and not args.live:
        print(json.dumps({"ok": False, "status": "live_opt_in_required"}))
        return 2
    if args.command == "evaluate":
        if not all((args.fixtures, args.manifest, args.gateway_root, args.report)):
            parser.error("evaluate requires fixtures, manifest, gateway-root and report")
        result = evaluate(args.gateway_root.resolve(), args.fixtures, args.manifest)
        write_safe_report(require_temporary_report_path(args.report), result, ())
        print(json.dumps({"status": result["status"], "report": str(args.report)}))
        return 0
    if args.command in {"prepare", "preflight"}:
        if not args.fixtures or not args.manifest:
            parser.error("fixtures and manifest are required")
        result = (prepare_fixtures(args.fixtures, args.manifest) if args.command == "prepare"
                  else preflight(args.fixtures, args.manifest))
        summary = {"prepared_cases": len(result["fixtures"]), "physical_calls": 0} if (
            args.command == "prepare"
        ) else result
        if args.report:
            write_safe_report(require_temporary_report_path(args.report), summary, ())
        print(json.dumps(summary))
        return 0 if summary.get("ok", True) else 1
    if not args.relay_ready or not args.phase:
        parser.error("relay-ready and phase are required")
    ready, log = open_relay(args.relay_ready)
    if args.command == "phase":
        log.select_phase(variant=args.variant, case_id=args.case_id, phase=args.phase)
        print(json.dumps({"phase": args.phase, "remaining": log.snapshot()["callsRemaining"]}))
        return 0
    if not args.gateway_root or not args.report or not args.ui_dist:
        parser.error("gateway-root, report, and ui-dist are required")
    if args.context_window < 0 or args.max_output < 1:
        parser.error("invalid context or output limit")
    if not (args.ui_dist / "index.html").is_file():
        parser.error("built WebUI index is required")
    log.select_phase(variant=args.variant, case_id=args.case_id, phase=args.phase)
    env = {**minimal_child_environment(), "TOKENRHYTHM_API_KEY": ready["client_key"],
           "OPENSQUILLA_LIVE_TRANSPORT": "1", "OPENSQUILLA_LIVE_RELAY_URL": ready["base_url"],
           "OPENSQUILLA_LIVE_RELAY_CLIENT_KEY": ready["client_key"],
           "OPENSQUILLA_LIVE_DISABLE_DOTENV": "1", "OPENSQUILLA_TURN_CALL_LOG": "0",
           "OPENSQUILLA_STATE_DIR": str(args.gateway_root / "state"),
           "OPENSQUILLA_USER_STATE_DIR": str(args.gateway_root / "user-state"),
           "OPENSQUILLA_LOG_DIR": str(args.gateway_root / "logs"),
           "OPENSQUILLA_TEST_PROFILE_LOCK_ROOT": "1", "OPENSQUILLA_OPENROUTER_LIVE_PRICING": "0",
           "OPENSQUILLA_CONTROL_UI_DIST": str(args.ui_dist.resolve())}
    from scripts.live_tokenrhythm_transport import install_from_env

    previous_logging = logging.root.manager.disable
    try:
        with patch.dict(os.environ, env, clear=True), contextlib.ExitStack() as stack:
            stack.callback(install_from_env())
            logging.disable(logging.CRITICAL)
            with (
                contextlib.redirect_stdout(io.StringIO()),
                contextlib.redirect_stderr(io.StringIO()),
            ):
                asyncio.run(serve_gateway(args, ready, log))
    except Exception as error:
        write_safe_report(require_temporary_report_path(args.report), {
            "status": "harness_failed", "error_type": type(error).__name__, **source_evidence(),
        }, (ready["client_key"],))
        print(json.dumps({"status": "harness_failed", "report": str(args.report)}))
        return 1
    finally:
        logging.disable(previous_logging)
    print(json.dumps({"status": "stopped", "report": str(args.report)}))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
