"""Bounded, opt-in real Gateway deliverable checks; never used by product code.

Run from an authorized ordinary validation checkout. The real credential is
read without terminal echo, held only by the existing local relay, and never
passed to the Agent subprocess. Results are evidence, not automatic fixes.
"""

from __future__ import annotations

import argparse
import asyncio
import contextlib
import csv
import getpass
import hashlib
import io
import json
import os
import posixpath
import re
import sqlite3
import subprocess
import sys
import time
import uuid
from decimal import Decimal
from html.parser import HTMLParser
from pathlib import Path
from urllib.parse import unquote, urlsplit

import httpx

ROOT = Path(__file__).resolve().parents[1]
sys.path[:0] = [str(ROOT), str(ROOT / "src")]

from opensquilla.artifacts import ArtifactStore  # noqa: E402
from opensquilla.gateway_client import GatewayRPCClient  # noqa: E402
from scripts.live_harness_security import child_environment, sanitize_report  # noqa: E402
from scripts.live_tokenrhythm_budget import (  # noqa: E402
    BudgetRejectedError,
    BudgetRelay,
    FunctionalRequestLog,
)
from scripts.smoke_v4_phase3_router import _free_port, _stop_gateway  # noqa: E402

MODEL = "deepseek-v4-pro-0813"
CASES = [
    (
        "html",
        "html-create",
        "做一个介绍北京四季天气的简洁 HTML 页面，标题是北京四季，"
        "包含春夏秋冬四块内容。写到 beijing-site/index.html，"
        "直接打开预览，不需要下载或发布。不要查询实时天气。",
        [".html"],
        False,
    ),
    (
        "html",
        "html-edit",
        "把刚才页面标题改成北京出行天气指南，增加一条秋季带薄外套的建议，继续打开预览，不导出。",
        [".html"],
        False,
    ),
    (
        "html",
        "html-export",
        "把当前页面导出成一个可以直接发给朋友、离线打开的单独 HTML 文件，给我下载。",
        [".html"],
        True,
    ),
    (
        "html",
        "html-after-export",
        "继续修改工作区里的页面，标题改成北京四季出行备忘，预览更新即可，不重新导出。",
        [".html"],
        False,
    ),
    (
        "multi",
        "html-multipage",
        "做一个小型咖啡店网站，三个 HTML 页面分别为首页、菜单、联系方式，"
        "放在 cafe/ 下，共用一个外置 CSS，每页能相互跳转。"
        "菜单有拿铁 28 元和美式 22 元。打开预览供我看，不导出。每页都给出文件入口。",
        [".html"],
        False,
    ),
    (
        "pdf",
        "pdf-create",
        "做一份两页中文 PDF 项目周报给我下载：项目名海风计划，"
        "第一页本周完成接口联调和首页开发，第二页下周计划测试和上线。"
        "使用简洁版式，不需要图片或联网。",
        [".pdf"],
        True,
    ),
    (
        "pdf",
        "pdf-edit",
        "修改刚才的 PDF，把第二页下周计划中的上线改为灰度发布，给我更新后的 PDF。",
        [".pdf"],
        True,
    ),
    (
        "pptx",
        "pptx-create",
        "做一个三页中文 PPTX 汇报文件给我下载，主题是咖啡店开业计划："
        "第一页目标，第二页预算（装修 3 万、设备 2 万），第三页进度。"
        "简单文字排版即可，不用图片、不联网。",
        [".pptx"],
        True,
    ),
    (
        "md",
        "markdown-create",
        "整理一份新同事入职清单，交付 onboarding.md 文件给我下载。"
        "包含一级标题、3 项待办复选框，以及账号/负责人两列表格。不要仅贴在聊天正文。",
        [".md"],
        True,
    ),
    (
        "csv",
        "csv-create",
        "生成一个 expenses.csv 给我下载，列为 date,category,amount；"
        "三行数据为 2026-09-01,coffee,28；2026-09-02,lunch,35；2026-09-03,transport,12。",
        [".csv"],
        True,
    ),
    (
        "mixed",
        "mixed-delivery",
        "做一组小型活动材料：活动名周末读书会、时间周六下午两点、地点城市书房。"
        "一个 HTML 活动页供预览（不导出网页）；"
        "另提供一页 PDF 邀请函和 README.md 活动说明供下载。"
        "内容保持一致，不联网、不用图片。",
        [".html", ".pdf", ".md"],
        True,
    ),
]


class BoundedRelay(BudgetRelay):
    """Limit requests and output size without rewriting model requests."""

    def __init__(self, *args, model=MODEL, **kwargs):
        super().__init__(*args, **kwargs)
        self.model = model
        self.calls = 0

    @contextlib.contextmanager
    def forward(self, body, headers=None):
        request = json.loads(body)
        with self._lock:
            if self.calls >= 90 or request.get("model") != self.model:
                raise BudgetRejectedError("acceptance_request_limit")
            limits = [
                request[name] for name in ("max_tokens", "max_completion_tokens") if name in request
            ]
            if (
                not limits
                or any(type(limit) is not int or not 1 <= limit <= 8192 for limit in limits)
                or any(limit != limits[0] for limit in limits)
                or type(request.get("n", 1)) is not int
                or request.get("n", 1) != 1
            ):
                raise BudgetRejectedError("acceptance_output_limit")
            self.calls += 1
        with contextlib.ExitStack() as stack:
            try:
                response = stack.enter_context(super().forward(body, headers))
            except BudgetRejectedError:
                # The parent can reject a reservation before any upstream dispatch.
                # Keep the relay counter aligned with requests the ledger accepted.
                with self._lock:
                    self.calls -= 1
                raise
            yield response


def rows(db_path, sql, values=()):
    if not db_path.is_file():
        return []
    with sqlite3.connect(db_path.as_uri() + "?mode=ro", uri=True) as db:
        db.row_factory = sqlite3.Row
        return [dict(row) for row in db.execute(sql, values)]


def listed_artifacts(media_root, session_id, session_key, *, page_size=100):
    """Read every public artifact page; the store excludes internal revision markers."""
    store = ArtifactStore(media_root)
    result = []
    seen = set()
    before = None
    while True:
        page = store.list_refs(session_id=session_id, limit=page_size, before=before)
        for ref in page.refs:
            if ref.id in seen or ref.session_key != session_key:
                raise RuntimeError("artifact catalog identity mismatch")
            seen.add(ref.id)
            result.append(ref.to_dict())
        if not page.has_more:
            return sorted(result, key=lambda ref: (ref["created_at"], ref["id"]))
        if not page.refs or page.refs[0].id == before:
            raise RuntimeError("artifact catalog cursor did not advance")
        before = page.refs[0].id


def snapshot(db, key):
    session = rows(
        db, "SELECT session_id, execution_workspace FROM sessions WHERE session_key=?", (key,)
    )
    documents = rows(db, "SELECT * FROM artifact_documents WHERE session_key=?", (key,))
    publications = rows(
        db,
        "SELECT p.* FROM document_publications p "
        "JOIN artifact_documents d USING(document_id) WHERE d.session_key=?",
        (key,),
    )
    revisions = rows(
        db,
        "SELECT r.* FROM artifact_revisions r "
        "JOIN artifact_documents d USING(document_id) WHERE d.session_key=?",
        (key,),
    )
    sources = rows(
        db,
        "SELECT w.* FROM artifact_working_files w "
        "JOIN artifact_documents d USING(document_id) WHERE d.session_key=?",
        (key,),
    )
    files = []
    binding = (
        json.loads(session[0]["execution_workspace"])
        if session and session[0]["execution_workspace"]
        else None
    )
    if binding:
        root = Path(binding["root"])
        for file in sorted(root.rglob("*")):
            if file.is_file() and not file.is_symlink() and file.stat().st_size < 20_000_000:
                data = file.read_bytes()
                files.append(
                    {
                        "path": str(file.relative_to(root)),
                        "size": len(data),
                        "sha256": hashlib.sha256(data).hexdigest(),
                    }
                )
    return {
        "session": session,
        "documents": documents,
        "publications": publications,
        "revisions": revisions,
        "sources": sources,
        "files": files,
        "workspace": binding,
        "artifacts": (
            listed_artifacts(db.parent.parent / "media", session[0]["session_id"], key)
            if session
            else []
        ),
    }


def validate_bytes(data, name):
    suffix = Path(name).suffix.lower()
    if suffix == ".docx":
        from docx import Document

        doc = Document(io.BytesIO(data))
        paragraphs = [paragraph.text for paragraph in doc.paragraphs]
        tables = [[[cell.text for cell in row.cells] for row in table.rows]
                  for table in doc.tables]
        return {"paragraphs": paragraphs, "tables": tables,
                "text": "\n".join(paragraphs + [" | ".join(row)
                                              for table in tables for row in table])}
    if suffix == ".pdf":
        from pypdf import PdfReader

        doc = PdfReader(io.BytesIO(data))
        return {
            "pages": len(doc.pages),
            "text": "\n".join(page.extract_text() or "" for page in doc.pages),
            "page_texts": [page.extract_text() or "" for page in doc.pages],
        }
    if suffix == ".pptx":
        from pptx import Presentation

        doc = Presentation(io.BytesIO(data))
        return {
            "slides": len(doc.slides),
            "text": "\n".join(
                shape.text for slide in doc.slides for shape in slide.shapes if shape.has_text_frame
            ),
            "slide_texts": [
                "\n".join(shape.text for shape in slide.shapes if shape.has_text_frame)
                for slide in doc.slides
            ],
        }
    if suffix in {".md", ".html", ".csv"}:
        return {"text": data.decode("utf-8-sig")}
    return {}


class _HtmlFacts(HTMLParser):
    """Static content and dependency facts; this does not claim rendered UI coverage."""

    def __init__(self, text):
        super().__init__(convert_charrefs=True)
        self.text = []
        self.headings = []
        self.links = []
        self.styles = []
        self.resources = []
        self.css = []
        self._heading = None
        self._ignored = None
        self.feed(text)

    def handle_starttag(self, tag, attrs):
        attrs = dict(attrs)
        if tag in {"script", "style"}:
            self._ignored = tag
        if tag in {"title", "h1", "h2", "h3", "h4", "h5", "h6"}:
            self._heading = []
        if tag == "a" and attrs.get("href"):
            self.links.append(attrs["href"])
        if tag == "link" and "stylesheet" in (attrs.get("rel") or "").split():
            if attrs.get("href"):
                self.styles.append(attrs["href"])
                self.resources.append(attrs["href"])
        if tag in {"script", "img", "iframe", "source"} and attrs.get("src"):
            self.resources.append(attrs["src"])
        if attrs.get("style"):
            self.css.append(attrs["style"])

    def handle_endtag(self, tag):
        if tag == self._ignored:
            self._ignored = None
        if tag in {"title", "h1", "h2", "h3", "h4", "h5", "h6"}:
            if self._heading is not None:
                self.headings.append("".join(self._heading))
            self._heading = None

    def handle_data(self, data):
        if self._ignored == "style":
            self.css.append(data)
        if self._ignored:
            return
        self.text.append(data)
        if self._heading is not None:
            self._heading.append(data)


def _compact(text):
    return re.sub(r"\s+|\x00", "", text)


def _target(source, reference):
    parsed = urlsplit(reference.strip())
    if parsed.scheme in {"data", "mailto", "tel"}:
        return None
    if parsed.scheme or parsed.netloc:
        return "external:" + reference
    if not parsed.path:
        return None
    return posixpath.normpath(posixpath.join(posixpath.dirname(source), unquote(parsed.path)))


def _css_references(text):
    return [
        match[0] or match[1]
        for match in re.findall(
            r"url\(\s*['\"]?([^)'\"]+)['\"]?\s*\)|@import\s+['\"]([^'\"]+)",
            text,
            re.IGNORECASE,
        )
    ]


def _verified_source_files(case, run_root):
    after = case.get("after", {})
    workspace = after.get("workspace") or {}
    files, unavailable = {}, []
    for item in after.get("files", []):
        name = item["path"]
        if Path(name).suffix.lower() not in {".html", ".htm", ".css", ".js", ".mjs"}:
            continue
        try:
            path = (Path(workspace["root"]) / name).resolve(strict=True)
            if not path.is_relative_to(run_root.resolve()):
                raise ValueError("source outside report run")
            data = path.read_bytes()
            if hashlib.sha256(data).hexdigest() != item.get("sha256"):
                raise ValueError("source no longer matches this historical snapshot")
            files[name] = data.decode("utf-8-sig")
        except (KeyError, OSError, ValueError):
            unavailable.append(name)
    return files, unavailable


def validate_case_content(case, run_root):
    """Check explicit brief facts against current bytes with recorded hash provenance."""
    checks = []

    def check(name, passed, *, actual=None, expected=None):
        checks.append(
            {
                "name": name,
                "status": "inconclusive" if passed is None else "passed" if passed else "failed",
                "actual": actual,
                "expected": expected,
            }
        )

    downloaded = {}
    for item in case.get("downloads", []):
        try:
            path = Path(item["path"]).resolve(strict=True)
            if not path.is_relative_to(run_root.resolve()):
                raise ValueError("download outside report run")
            data = path.read_bytes()
            valid = hashlib.sha256(data).hexdigest() == item.get("sha256")
            check("download_hash", valid, actual=item["name"])
            if valid:
                parsed = validate_bytes(data, item["name"])
                downloaded.setdefault(Path(item["name"]).suffix.lower(), []).append(parsed)
        except Exception as exc:  # A corrupt deliverable must not hide later case evidence.
            check("download_readable", False, actual=type(exc).__name__)

    case_id = case["id"]
    if case_id in {"pdf-create", "pdf-edit", "pptx-create", "markdown-create", "csv-create"}:
        suffix = {
            "pdf-create": ".pdf",
            "pdf-edit": ".pdf",
            "pptx-create": ".pptx",
            "markdown-create": ".md",
            "csv-create": ".csv",
        }[case_id]
        check("required_download", bool(downloaded.get(suffix)), expected=suffix)
        for parsed in downloaded.get(suffix, []):
            text = _compact(parsed.get("text", ""))
            if suffix == ".pdf":
                check("pdf_page_count", parsed["pages"] == 2, actual=parsed["pages"], expected=2)
                change = "灰度发布" if case_id == "pdf-edit" else "上线"
                check(
                    "pdf_required_content",
                    all(
                        word in text
                        for word in ("海风计划", "接口联调", "首页开发", "测试", change)
                    ),
                )
                if case_id == "pdf-edit":
                    check("pdf_release_updated", "灰度发布" in text and "上线" not in text)
                if parsed["pages"] == 2:
                    pages = [_compact(page) for page in parsed["page_texts"]]
                    check(
                        "pdf_page_sections",
                        all(word in pages[0] for word in ("接口联调", "首页开发"))
                        and all(word in pages[1] for word in ("测试", change)),
                    )
            elif suffix == ".pptx":
                check(
                    "pptx_slide_count", parsed["slides"] == 3, actual=parsed["slides"], expected=3
                )
                check(
                    "pptx_sections",
                    all(word in text for word in ("咖啡店", "目标", "预算", "进度")),
                )
                check(
                    "pptx_budget",
                    bool(re.search(r"装修[^\d]{0,10}(?:3万|30,?000)", text))
                    and bool(re.search(r"设备[^\d]{0,10}(?:2万|20,?000)", text)),
                )
                if parsed["slides"] == 3:
                    check(
                        "pptx_section_order",
                        all(
                            word in page
                            for word, page in zip(
                                ("目标", "预算", "进度"), parsed["slide_texts"], strict=True
                            )
                        ),
                    )
            elif suffix == ".md":
                raw = parsed["text"]
                check("markdown_heading", bool(re.search(r"^#\s+\S", raw, re.MULTILINE)))
                count = len(re.findall(r"^\s*[-*+]\s+\[[ xX]\]\s+\S", raw, re.MULTILINE))
                check("markdown_checkboxes", count == 3, actual=count, expected=3)
                lines = [
                    line.strip().strip("|").split("|") for line in raw.splitlines() if "|" in line
                ]
                check(
                    "markdown_table",
                    len(lines) >= 3
                    and any(
                        len(header) == len(separator) == 2
                        and "账号" in header[0]
                        and "负责人" in header[1]
                        and all(re.fullmatch(r"\s*:?-{3,}:?\s*", col) for col in separator)
                        for header, separator in zip(lines, lines[1:], strict=False)
                    ),
                )
            else:
                actual = list(csv.reader(io.StringIO(parsed["text"])))
                expected = [
                    ["date", "category", "amount"],
                    ["2026-09-01", "coffee", "28"],
                    ["2026-09-02", "lunch", "35"],
                    ["2026-09-03", "transport", "12"],
                ]
                check("csv_rows", actual == expected, actual=actual, expected=expected)

    if case_id.startswith("html-") or case_id == "mixed-delivery":
        if case_id == "html-export":
            files = {
                f"download-{index}.html": item["text"]
                for index, item in enumerate(downloaded.get(".html", []))
            }
            unavailable = []
        else:
            files, unavailable = _verified_source_files(case, run_root)
        check("source_snapshot_available", None if unavailable else True, actual=unavailable)
        pages = {
            name: _HtmlFacts(text)
            for name, text in files.items()
            if Path(name).suffix.lower() in {".html", ".htm"}
        }
        if not unavailable:
            check(
                "html_page_count",
                len(pages) == (3 if case_id == "html-multipage" else 1),
                actual=len(pages),
                expected=3 if case_id == "html-multipage" else 1,
            )
            if case_id == "html-multipage" and pages:
                targets = {
                    name: {_target(name, link) for link in page.links}
                    for name, page in pages.items()
                }
                check(
                    "html_page_links",
                    all(set(pages) - {name} <= links for name, links in targets.items()),
                )
                styles = [
                    {_target(name, ref) for ref in page.styles} for name, page in pages.items()
                ]
                shared = set.intersection(*styles)
                check(
                    "html_shared_css",
                    any(name in files and name.endswith(".css") for name in shared if name),
                )
                text = _compact("".join(part for page in pages.values() for part in page.text))
                check(
                    "html_menu_prices",
                    bool(re.search(r"拿铁[^\d]{0,12}28(?!\d)", text))
                    and bool(re.search(r"美式[^\d]{0,12}22(?!\d)", text)),
                )
            for name, page in pages.items():
                references = page.resources + [
                    ref for css in page.css for ref in _css_references(css)
                ]
                dependencies = {_target(name, ref) for ref in references} - {None}
                if case_id == "html-export":
                    check("html_standalone", not dependencies, actual=sorted(dependencies))
                else:
                    check(
                        "html_dependencies",
                        dependencies <= set(files),
                        actual=sorted(dependencies - set(files)),
                    )
                if case_id not in {"html-multipage", "mixed-delivery"}:
                    title = {
                        "html-create": "北京四季",
                        "html-edit": "北京出行天气指南",
                        "html-export": "北京出行天气指南",
                        "html-after-export": "北京四季出行备忘",
                    }[case_id]
                    check(
                        "html_heading",
                        any(title in _compact(h) for h in page.headings),
                        expected=title,
                    )
                    text = _compact("".join(page.text))
                    check("html_seasons", all(word in text for word in ("春", "夏", "秋", "冬")))
                    if case_id != "html-create":
                        check("html_autumn_advice", "秋" in text and "薄外套" in text)
            for name, text in files.items():
                if name.endswith(".css"):
                    dependencies = {_target(name, ref) for ref in _css_references(text)} - {None}
                    check(
                        "css_dependencies",
                        dependencies <= set(files),
                        actual=sorted(dependencies - set(files)),
                    )
        if case_id == "mixed-delivery":
            check("mixed_download_types", all(downloaded.get(ext) for ext in (".pdf", ".md")))
            for pdf in downloaded.get(".pdf", []):
                check("pdf_page_count", pdf["pages"] == 1, actual=pdf["pages"], expected=1)
            materials = [
                item["text"] for ext in (".pdf", ".md") for item in downloaded.get(ext, [])
            ]
            if not unavailable:
                materials.extend("".join(page.text) for page in pages.values())
            for text in materials:
                compact = _compact(text)
                check(
                    "mixed_event_consistency",
                    "周末读书会" in compact
                    and "城市书房" in compact
                    and "周六" in compact
                    and bool(re.search(r"14[:：]00|下午(?:两|二|2)点", compact)),
                )

    if case_id not in {row[1] for row in CASES} or not checks:
        check("case_requirements_known", None, actual=case_id)
    status = (
        "failed"
        if any(item["status"] == "failed" for item in checks)
        else (
            "inconclusive" if any(item["status"] == "inconclusive" for item in checks) else "passed"
        )
    )
    return {
        "id": case_id,
        "status": status,
        "checks": checks,
        "transport_failures": case.get("failures", []),
    }


def plan_cases(repetitions=3):
    """Independent sessions, including repeated core regressions; no silent retry."""
    if type(repetitions) is not int or not 1 <= repetitions <= 10:
        raise ValueError("plan repetitions must be between 1 and 10")
    return [
        {"id": f"{kind}-{number}", "kind": kind, "number": number}
        for kind, count in (("plan-first", repetitions), ("plan-repair", repetitions),
                            ("plan-unverifiable", 1), ("plan-flexible", 1),
                            ("plan-discussion", 1))
        for number in range(1, count + 1)
    ]


def plan_fixture(case):
    """Synthetic, case-local facts. These are not a product-wide year validator."""
    year = 2020 + case["number"]
    return {
        "title": "海岬市示例统计简报", "year": str(year),
        "rows": [["海岬市", str(310 + case["number"]), "万人", "年末常住人口"],
                 ["青湾市", str(420 + case["number"]), "万人", "年末常住人口"]],
        "verifiable": case["kind"] != "plan-unverifiable",
    }


def plan_prompt(case, fixture):
    if case["kind"] == "plan-first":
        return "帮我写个文档介绍华北"
    if case["kind"] == "plan-discussion":
        return "解释一下统计数据的年份和统计口径分别是什么意思，简单举例即可，不需要制定实施方案。"
    facts = "；".join("，".join(row) for row in fixture["rows"])
    source = (f"以下是本次练习可核实的全部资料，均为虚构教学数据："
              f"数据年份 {fixture['year']}，{facts}。") if fixture["verifiable"] else (
        f"以下仅是没有来源、年份未核实的草稿数字：{facts}。没有可核实的原始资料；"
        "这是虚构城市，不得联网猜测或编造年份。"
    )
    flexible = ("已有内容可以复用，生成与校验可以合并，不必机械照搬步骤；"
                "以交付内容满足要求为准。") if case["kind"] == "plan-flexible" else ""
    return (
        f"帮我完成一份 DOCX 统计简报供下载，标题为《{fixture['title']}》，文件名 report.docx。"
        "当前工作区已提供 report.docx 草稿及 source.json；"
        "请先检查草稿和资料，再规划如何依据要求修改。"
        "包含一个标题、一段说明两城市人口数据的正文和一张两城市人口表。"
        "正文和人口表中的每项人口数据都必须明确数据年份、单位与统计口径。"
        "表格可用统一表头说明共同年份与口径，不必机械重复；不能用文末笼统免责声明代替年份。"
        + source + flexible + "请核对实际文件；遇到无法核实的信息应明确说明未满足项。"
    )


def seed_plan_draft(workspace, fixture, evidence_dir):
    """Place a valid but semantically incomplete user input before the planning turn.

    No session/plan rows are manufactured. Keep pre-execution bytes outside the
    Agent workspace so later repairs cannot erase the original defect evidence.
    """
    from docx import Document

    workspace = Path(workspace).resolve(strict=True)
    target = workspace / "report.docx"
    source_path = workspace / "source.json"
    if target.exists() or source_path.exists():
        raise ValueError("fixture input already exists; do not overwrite it")
    document = Document()
    document.add_heading(fixture["title"], 0)
    document.add_paragraph("虚构教学数据；" + "；".join(
        f"{row[0]}{row[3]}为{row[1]}{row[2]}" for row in fixture["rows"]
    ) + "。")
    table = document.add_table(rows=1, cols=4)
    for cell, value in zip(table.rows[0].cells, ["城市", "人口", "单位", "统计口径"]):
        cell.text = value
    for row in fixture["rows"]:
        for cell, value in zip(table.add_row().cells, row):
            cell.text = value
    document.save(target)
    evidence_dir.mkdir(parents=True, exist_ok=True)
    data = target.read_bytes()
    (evidence_dir / "before.docx").write_bytes(data)
    source = {key: value for key, value in fixture.items()
              if key != "year" or fixture["verifiable"]}
    source_data = json.dumps(source, ensure_ascii=False).encode("utf-8")
    source_path.write_bytes(source_data)
    (evidence_dir / "source-before.json").write_bytes(source_data)
    return {"path": str(target), "sha256": hashlib.sha256(data).hexdigest(),
            "bytes": len(data), "parsed": validate_bytes(data, target.name),
            "source": {"path": str(source_path), "sha256": hashlib.sha256(source_data).hexdigest(),
                       "bytes": len(source_data)}}


def read_plan_inputs(client, session_key, evidence_dir, run, label):
    """Record both supplied inputs through the session's public file authority."""
    names = {"report.docx", "source.json"}
    headers = {"x-opensquilla-session-key": session_key}
    response = client.post("/api/v1/workspace-files/resolve",
                           json={"paths": sorted(names)}, headers=headers)
    response.raise_for_status()
    resolution = response.json()
    files = {}
    for item in resolution.get("files", []):
        name, url = item.get("path"), item.get("contentUrl", "")
        if name not in names or name in files:
            raise RuntimeError("fixture file resolution returned unexpected inputs")
        if not url.startswith("/api/v1/workspace-files/content?"):
            raise RuntimeError("fixture file resolution returned an unexpected URL")
        content = client.get(url, headers=headers)
        content.raise_for_status()
        target = evidence_dir / (label + "-" + name)
        target.write_bytes(content.content)
        files[name] = {"sha256": hashlib.sha256(content.content).hexdigest(),
                       "bytes": len(content.content), "evidence_path": str(target.relative_to(run))}
    if set(files) != names:
        raise RuntimeError("fixture inputs are unavailable through the public file API")
    return {"resolution": resolution, "files": files}


def _fixture_has_number(text, expected):
    tokens = re.findall(
        r"(?<![\d.,A-Za-z])[-+]?(?:\d+(?:,\d{3})*(?:\.\d+)?|\.\d+)"
        r"(?:[eE][-+]?\d+)?(?![\d.,A-Za-z])", text,
    )
    return any(Decimal(token.replace(",", "")) == Decimal(expected) for token in tokens)


def check_plan_document(data, fixture):
    """Check this fixture's facts in their paragraph/table context, not document-wide."""
    parsed = validate_bytes(data, "report.docx")
    table_contexts = []
    cities = [row[0] for row in fixture["rows"]]
    for table in parsed["tables"]:
        header_rows = []
        for row in table:
            text = " | ".join(row)
            if any(city in text for city in cities):
                break
            header_rows.append(text)
        table_contexts.append((table, " | ".join(header_rows)))
    missing = []
    missing_body = []
    for expected in fixture["rows"]:
        city, number, unit, basis = expected
        candidates = [(" | ".join(row), header) for table, header in table_contexts
                      for row in table if city in " | ".join(row)]
        if not any(_fixture_has_number(row, number)
                   and _fixture_has_number(row + " | " + header, fixture["year"])
                   and unit in row + header and basis in row + header
                   for row, header in candidates):
            missing.append(expected[0])
        if not any(city in paragraph and unit in paragraph and basis in paragraph
                   and _fixture_has_number(paragraph, number)
                   and _fixture_has_number(paragraph, fixture["year"])
                   for paragraph in parsed["paragraphs"]):
            missing_body.append(expected[0])
    return {
        "structure_passed": bool(parsed["tables"] and any(parsed["paragraphs"])),
        "fixture_requirements_passed": (
            not missing and not missing_body and fixture["title"] in parsed["text"]
        ),
        "rows_missing_required_fact": missing, "body_missing_required_fact": missing_body,
        "parsed": parsed,
    }


def validate_plan_case(case, run_root):
    checks = []
    failures = list(case.get("failures", []))
    if case.get("status") == "acceptance_limit":
        return {"id": case["id"], "status": "inconclusive", "checks": [],
                "transport_failures": failures, "reason": "acceptance_limit",
                "limitations": ["The harness stopped this task; not a natural model completion."]}
    if case["kind"] == "plan-first":
        stages = case.get("stages", [])
        proposal = (stages[-1].get("bootstrap", {}).get("currentPlan") or {}) if stages else {}
        submitted = bool(proposal.get("revisionId")) and bool(
            stages and stages[-1].get("task", {}).get("status") == "succeeded"
        )
        if not submitted:
            failures.append("first_turn_missing_proposal")
        return {"id": case["id"], "status": "failed" if failures else "passed",
                "checks": [{"submitted_proposal": submitted}], "transport_failures": failures,
                "limitations": ["Proposal admission evidence does not assess proposal quality."]}
    for item in case.get("outputs", []):
        path = (Path(run_root) / item["evidence_path"]).resolve()
        if not path.is_relative_to(Path(run_root).resolve()):
            failures.append("evidence_path_outside_run")
            continue
        try:
            data = path.read_bytes()
            if hashlib.sha256(data).hexdigest() != item["sha256"]:
                failures.append("evidence_hash_mismatch")
                continue
            check = {"kind": item.get("kind", "workspace"), "sha256": item["sha256"],
                     "evidence_path": item["evidence_path"],
                     "artifact_id": item.get("artifact_id")}
            try:
                check.update(check_plan_document(data, case["fixture"]))
            except Exception as exc:
                # A bad earlier publication is useful repair evidence, not proof
                # that a later repaired delivery is still bad.
                check.update(structure_passed=False, fixture_requirements_passed=False,
                             parse_error=type(exc).__name__)
            checks.append(check)
        except Exception as exc:
            failures.append("evidence_unreadable:" + type(exc).__name__)
    # The recorder appends downloads in public catalog creation order. Preserve
    # every revision above; select the last delivery, never any passing revision.
    final_workspace = next((check for check in reversed(checks)
                            if check["kind"] == "workspace"), None)
    final_download = next((check for check in reversed(checks)
                           if check["kind"] == "download"), None)
    for check in checks:
        check["final"] = check is final_workspace or check is final_download
    if case["kind"] not in {"plan-discussion", "plan-unverifiable"}:
        if not final_workspace:
            failures.append("final_workspace_missing")
        if not final_download:
            failures.append("requested_download_missing")
        if not all(check and check["fixture_requirements_passed"]
                   for check in (final_workspace, final_download)):
            failures.append("fixture_content_not_satisfied")
    if final_download and (not final_workspace
                           or final_download["sha256"] != final_workspace["sha256"]):
        failures.append("final_download_workspace_mismatch")
    # Tool activity or a successful task cannot settle truthfulness of the final
    # explanation. Preserve a review boundary instead of a keyword-based verdict.
    return {"id": case["id"], "status": "failed" if failures else "inconclusive",
            "checks": checks, "transport_failures": failures,
            "review_required": [
                "Inspect final reply and tool results: was verification and reporting truthful?",
                "Check missing requirements are repaired or explicitly left unresolved.",
                "Check update_plan reports actual milestones without a required 4/4 sequence.",
                "Check the final reply actually links the selected final artifact version.",
            ]}


async def _rpc_call(base_url, method, params):
    client = GatewayRPCClient()
    await client.connect(base_url)
    try:
        return await client.call(method, params)
    finally:
        await client.close()


def plan_rpc(client, method, params):
    return asyncio.run(_rpc_call(str(client.base_url), method, params))


def plan_clarification_answers(request, case=None):
    """Supply declared preferences or fixture facts without inventing missing sources."""
    case_kind = (case or {}).get("kind", "plan-first")
    defaults = "面向普通读者，DOCX 文档，约 2500 字，华北范围为北京、天津、河北、山西、内蒙古。"
    if case_kind in {"plan-repair", "plan-flexible", "plan-unverifiable"}:
        defaults = (
            "report.docx 草稿及 source.json 已在当前工作区，请先读取；"
            "按原请求要求修改，版式等非关键细节可自行合理决定。"
        )
        defaults += ("没有额外来源或可核实年份；不要猜测或编造，不能满足的要求请明确说明。"
                     if case_kind == "plan-unverifiable" else
                     "可核实资料已在原请求和 source.json 中提供，没有额外资料。")
    fields = request.get("clarify_schema", {}).get("fields", [])
    if not fields:
        raise ValueError("clarification missing public field schema")
    answers = {}
    for field in fields:
        name = field["name"]
        kind = field.get("type", "string")
        if case_kind != "plan-first":
            if kind not in {"string", "enum", "choice"} or (
                kind in {"enum", "choice"} and not field.get("allow_other")
            ):
                raise AcceptanceLimitError("fixture clarification requires an undeclared decision")
            answers[name] = defaults
            continue
        if kind == "bool":
            answers[name] = False
        elif kind == "int":
            answers[name] = 2500
        elif kind in {"enum", "choice"} and not field.get("allow_other"):
            choices = field.get("choices", [])
            preferred = next((choice for choice in choices if re.search(
                r"普通|通俗|DOCX|Word|2500|京津冀晋蒙|内蒙古", choice, re.IGNORECASE,
            )), None)
            if preferred is None:
                raise ValueError("clarification choices require manual preference selection")
            answers[name] = preferred
        else:
            answers[name] = defaults
    return answers


def run_plan_cases(client, run, report, selected, request_log, *, key="", phase=None):
    """Drive real public admission/approval; append every attempt before dispatch."""
    completed = {case["id"] for case in report["cases"]}
    for definition in plan_cases(report["repetitions"]):
        if definition["id"] not in selected or definition["id"] in completed:
            continue
        case_id = definition["id"]
        request_log.select_phase(variant="new", case_id=case_id, phase=phase)
        fixture = plan_fixture(definition)
        entry = {**definition, "session_key": "agent:main:webchat:qa" + uuid.uuid4().hex,
                 "fixture": fixture, "prompt": plan_prompt(definition, fixture),
                 "stages": [], "outputs": [], "failures": [], "status": "running",
                 "budget_config": dict(report.get("budget_config") or acceptance_budgets(True))}
        # An interrupted/failed attempt remains in evidence and is not retried on
        # resume; a new run root makes any intentional replay explicit.
        report["cases"].append(entry)
        save_report(run, report, key)
        evidence_dir = run / "plan-evidence" / case_id
        evidence_dir.mkdir(parents=True, exist_ok=True)

        def persist():
            save_report(run, report, key)

        def turn(method, params, label):
            stage = {"label": label, "method": method, "params": params, "snapshots": []}
            entry["stages"].append(stage)
            persist()
            stage["accepted"] = plan_rpc(client, method, params)
            task_id = (stage["accepted"].get("taskId") or stage["accepted"].get("task_id")
                       or stage["accepted"].get("turn_id") or stage["accepted"].get("runId"))
            if not task_id:
                raise RuntimeError("admission returned no task identity")
            persist()
            deadline_seconds = entry["budget_config"]["harness_turn_deadline_s"]
            deadline = time.monotonic() + deadline_seconds
            while time.monotonic() < deadline:
                snapshot_value = plan_rpc(client, "sessions.bootstrap", {
                    "key": entry["session_key"], "limit": 500,
                })
                # Preserve full replies including tool outcomes and final text.
                stage["bootstrap"] = snapshot_value
                tasks = snapshot_value.get("tasks", [])
                tasks = [*tasks, snapshot_value.get("last_task") or {},
                         snapshot_value.get("active_task") or {}]
                task = next((item for item in tasks if item.get("task_id") == task_id), {})
                observation = {"task": task, "plan": snapshot_value.get("currentPlan"),
                               "run": snapshot_value.get("activePlanRun")}
                if not stage["snapshots"] or observation != stage["snapshots"][-1]:
                    stage["snapshots"].append(observation)
                    persist()
                if task.get("status") in {
                    "succeeded", "failed", "cancelled", "timeout", "abandoned",
                }:
                    stage["task"] = task
                    persist()
                    if task["status"] != "succeeded":
                        entry["failures"].append(label + "_turn_not_succeeded")
                    return snapshot_value
                pending_inputs = snapshot_value.get("session", {}).get("pendingUserInputs", [])
                if pending_inputs:
                    answers = stage.setdefault("clarifications", [])
                    answered = {item["request"]["request_id"] for item in answers}
                    for request in pending_inputs:
                        if request["request_id"] in answered:
                            continue
                        if len(answers) >= 3:
                            raise RuntimeError("clarification round limit reached")
                        answer = {"request": request,
                                  "fields": plan_clarification_answers(request, definition)}
                        answers.append(answer)
                        persist()
                        answer["response"] = plan_rpc(client, "chat.clarify_submit", {
                            "sessionKey": entry["session_key"],
                            "requestId": request["request_id"], "fields": answer["fields"],
                        })
                        persist()
                time.sleep(0.5)
            stage["abort"] = plan_rpc(client, "chat.abort", {
                "sessionKey": entry["session_key"], "runId": task_id,
            })
            stage["acceptance_limit"] = {"kind": "wall_time", "seconds": deadline_seconds}
            raise AcceptanceLimitError(label + " exceeded the harness deadline")

        try:
            supplied_draft = definition["kind"] in {
                "plan-repair", "plan-flexible", "plan-unverifiable",
            }
            if supplied_draft:
                # Create only a normal empty session. Its managed workspace exists
                # before any Agent task, so planning can inspect the supplied input.
                creation_params = {"agentId": "main", "kind": "webchat"}
                entry["session_creation"] = {"method": "sessions.create", "params": creation_params}
                persist()
                created = plan_rpc(client, "sessions.create", creation_params)
                entry["session_creation"]["response"] = created
                entry["session_key"] = created["key"]
                persist()
                initial = plan_rpc(client, "sessions.bootstrap", {"key": entry["session_key"]})
                entry["initial_bootstrap"] = initial
                workspace = Path(initial["session"]["workspace"]).resolve(strict=True)
                if not workspace.is_relative_to((run / "profile").resolve()):
                    raise RuntimeError("plan workspace escaped isolated profile")
                if any(previous.get("workspace") == str(workspace)
                       for previous in report["cases"] if previous is not entry):
                    raise RuntimeError("independent cases unexpectedly share a workspace")
                entry["workspace"] = str(workspace)
                entry["seeded_draft"] = seed_plan_draft(workspace, fixture, evidence_dir)
                persist()
                entry["inputs_before_planning"] = read_plan_inputs(
                    client, entry["session_key"], evidence_dir, run, "planning-input",
                )
                for name, expected in (("report.docx", entry["seeded_draft"]),
                                       ("source.json", entry["seeded_draft"]["source"])):
                    actual = entry["inputs_before_planning"]["files"][name]["sha256"]
                    if actual != expected["sha256"]:
                        raise RuntimeError("public fixture input hash mismatch")
                mode_params = {"sessionKey": entry["session_key"], "mode": "plan",
                               "expectedRevision": initial["collaboration"]["revision"]}
                entry["mode_selection"] = {"method": "plans.setMode", "params": mode_params}
                persist()
                entry["mode_selection"]["response"] = plan_rpc(client, "plans.setMode", mode_params)
                mode_response = entry["mode_selection"]["response"]
                if mode_response.get("collaboration", {}).get("mode") != "plan":
                    raise RuntimeError("public mode selection did not enter Plan")
                persist()
            send_params = {"sessionKey": entry["session_key"], "message": entry["prompt"],
                           "intent": "continue" if supplied_draft else "new_chat",
                           "clientRequestId": uuid.uuid4().hex}
            if not supplied_draft:
                send_params["collaborationMode"] = "plan"
            planned = turn("chat.send", send_params, "planning")
            proposal = planned.get("currentPlan")
            if definition["kind"] == "plan-discussion":
                if proposal:
                    entry["failures"].append("discussion_submitted_unrequested_plan")
                entry["status"] = "finished"
                continue
            if not proposal or not proposal.get("revisionId"):
                entry["failures"].append("first_turn_missing_proposal")
                entry["status"] = "finished"
                continue
            if definition["kind"] == "plan-first":
                entry["workspace"] = planned.get("session", {}).get("workspace")
                entry["status"] = "finished"
                continue
            if Path(planned["session"]["workspace"]).resolve(strict=True) != workspace:
                raise RuntimeError("planning changed the supplied input workspace")
            entry["inputs_after_planning"] = read_plan_inputs(
                client, entry["session_key"], evidence_dir, run, "proposed-input",
            )
            if any(entry["inputs_after_planning"]["files"][name]["sha256"]
                   != initial_file["sha256"]
                   for name, initial_file in entry["inputs_before_planning"]["files"].items()):
                entry["failures"].append("planning_modified_supplied_input_before_approval")
                entry["status"] = "finished"
                continue
            persist()
            executed = turn("plans.implement", {
                "sessionKey": entry["session_key"], "planRevisionId": proposal["revisionId"],
                "intent": "continue", "clientRequestId": uuid.uuid4().hex,
            }, "implementation")
            # Reopen the actual output through the ordinary public file API.
            response = client.post("/api/v1/workspace-files/resolve",
                                   json={"paths": ["report.docx"]},
                                   headers={"x-opensquilla-session-key": entry["session_key"]})
            response.raise_for_status()
            entry["workspace_resolution"] = response.json()
            for item in entry["workspace_resolution"].get("files", []):
                url = item.get("contentUrl", "")
                if not url.startswith("/api/v1/workspace-files/content?"):
                    continue
                output = client.get(
                    url, headers={"x-opensquilla-session-key": entry["session_key"]},
                )
                output.raise_for_status()
                target = evidence_dir / "after.docx"
                target.write_bytes(output.content)
                entry["outputs"].append({"kind": "workspace",
                                         "evidence_path": str(target.relative_to(run)),
                                         "sha256": hashlib.sha256(output.content).hexdigest(),
                                         "bytes": len(output.content)})
            catalog = snapshot(run / "profile" / "state" / "sessions.db", entry["session_key"])
            entry["artifact_catalog"] = catalog["artifacts"]
            for artifact in catalog["artifacts"]:
                if Path(artifact["name"]).suffix.lower() != ".docx":
                    continue
                response = client.get("/api/v1/artifacts/" + artifact["id"],
                                      params={"sessionKey": entry["session_key"]})
                response.raise_for_status()
                data = response.content
                digest = hashlib.sha256(data).hexdigest()
                if digest != artifact["sha256"]:
                    entry["failures"].append("download_hash_mismatch")
                target = evidence_dir / ("download-" + artifact["id"] + ".docx")
                target.write_bytes(data)
                entry["outputs"].append({"kind": "download", "artifact_id": artifact["id"],
                                         "created_at": artifact.get("created_at"),
                                         "evidence_path": str(target.relative_to(run)),
                                         "sha256": digest, "bytes": len(data)})
            entry["status"] = "finished"
            entry["final_bootstrap"] = executed
        except Exception as exc:
            entry["status"] = ("acceptance_limit" if isinstance(exc, AcceptanceLimitError)
                               else "error")
            entry["failures"].append(type(exc).__name__ + ":" + str(exc))
            stage = entry["stages"][-1] if entry["stages"] else {}
            accepted = stage.get("accepted", {})
            unfinished_id = (accepted.get("taskId") or accepted.get("task_id")
                             or accepted.get("turn_id") or accepted.get("runId"))
            if unfinished_id and not stage.get("task") and not stage.get("abort"):
                with contextlib.suppress(Exception):
                    stage["abort"] = plan_rpc(client, "chat.abort", {
                        "sessionKey": entry["session_key"], "runId": unfinished_id,
                    })
        finally:
            entry["assessment"] = validate_plan_case(entry, run)
            entry["physical_requests"] = [item for item in request_log.snapshot()["requests"]
                                           if item["case_id"] == case_id]
            if any(item["model"] != report["model"] for item in entry["physical_requests"]):
                entry["failures"].append("unexpected_physical_model")
                entry["assessment"] = validate_plan_case(entry, run)
            persist()
            print(json.dumps({"case": case_id, "status": entry["status"],
                              "assessment": entry["assessment"]["status"],
                              "failures": entry["failures"]}, ensure_ascii=False), flush=True)
        if entry["status"] == "acceptance_limit":
            break  # The caller's finally stops this harness-owned Gateway.
    return plan_report_exit_code(report)


def plan_report_exit_code(report):
    """A transport success is never presented as automatic semantic acceptance."""
    if any(case.get("status") == "acceptance_limit" for case in report["cases"]):
        return 2
    if any(case.get("failures") or case.get("assessment", {}).get("status") == "failed"
           for case in report["cases"]):
        return 1
    return (0 if report["cases"] and all(
        case.get("assessment", {}).get("status") == "passed" for case in report["cases"]
    ) else 2)  # Human review of implementation behavior is still required.


class AcceptanceLimitError(RuntimeError):
    """A test boundary stopped the run, rather than the Agent finishing its task."""


def acceptance_budgets(plan_suite):
    """Keep Plan behavior on product defaults, bounded by external test controls."""
    from opensquilla.engine.types import AgentConfig

    return {
        "agent_max_iterations": 0 if plan_suite else 16,
        "agent_runtime_timeout_seconds": None if plan_suite else 240,
        "agent_max_provider_retries": None if plan_suite else 0,
        "effective_agent_max_provider_retries": (
            AgentConfig().max_provider_retries if plan_suite else 0
        ),
        "provider_retry_policy": "product_default" if plan_suite else "smoke_override",
        "turn_hard_deadline_s": None if plan_suite else 270,
        "llm_request_timeout_seconds": 90,
        "harness_turn_deadline_s": 600 if plan_suite else 285,
    }


def configure_acceptance_budgets(config_text, budgets):
    for name in ("agent_max_iterations", "agent_runtime_timeout_seconds",
                 "agent_max_provider_retries", "turn_hard_deadline_s"):
        value = budgets[name]
        config_text = re.sub(rf"^{name} = [^\n]*\n",
                             "" if value is None else f"{name} = {value}\n",
                             config_text, flags=re.MULTILINE)
    return config_text


def isolated_windows_home(run, *, platform=None):
    """CLI imports need a home on Windows; keep every fallback inside this run."""
    if (os.name if platform is None else platform) != "nt":
        return {}
    private_home = Path(run).resolve() / "host-home"
    roaming = private_home / "AppData" / "Roaming"
    local = private_home / "AppData" / "Local"
    for directory in (private_home, roaming, local):
        directory.mkdir(parents=True, exist_ok=True)
    return {"USERPROFILE": str(private_home), "APPDATA": str(roaming), "LOCALAPPDATA": str(local)}


def prepare_gateway_bootstrap(run):
    """Install the relay gate in the Gateway only, before importing its CLI.

    Ordinary tool Python processes inherit source paths, not sitecustomize.
    Keep old transport directories as evidence, outside the new startup path.
    """
    launcher = Path(run) / "gateway_launcher.py"
    launcher.write_text(
        "try:\n"
        "    from scripts.live_tokenrhythm_transport import install_from_env\n"
        "    install_from_env()\n"
        "except BaseException:\n"
        "    raise SystemExit('acceptance transport unavailable')\n"
        "import runpy\n"
        "runpy.run_module('opensquilla.cli.main', run_name='__main__')\n",
        encoding="utf-8",
    )
    return launcher, os.pathsep.join(map(str, [ROOT / "src", ROOT]))


def validate_report(path):
    """Re-read evidence without network, credentials, execution, or report mutation."""
    path = Path(path).resolve(strict=True)
    raw = path.read_bytes()
    report = json.loads(raw)
    validate = validate_plan_case if report.get("suite") == "plan" else validate_case_content
    cases = [validate(case, path.parent) for case in report["cases"]]
    summary = {
        status: sum(case["status"] == status for case in cases)
        for status in ("passed", "failed", "inconclusive")
    }
    status = (
        "failed"
        if summary["failed"]
        else ("inconclusive" if summary["inconclusive"] or not cases else "passed")
    )
    return {
        "source_report": str(path),
        "source_sha256": hashlib.sha256(raw).hexdigest(),
        "status": status,
        "summary": summary,
        "cases": cases,
        "limitations": [
            "Static content checks do not prove rendering, interaction, or visual layout.",
            "Changed source bytes cannot prove the contents of an earlier snapshot.",
        ],
    }


def save_report(run, report, key=""):
    temporary = run / "report.json.tmp"
    temporary.write_text(
        json.dumps(
            sanitize_report(report, secrets={"TOKENRHYTHM_API_KEY": key}),
            ensure_ascii=False,
            indent=2,
        ),
        encoding="utf-8",
    )
    temporary.chmod(0o600)
    temporary.replace(run / "report.json")


def restored_sessions(report):
    groups = {case_id: group for group, case_id, *_rest in CASES}
    sessions = dict(report.get("sessions", {}))
    for entry in report.get("cases", []):
        if entry["id"] in groups:
            sessions[groups[entry["id"]]] = entry["session_key"]
    pending = report.get("pending_case")
    if pending and pending["id"] in groups:
        sessions[groups[pending["id"]]] = pending["session_key"]
    return sessions


def load_report(run, *, resume, model=MODEL):
    path = run / "report.json"
    if not resume:
        if path.exists():
            raise ValueError("run already has a report; use --resume")
        return {"model": model, "cases": [], "attempts": []}
    if not path.is_file():
        raise ValueError("--resume requires an existing report.json")
    report = json.loads(path.read_text(encoding="utf-8"))
    if report.get("model") != model or not isinstance(report.get("cases"), list):
        raise ValueError("resume report model or cases do not match this runner")
    prior = {
        name: report[name]
        for name in (
            "source",
            "source_commit",
            "provider_calls",
            "runner_error",
            "http_rejection",
            "blocked",
            "provider_probe_status",
            "budget_config",
        )
        if name in report
    }
    prior["case_ids"] = [entry["id"] for entry in report["cases"]]
    prior["archived_at"] = time.time()
    report.setdefault("attempts", []).append(prior)
    for name in ("runner_error", "http_rejection", "blocked"):
        report.pop(name, None)
    return report


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--run-root", type=Path)
    parser.add_argument("--validate-report", type=Path)
    parser.add_argument("--enable-live", action="store_true")
    parser.add_argument("--cases", nargs="*", default=[])
    parser.add_argument("--resume", action="store_true")
    parser.add_argument("--plan-suite", action="store_true",
                        help="Plan first-turn, approved execution, repair and discussion cases")
    parser.add_argument("--repetitions", type=int, default=3)
    parser.add_argument("--model", default=MODEL)
    parser.add_argument("--run-mode", choices=("safe", "full"), default="safe")
    parser.add_argument("--relay-ready", type=Path,
                        help="Existing functional relay ready JSON; no desktop credential read")
    parser.add_argument("--relay-phase", help="Existing relay budget allocation, if required")
    args = parser.parse_args()
    if args.validate_report is not None:
        if args.enable_live or args.run_root or args.resume or args.cases or args.plan_suite:
            parser.error("--validate-report is an independent offline mode")
        result = validate_report(args.validate_report)
        print(json.dumps(result, ensure_ascii=False, indent=2))
        return {"passed": 0, "failed": 1, "inconclusive": 2}[result["status"]]
    if args.run_root is None:
        parser.error("live runs require --run-root")
    if not args.enable_live:
        parser.error("real provider tests require --enable-live")
    run = args.run_root.resolve()
    if ".codex" in run.parts or str(run).startswith(
        ("/tmp/", "/private/tmp/", "/private/var/folders/")
    ):
        parser.error("use an authorized ordinary validation directory")
    if not re.fullmatch(r"[A-Za-z0-9][A-Za-z0-9._:/+-]{0,199}", args.model):
        parser.error("invalid model identifier")
    known_cases = ({case["id"] for case in plan_cases(args.repetitions)} if args.plan_suite
                   else {case_id for _group, case_id, *_rest in CASES})
    if set(args.cases) - known_cases:
        parser.error("unknown case id")
    run.mkdir(mode=0o700, parents=True, exist_ok=True)
    report = load_report(run, resume=args.resume, model=args.model)
    suite = "plan" if args.plan_suite else "deliverable"
    if args.resume and report.get("suite", "deliverable") != suite:
        parser.error("cannot resume a different acceptance suite")
    if args.resume and report.get("run_mode", "safe") != args.run_mode:
        parser.error("cannot silently change run mode when resuming")
    report.update(suite=suite, repetitions=args.repetitions, run_mode=args.run_mode)
    report["budget_config"] = acceptance_budgets(args.plan_suite)
    report["source"] = str(ROOT)
    report["source_commit"] = subprocess.check_output(
        ["git", "rev-parse", "HEAD"], cwd=ROOT, text=True
    ).strip()
    report["source_diff_sha256"] = hashlib.sha256(subprocess.check_output(
        ["git", "diff", "HEAD", "--", "src", "scripts"], cwd=ROOT,
    )).hexdigest()
    sessions = restored_sessions(report)
    report["sessions"] = sessions
    completed = {entry["id"] for entry in report["cases"]}
    selected = set(args.cases) or known_cases
    save_report(run, report)
    if selected <= completed:
        return (plan_report_exit_code(report) if args.plan_suite
                else int(any(entry.get("failures") for entry in report["cases"])))
    relay = None
    if args.relay_ready:
        from scripts.live_tokenrhythm_transport import RelayTarget

        ready = json.loads(args.relay_ready.read_text(encoding="utf-8"))
        if ready.get("mode") != "functional" or ready.get("enabled") is not True:
            raise ValueError("functional_relay_required")
        RelayTarget(str(ready["base_url"]), str(ready["client_key"]))
        relay_url, client_key = ready["base_url"], ready["client_key"]
        key = client_key
        request_log = FunctionalRequestLog(Path(ready["request_log"]), enabled=True,
                                           max_calls=int(ready.get("max_calls", 60)))
        report["external_relay"] = True
    else:
        key = getpass.getpass("TokenRhythm test key (not saved): ")
        if not key.startswith("sk_"):
            raise ValueError("unexpected credential format")
        request_log = FunctionalRequestLog(run / "requests.sqlite3", enabled=True)
        relay = BoundedRelay(None, {}, api_key=key, model=args.model, request_log=request_log)
        relay.calls = max(len(request_log.snapshot()["requests"]),
                          int(report.get("provider_calls", 0)))
        relay_url = relay.start()
        client_key = relay.client_key
    request_log.select_phase(variant="new", case_id="setup", phase=args.relay_phase)
    report["budget_config"]["physical_request_max_calls"] = request_log.max_calls
    proc = None
    try:
        # Credential validation is a read-only catalog request before spending tokens.
        with httpx.Client(timeout=30, trust_env=False) as http:
            response = http.get(
                relay_url + "/models", headers={"Authorization": "Bearer " + client_key}
            )
            report["provider_probe_status"] = response.status_code
            print(
                json.dumps({"phase": "provider_probe", "http_status": response.status_code}),
                flush=True,
            )
            if response.status_code != 200:
                report["blocked"] = "provider_auth_or_catalog_failed"
                return 2
            available = {row["id"] for row in response.json().get("data", [])}
            if args.model not in available:
                report["blocked"] = "configured_model_not_available"
                return 2
        profile = run / "profile"
        profile.mkdir(mode=0o700, exist_ok=True)
        launcher, pythonpath = prepare_gateway_bootstrap(run)
        report["transport_bootstrap"] = "gateway_only"
        config = run / "gateway.toml"
        config.write_text(configure_acceptance_budgets(
            ('''host = "127.0.0.1"
debug = false
llm_request_timeout_seconds = 90
agent_runtime_timeout_seconds = 240
agent_max_iterations = 16
agent_max_provider_retries = 0
[auth]
mode = "none"
[control_ui]
enabled = true
[rate_limit]
enabled = false
[privacy]
disable_network_observability = true
[naming]
enabled = false
[memory]
source = "state"
[sandbox]
run_mode = "safe"
approvals_reviewer = "user"
[tools]
profile = "full"
deny = ["sessions_spawn", "sessions_send"]
[task_runtime]
turn_hard_deadline_s = 270
[llm]
provider = "tokenrhythm"
model = "'''
            + args.model
            + """"
api_key_env = "TOKENRHYTHM_API_KEY"
base_url = "https://tokenrhythm.studio/v1"
max_tokens = 8192
thinking = "off"
[squilla_router]
enabled = false
""").replace('run_mode = "safe"', 'run_mode = "' + args.run_mode + '"'),
            report["budget_config"],
        ))
        env = child_environment(
            "tokenrhythm", {"TOKENRHYTHM_API_KEY": client_key}, base_environment=os.environ
        )
        env.update(isolated_windows_home(run))
        env.update(
            {
                "PYTHONPATH": pythonpath,
                "PATH": str(Path(sys.executable).parent)
                + os.pathsep
                + env.get("PATH", "/usr/bin:/bin"),
                "OPENSQUILLA_GATEWAY_CONFIG_PATH": str(config),
                "OPENSQUILLA_STATE_DIR": str(profile),
                "OPENSQUILLA_USER_STATE_DIR": str(run / "user-state"),
                "OPENSQUILLA_TEST_PROFILE_LOCK_ROOT": "1",
                "OPENSQUILLA_MEMORY_DREAM_DISABLED": "1",
                "OPENSQUILLA_TURN_CALL_LOG": "1",
                "OPENSQUILLA_TURN_CALL_LOG_DIR": str(run / "turn-calls"),
                "OPENSQUILLA_LIVE_TRANSPORT": "1",
                "OPENSQUILLA_LIVE_RELAY_URL": relay_url,
                "OPENSQUILLA_LIVE_RELAY_CLIENT_KEY": client_key,
                "OPENSQUILLA_LIVE_DISABLE_DOTENV": "1",
            }
        )
        port = _free_port()
        with (run / "gateway.log").open("a" if args.resume else "w") as log:
            proc = subprocess.Popen(
                [
                    sys.executable,
                    str(launcher),
                    "gateway",
                    "run",
                    "--port",
                    str(port),
                    "--bind",
                    "127.0.0.1",
                ],
                cwd=run,
                env=env,
                stdout=log,
                stderr=log,
            )
        db = profile / "state" / "sessions.db"
        client = httpx.Client(base_url=f"http://127.0.0.1:{port}", timeout=15, trust_env=False)
        for _ in range(90):
            try:
                if client.get("/api/system/status").status_code == 200:
                    break
            except httpx.HTTPError:
                pass
            if proc.poll() is not None:
                raise RuntimeError("gateway startup failed; inspect gateway.log")
            time.sleep(0.5)
        else:
            raise RuntimeError("gateway readiness timeout")
        print(json.dumps({"phase": "gateway_ready", "url": str(client.base_url)}), flush=True)
        report["gateway_url"] = str(client.base_url)
        if args.plan_suite:
            try:
                return run_plan_cases(client, run, report, selected, request_log,
                                      key=key, phase=args.relay_phase)
            finally:
                client.close()
        downloaded = {
            item["artifact_id"]: (entry["session_key"], item["sha256"])
            for entry in report["cases"]
            for item in entry.get("downloads", [])
        }
        for group, case_id, prompt, extensions, delivery in CASES:
            if case_id not in selected or case_id in completed:
                continue
            key_session = sessions.setdefault(
                group, "agent:main:webchat:qa" + uuid.uuid4().hex[:10]
            )
            request_log.select_phase(variant="new", case_id=case_id, phase=args.relay_phase)
            pending = report.get("pending_case")
            if pending and pending["id"] != case_id:
                raise RuntimeError("resume the pending case before selecting another case")
            old = pending["before"] if pending else snapshot(db, key_session)
            if not pending:
                pending = {
                    "id": case_id,
                    "session_key": key_session,
                    "before": old,
                    "client_request_id": uuid.uuid4().hex,
                    "intent": "continue" if old["session"] else "new_chat",
                }
                report["pending_case"] = pending
                save_report(run, report, key)
            accepted = pending.get("accepted")
            if not accepted:
                response = client.post(
                    "/api/chat",
                    json={
                        "sessionKey": key_session,
                        "message": prompt,
                        "intent": pending["intent"],
                        "clientRequestId": pending["client_request_id"],
                    },
                )
                if response.is_error:
                    try:
                        rejection = response.json()
                    except ValueError:
                        rejection = {"message": response.text[:2000]}
                    report["http_rejection"] = sanitize_report(
                        {"case_id": case_id, "status": response.status_code, "response": rejection},
                        secrets={"TOKENRHYTHM_API_KEY": key},
                    )
                    response.raise_for_status()
                accepted = response.json()
                pending["accepted"] = accepted
                save_report(run, report, key)
            task_id = accepted.get("taskId") or accepted.get("task_id")
            if not task_id:
                raise RuntimeError("admission did not return a task identity")
            print(
                json.dumps(
                    {
                        "case": case_id,
                        "phase": "accepted",
                        "session": key_session,
                        "receipt": accepted,
                    },
                    ensure_ascii=False,
                ),
                flush=True,
            )
            started = time.monotonic()
            task = None
            while time.monotonic() - started < 285:
                tasks = rows(
                    db,
                    "SELECT task_id,status,terminal_reason,error_class,created_at "
                    "FROM agent_tasks WHERE session_key=? AND task_id=?",
                    (key_session, task_id),
                )
                if tasks and tasks[0]["status"] in {
                    "succeeded",
                    "failed",
                    "cancelled",
                    "timeout",
                    "abandoned",
                }:
                    task = tasks[0]
                    break
                time.sleep(1)
            if task is None:
                raise RuntimeError("turn did not settle within deadline: " + case_id)
            history = client.get(
                "/api/chat/history", params={"sessionKey": key_session, "limit": 500}
            ).json()
            current = snapshot(db, key_session)
            entry = {
                "id": case_id,
                "session_key": key_session,
                "prompt": prompt,
                "task": task,
                "seconds": round(time.monotonic() - started, 1),
                "before": old,
                "after": current,
                "downloads": [],
                "failures": [],
            }
            (run / (case_id + ".history.json")).write_text(
                json.dumps(
                    sanitize_report(history, secrets={"TOKENRHYTHM_API_KEY": key}),
                    ensure_ascii=False,
                    indent=2,
                )
            )
            if task["status"] != "succeeded":
                entry["failures"].append("turn_not_completed")
            existing = {p["publication_id"] for p in old["publications"]}
            new = [p for p in current["publications"] if p["publication_id"] not in existing]
            if not delivery and new:
                entry["failures"].append("unexpected_publication")
            previous_artifacts = {ref["id"] for ref in old.get("artifacts", [])}
            new_artifacts = [
                ref for ref in current["artifacts"] if ref["id"] not in previous_artifacts
            ]
            if not delivery and new_artifacts:
                entry["failures"].append("unexpected_listed_artifact")
            for artifact in new_artifacts:
                artifact_id = artifact["id"]
                data_response = client.get(
                    "/api/v1/artifacts/" + artifact_id, params={"sessionKey": key_session}
                )
                data_response.raise_for_status()
                data = data_response.content
                name = artifact["name"]
                target = run / "downloads" / case_id / Path(name).name
                target.parent.mkdir(parents=True, exist_ok=True)
                target.write_bytes(data)
                info = {
                    "artifact_id": artifact_id,
                    "name": name,
                    "path": str(target),
                    "sha256": hashlib.sha256(data).hexdigest(),
                    "bytes": len(data),
                    "mime": data_response.headers.get("content-type"),
                    "parsed": validate_bytes(data, name),
                }
                if info["sha256"] != artifact["sha256"]:
                    entry["failures"].append("download_hash_mismatch")
                if len(data) != artifact["size"]:
                    entry["failures"].append("download_size_mismatch")
                if info["mime"].split(";", 1)[0] != artifact["mime"]:
                    entry["failures"].append("download_mime_mismatch")
                entry["downloads"].append(info)
                downloaded[artifact_id] = (key_session, info["sha256"])
            wanted_delivery = set(extensions) - (
                {".html"} if case_id == "mixed-delivery" else set()
            )
            if delivery and not wanted_delivery <= {
                Path(d["name"]).suffix for d in entry["downloads"]
            }:
                entry["failures"].append("missing_download_type")
            if ".html" in extensions and not delivery and not current["sources"]:
                entry["failures"].append("missing_source_preview")
            for source in current["sources"]:
                response = client.get(
                    f"/api/v1/artifact-documents/{source['document_id']}/working-file",
                    headers={"x-opensquilla-session-key": key_session},
                )
                if response.status_code != 200:
                    entry["failures"].append("current_source_unreadable")
            for artifact_id, (owner, expected) in downloaded.items():
                check = client.get("/api/v1/artifacts/" + artifact_id, params={"sessionKey": owner})
                if (
                    check.status_code != 200
                    or hashlib.sha256(check.content).hexdigest() != expected
                ):
                    entry["failures"].append("previous_download_changed")
            report["cases"].append(entry)
            report.pop("pending_case", None)
            (run / "report.json").write_text(
                json.dumps(
                    sanitize_report(report, secrets={"TOKENRHYTHM_API_KEY": key}),
                    ensure_ascii=False,
                    indent=2,
                )
            )
            print(
                json.dumps(
                    {
                        "case": case_id,
                        "phase": "settled",
                        "status": task["status"],
                        "failures": entry["failures"],
                        "files": [f["path"] for f in current["files"]],
                        "downloads": [d["name"] for d in entry["downloads"]],
                    },
                    ensure_ascii=False,
                ),
                flush=True,
            )
        client.close()
        return int(any(row["failures"] for row in report["cases"]))
    except Exception as exc:
        report["runner_error"] = str(exc)
        print(
            json.dumps(sanitize_report({"error": str(exc)}, secrets={"TOKENRHYTHM_API_KEY": key})),
            flush=True,
        )
        return 2
    finally:
        if proc is not None:
            _stop_gateway(proc)
        if relay is not None:
            relay.close()
        report["provider_calls"] = len(request_log.snapshot()["requests"])
        save_report(run, report, key)


if __name__ == "__main__":
    raise SystemExit(main())
